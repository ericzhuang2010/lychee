#!/usr/bin/env python3
"""Build and validate the AgriGenomics India 2026 abstract submission file."""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


HERE = Path(__file__).resolve().parent
TEXT_PATH = HERE / "agrigenomics_india_2026_abstract.txt"
OUTPUT_PATH = HERE / "agrigenomics_india_2026_abstract.docx"

EMAIL = "eric.zhuang@nyulangone.org"
ORCID = "0009-0001-9050-0214"
REQUIRED_LABELS = ("Objective:", "Methodology:", "Key Results:", "Conclusion:")
SPECIES = ("Peronophythora litchii",)


def parse_source(text: str) -> tuple[str, str, str, str]:
    title_match = re.search(r"^Title: (.+)$", text, re.MULTILINE)
    author_match = re.search(r"^Author: (.+)$", text, re.MULTILINE)
    affiliation_match = re.search(r"^Affiliation: (.+)$", text, re.MULTILINE)
    abstract_match = re.search(r"^Abstract \(\d+ words\):\n\n(.+)$", text, re.MULTILINE | re.DOTALL)
    if not all((title_match, author_match, affiliation_match, abstract_match)):
        raise ValueError("The source text is missing title, author, affiliation, or abstract metadata")
    return (
        title_match.group(1).strip(),
        author_match.group(1).strip(),
        affiliation_match.group(1).strip(),
        abstract_match.group(1).strip(),
    )


def validate(title: str, author: str, affiliation: str, abstract: str) -> int:
    if author != "Eric Zhuang":
        raise ValueError(f"Unexpected author: {author}")
    if affiliation != "NYU Langone Health, New York, NY, USA":
        raise ValueError(f"Unexpected affiliation: {affiliation}")
    if "\n" in abstract:
        raise ValueError("The conference abstract must be a single paragraph")
    positions = []
    for label in REQUIRED_LABELS:
        if abstract.count(label) != 1:
            raise ValueError(f"Required label must appear exactly once: {label}")
        positions.append(abstract.index(label))
    if positions != sorted(positions):
        raise ValueError("Required abstract sections are not in the expected order")
    word_count = len(abstract.split())
    if word_count > 200:
        raise ValueError(f"Abstract has {word_count} words; the conference limit is 200")
    if "validation" in abstract.lower():
        raise ValueError("Use external evaluation or cross-context support, not unqualified validation")
    if "CRISPR" in title or "artificial intelligence" in abstract.lower():
        raise ValueError("The title or abstract attributes methods that this study did not use")
    return word_count


def set_font(run, size: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_text_with_species_italics(paragraph, text: str, size: float, bold: bool = False) -> None:
    pattern = re.compile("(" + "|".join(re.escape(item) for item in SPECIES) + ")")
    for part in pattern.split(text):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_font(run, size=size, bold=bold, italic=part in SPECIES)


def build_document(title: str, author: str, affiliation: str, abstract: str, word_count: int) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(8)
    add_text_with_species_italics(title_paragraph, title, size=14, bold=True)

    author_paragraph = document.add_paragraph()
    author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_paragraph.paragraph_format.space_after = Pt(2)
    set_font(author_paragraph.add_run(author), size=12, bold=True)

    affiliation_paragraph = document.add_paragraph()
    affiliation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation_paragraph.paragraph_format.space_after = Pt(2)
    set_font(affiliation_paragraph.add_run(affiliation), size=11)

    contact_paragraph = document.add_paragraph()
    contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_paragraph.paragraph_format.space_after = Pt(12)
    set_font(contact_paragraph.add_run(f"Email: {EMAIL} | ORCID: {ORCID}"), size=10)

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(4)
    set_font(heading.add_run("Abstract"), size=12, bold=True)

    abstract_paragraph = document.add_paragraph()
    abstract_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_paragraph.paragraph_format.line_spacing = 1.15
    abstract_paragraph.paragraph_format.space_after = Pt(8)
    add_text_with_species_italics(abstract_paragraph, abstract, size=12)

    count_paragraph = document.add_paragraph()
    count_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(count_paragraph.add_run(f"Abstract word count: {word_count}"), size=10, italic=True)

    document.core_properties.title = title
    document.core_properties.author = author
    document.core_properties.subject = "AgriGenomics India 2026 conference abstract"
    document.core_properties.keywords = "lychee; Peronophythora litchii; transcriptomics; plant-pathogen interaction"
    return document


def main() -> None:
    source = TEXT_PATH.read_text(encoding="utf-8")
    title, author, affiliation, abstract = parse_source(source)
    word_count = validate(title, author, affiliation, abstract)
    document = build_document(title, author, affiliation, abstract, word_count)
    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")
    print(f"Validated abstract word count: {word_count}/200")


if __name__ == "__main__":
    main()
