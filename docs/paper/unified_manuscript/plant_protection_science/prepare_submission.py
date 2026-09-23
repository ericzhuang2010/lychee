#!/usr/bin/env python3
"""Build the Plant Protection Science submission package from manuscript v3.

The script uses the journal's 2026 CAAS Word templates retained under
source/templates. It creates the blinded manuscript, title page, cover letter,
and prefilled declaration. The declaration intentionally leaves the date and
handwritten signature blank.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
UNIFIED = ROOT.parent
TEMPLATES = ROOT / "source" / "templates"
FIGURES = ROOT / "figures"

MANUSCRIPT_TEMPLATE = TEMPLATES / "Manuscript-file_template_CAAS.docx"
TITLE_TEMPLATE = TEMPLATES / "Title_page_template_CAAS.docx"
LETTER_TEMPLATE = TEMPLATES / "Acompanying-letter_template_CAAS.docx"
DECLARATION_TEMPLATE = TEMPLATES / "Authors_Declaration_2026.docx"

MANUSCRIPT_OUT = ROOT / "PPS_blinded_manuscript.docx"
TITLE_OUT = ROOT / "PPS_title_page.docx"
LETTER_OUT = ROOT / "PPS_cover_letter.docx"
DECLARATION_OUT = ROOT / "PPS_Authors_Declaration_SIGN_AND_SCAN.docx"

TITLE = (
    "Differences among lychee cultivars in responses to Peronophythora litchii "
    "implicate cell wall and carbohydrate remodeling"
)
AUTHOR = "Eric Zhuang"
AFFILIATION = "NYU Langone Health, New York, New York, USA"
EMAIL = "eric.zhuang@nyulangone.org"
ORCID = "0009-0001-9050-0214"

ABSTRACT = (
    "Litchi downy blight, caused by the oomycete Peronophythora litchii C.C. Chen ex W.H. Ko, "
    "H.S. Chang, H.J. Su, C.C. Chen & L.S. Leu, damages lychee (Litchi chinensis Sonn.) leaves, "
    "flowers and fruit. We reanalysed five public RNA-sequencing cohorts to identify cultivar-dependent "
    "infection responses and test their transfer across tissues and times. Discovery in Guiwei and "
    "Yurong1 leaves tested 19 445 expressed genes. Two hundred and six genes passed genome-wide testing, "
    "mappability and gene-model quality control; 16 remained stable under alternative quantification, "
    "statistical analysis, expression filtering, leave-one-library-out analysis and mapping checks. "
    "Twelve internally robust genes formed the highest supported tier and included candidates involved "
    "in sugar transport, starch metabolism, cell-wall remodelling, lipid signalling and RNA metabolism. "
    "In an independent pericarp time course, LITCHI005518, encoding an alpha-xylosidase, and LITCHI001963 "
    "showed same-direction cultivar-dependent effects, whereas five genes reversed direction. A frozen "
    "206-gene signature changed between leaf and fruit and across infection times. Most of 18 earlier "
    "candidates responded within at least one cultivar, but none passed the genome-wide cultivar-interaction "
    "criterion. These results distinguish general infection responses from cultivar-dependent effects and "
    "prioritise testable targets for experimental validation."
)

KEYWORDS = (
    "oomycete disease; host response; differential expression; transcript usage; "
    "reproducible bioinformatics; crop improvement"
)

ITALIC_TERMS = (
    "Peronophythora litchii",
    "Litchi chinensis",
    "P. litchii",
    "in silico",
)


INTRODUCTION = [
    (
        "Lychee (Litchi chinensis Sonn.) is an important subtropical fruit crop. Litchi downy blight, "
        "caused by Peronophythora litchii, affects young leaves, inflorescences and fruit, and cultivars "
        "differ in disease response. Published work has described pathogen virulence machinery, host "
        "ethylene disruption and cultivar-specific early transcriptional programmes (Sun et al. 2017; "
        "Sun et al. 2019; Li et al. 2023). However, most RNA-sequencing studies analyse a single cohort, "
        "so prominent candidates can reflect tissue, time or sampling context rather than stable cultivar differences."
    ),
    (
        "A cultivar-by-infection interaction asks whether the infection response differs between cultivars; "
        "it is not equivalent to testing infection within each cultivar. We first explored Guiwei and Yurong1 "
        "leaves and nominated 18 defence-related genes. We then registered a separate genome-wide analysis "
        "before computing confirmatory outcomes. Dataset roles, models, thresholds, internal robustness gates, "
        "external evaluations and evidence tiers were fixed in advance."
    ),
    (
        "Here we integrate five public cohorts to identify quality-controlled cultivar-dependent responses, "
        "evaluate their robustness and transport across biological contexts, and distinguish supported, null "
        "and opposite-direction outcomes. We also assess transcript usage, pathway and promoter evidence, and "
        "provide a reproducible candidate resource for targeted validation."
    ),
]


METHODS = [
    (
        "Study design and data. The exploratory stage preceded registration. The confirmatory protocol was "
        "frozen on 18 August 2026 and amendments were time-stamped. Five public cohorts had non-interchangeable "
        "roles (Table 1; Figure 1A): PRJNA830488/GSE201243 was the discovery cohort; PRJNA450886 was the primary "
        "cross-context evaluation; PRJNA922966/GSE222651 tested generic infection transfer; PRJNA922965/GSE222650 "
        "provided small-RNA data; and PRJNA1090613/GSE262200 remained exploratory because time and resistance "
        "metadata were unresolved. All inferences are conditional on independence of deposited libraries, which "
        "could not be verified from source-tree, pooling or extraction metadata."
    ),
    (
        "Preprocessing and discovery. The host SCAU_Lch_v2.0 assembly (Hu et al. 2022) and pathogen reference were combined. Reads were trimmed with fastp "
        "(Chen et al. 2018), aligned with STAR (Dobin et al. 2013), counted with featureCounts (Liao et al. 2014) "
        "and independently quantified with Salmon (Patro et al. 2017). Libraries required at least 10 million "
        "surviving read pairs and 40% uniquely aligned reads. DESeq2 fitted gene counts as cultivar + treatment + "
        "cultivar:treatment (Love et al. 2014). Genes with counts of at least 10 in three libraries were tested; "
        "candidates required genome-wide Benjamini-Hochberg q < 0.05 and absolute interaction log2 fold change "
        "at least log2(1.5) (Benjamini & Hochberg 1995). Exon mappability was assessed with GenMap (Pockrandt et al. 2020)."
    ),
    (
        "Robustness and external evaluation. Frozen genes were retested using Salmon gene-level counts, edgeR "
        "quasi-likelihood inference (Robinson et al. 2010), a counts-per-million filter and every leave-one-library-out "
        "refit, followed by an observed-mapping check. The primary external model estimated the Heiye-minus-Guiwei "
        "difference in infection response at 24 h in pericarp. Passing results required family-adjusted q < 0.05, "
        "absolute log2 fold change at least log2(1.5), a confidence interval excluding zero and discovery-direction "
        "agreement; significant opposite-direction effects were recorded separately."
    ),
    (
        "Additional analyses. A frozen signed score summarised the 206 genes across external contrasts. Pathways "
        "were mapped through one-to-one Oryza orthologues to Plant Reactome (Naithani et al. 2024). Transcript usage "
        "was tested with DRIMSeq and stageR (Nowicka & Robinson 2016; Van den Berge et al. 2017). Annotation used "
        "Swiss-Prot DIAMOND matches, InterPro and Oryza orthology (Buchfink et al. 2015; Jones et al. 2014). "
        "Promoter analysis tested JASPAR plant motifs "
        "against expression- and GC-matched backgrounds (Rauluseviciute et al. 2024); FIMO provided sensitivity "
        "testing (Grant et al. 2011). The PmiREN archive was audited before small-RNA integration (Guo et al. 2020)."
    ),
    (
        "Evidence integration and reproducibility. Deterministic rules combined discovery, mapping, internal "
        "robustness, primary external evidence and orthogonal evidence without additive scoring. Tier A required "
        "all classes; Tier B required discovery and complete internal robustness with partial or non-testable "
        "external evidence; Tier C was reserved for a robust, externally supported pathway. Snakemake workflows "
        "enforced staged analysis and SHA-256 manifests (Mölder et al. 2021). Complete outputs, source data and exact "
        "software versions are supplied in the supplementary archive."
    ),
]


RESULTS = [
    (
        "Discovery quality and genome-wide interaction. All 12 discovery libraries passed technical gates: read "
        "survival was 99.56-99.66%, unique alignment 84.44-86.22% and Salmon mapping 80.77-82.56% (Figure 1B). "
        "Principal components separated cultivar and infection status (Figure 1C). Among 19 445 expressed genes, "
        "262 met the registered interaction threshold; uniform mappability and gene-model control retained 206 "
        "candidates (Figure 2A,C). Recomputing adjustment only among 13 602 QC-eligible genes recovered all 206, "
        "showing that post-test quality control produced a conservative subset."
    ),
    (
        "Exploratory candidates and internal robustness. Thirteen of 18 earlier candidates were infection responsive "
        "in Guiwei and 16 in Yurong1 at genome-wide q < 0.05, but none met the separate cultivar-interaction criterion "
        "(Figure 2B). Two failed uniform mappability. Of the 206 frozen genes, 196 passed independent quantification, "
        "19 passed the edgeR gate, 205 passed expression filtering and 125 passed all leave-one-out refits. Eighteen "
        "passed these four gates jointly and 16 remained after observed-mapping checks (Figure 3A). DESeq2 and edgeR "
        "interaction estimates correlated at r = 0.995 genome-wide."
    ),
    (
        "Cross-context outcomes. In the primary pericarp evaluation, 184 candidates were measurable. LITCHI005518 "
        "and LITCHI001963 met every criterion in the discovery direction, while five genes met magnitude and "
        "significance criteria in the opposite direction (Table 2; Figure 3C). Discovery and external effects were "
        "uncorrelated (Pearson r = -0.05), and 95 of 184 shared direction. The 16 internally robust genes and two "
        "cross-context-supported genes did not overlap; for sets of these sizes, the probability of zero overlap "
        "under independence is 0.85 (Figure 3B)."
    ),
    (
        "Signatures, pathways and transcript usage. The frozen 206-gene score tracked infection positively in leaf "
        "(0.107, q = 0.0015) and negatively in fruit (-0.181, q < 0.001), with a significant tissue interaction "
        "(0.289, q < 0.001). In pericarp it changed from 0.109 at 24 h to -0.240 at 48 h (q = 0.043), demonstrating "
        "tissue and temporal dependence (Figure 4B). Six pathways were discovered, but none passed primary external "
        "criteria; circadian rhythm alone passed all internal gates (Figure 4A). Conditional analysis identified 225 "
        "transcript-usage events across 152 genes; none passed the primary cross-context threshold (Figure 5)."
    ),
    (
        "Orthogonal evidence and tiers. Family-level labels were assigned to 150 candidates. No motif passed the "
        "registered 80-of-100 matched-background rule, and small-RNA coherence was not testable because PmiREN had no "
        "exact Litchi entry (Figure 6A,B). The frozen signature used apeglm-shrunken weights (Zhu et al. 2019), and "
        "pathway robustness included competitive testing with camera (Wu & Smyth 2012). Evidence integration produced "
        "no Tier A gene or Tier C pathway, placed 12 "
        "internally robust genes in Tier B, retained 191 entities for exploratory follow-up and excluded 65 through "
        "predefined quality or direction rules (Figure 6C). Tier B included genes linked to sugar transport, starch "
        "and cell-wall metabolism, membrane and lipid transport, chaperone regulation, RNA metabolism and specialised metabolism."
    ),
]


DISCUSSION = [
    (
        "The analysis separates infection responsiveness from cultivar-dependent response. Most exploratory genes "
        "responded within at least one cultivar, yet none passed the genome-wide interaction criterion. This is not "
        "a contradiction: a strong response within a cultivar can coexist with an imprecise difference between "
        "cultivars, especially with three deposited libraries per cell. Mapping control also identified two signals "
        "that could not be assigned unambiguously to one locus."
    ),
    (
        "The prioritised genes point to cell-wall and carbohydrate dynamics. Tier B contains a phosphomannomutase/"
        "phosphoglucomutase, a starch-branching enzyme, an ERD6-like sugar transporter and a serine carboxypeptidase; "
        "the cross-context set includes alpha-xylosidase LITCHI005518, which acts on xyloglucan. This convergence is "
        "consistent with previous lychee work emphasising sugar metabolism and wall reinforcement (Sun et al. 2019). "
        "Other Tier B annotations implicate DIR1-like lipid transfer, BAG-family chaperone regulation and RNA metabolism "
        "(Berka et al. 2022); "
        "these remain hypotheses rather than demonstrated resistance mechanisms."
    ),
    (
        "The external results show that cultivar-associated responses depend strongly on tissue and time. Two genes "
        "transported across different cultivar and tissue contrasts, whereas five reversed direction and the broader "
        "signature changed between leaf and fruit and between 24 and 48 h. Thus below-threshold or opposite-direction "
        "external outcomes should not be described as failed technical replication; they delimit the contexts in which "
        "a response may be useful."
    ),
    (
        "Limitations follow from the public designs. Biological independence of deposited libraries is undocumented; "
        "inoculum, developmental stage and wounding metadata are incomplete; one reference assembly was used; and no "
        "same-tissue, same-time independent cohort was available. None of the candidates has been perturbed experimentally. "
        "Independent infections, qPCR and functional analysis of LITCHI005518 and the Tier B genes are therefore required "
        "before causal or breeding claims."
    ),
]


CONCLUSION = (
    "A registered multi-cohort analysis identified 206 quality-controlled cultivar-interaction candidates, a stable "
    "core of 16 genes, 12 internally robust Tier B genes and two genes supported across an independent tissue and "
    "cultivar contrast. The results implicate cell-wall and carbohydrate remodelling while showing that broader "
    "responses vary with tissue and infection time. The released evidence tables and workflows provide a transparent "
    "basis for targeted experimental validation rather than causal claims."
)


FIGURE_CAPTIONS = [
    "Figure 1. Study cohorts and discovery quality. (A) Five public cohorts with fixed roles. (B) Per-library read survival, STAR unique alignment and Salmon mapping; all 12 discovery libraries passed. (C) Principal components separate cultivar and infection status.",
    "Figure 2. Genome-wide discovery and reassessment of exploratory candidates. (A) Interaction effects across 19 445 genes; 262 met the statistical threshold and 206 remained after quality control. (B) Interaction estimates for 18 exploratory candidates. (C) Attrition to the internally robust set.",
    "Figure 3. Internal robustness and primary external evaluation. (A) Retention by each robustness gate. (B) The 16 internally robust and two cross-context-supported genes do not overlap. (C) Discovery effects against 24-h pericarp effects; the diagonal is a visual reference.",
    "Figure 4. Pathway and signature evaluation. (A) Six discovery pathways; circadian rhythm passed all internal gates, but none passed external criteria. (B) The frozen 206-gene signature across six non-exploratory external contrasts; filled points denote q < 0.05.",
    "Figure 5. Conditional transcript usage. (A) Discovery gate and 225 events across 152 genes. (B) External follow-up separated by prespecified study role.",
    "Figure 6. Annotation, orthogonal evidence and prioritisation. (A) Annotation of 206 candidates. (B) Motif robustness and small-RNA reference gate. (C) Deterministic tiers over 268 frozen entities.",
]


FIGURE_FILES = [
    "Figure_1_study_design_and_QC.tif",
    "Figure_2_discovery_and_legacy_audit.tif",
    "Figure_3_robustness_and_external_evaluation.tif",
    "Figure_4_pathway_and_signature_evaluation.tif",
    "Figure_5_transcript_usage.tif",
    "Figure_6_orthogonal_evidence_and_tiers.tif",
]

FIGURE_ALT_TEXT = [
    "Study-design diagram and quality-control summaries for five public cohorts and twelve discovery libraries.",
    "Genome-wide interaction effects, exploratory-candidate audit and attrition to the internally robust gene set.",
    "Internal robustness gates, overlap with cross-context support and comparison of discovery with pericarp effects.",
    "Discovery-pathway robustness and frozen 206-gene signature results across external contrasts.",
    "Conditional transcript-usage discovery counts and external follow-up outcomes by study role.",
    "Candidate annotation, motif and small-RNA evidence gates, and final evidence-tier counts.",
]


TABLE1 = [
    ["Accession", "Design", "Fixed role"],
    ["PRJNA830488 / GSE201243", "Guiwei and Yurong1 leaf; 24 h", "Discovery"],
    ["PRJNA450886", "Guiwei and Heiye pericarp; 6, 24, 48 h", "Primary external"],
    ["PRJNA922966 / GSE222651", "Feizixiao leaf and fruit; 24 h", "Generic transfer"],
    ["PRJNA922965 / GSE222650", "Feizixiao small RNA; 24 h", "Orthogonal modality"],
    ["PRJNA1090613 / GSE262200", "Guiwei and SFZ leaf", "Exploratory only"],
]


TABLE2 = [
    ["Gene", "Annotation", "Discovery log2FC", "Discovery q", "External log2FC (95% CI)", "External q", "Outcome"],
    ["LITCHI005518", "Alpha-xylosidase 1", "+0.64", "0.025", "+0.67 (+0.32 to +1.02)", "0.007", "Supported"],
    ["LITCHI001963", "Uncharacterised; 1:1 Oryza orthologue", "-1.41", "0.024", "-0.94 (-1.43 to -0.46)", "0.007", "Supported"],
    ["LITCHI030761", "Caffeoylshikimate esterase", "-0.92", "0.040", "+1.31 (+0.78 to +1.85)", "0.0002", "Reversed"],
    ["LITCHI029534", "Thiamine thiazole synthase 2", "+1.48", "0.002", "-1.04 (-1.61 to -0.48)", "0.010", "Reversed"],
    ["LITCHI021860", "Phosphatidylinositol 4-kinase gamma 2", "-0.65", "0.032", "+0.97 (+0.39 to +1.54)", "0.029", "Reversed"],
    ["LITCHI008313", "Beta-glucosidase 44", "+0.90", "0.045", "-1.12 (-1.81 to -0.44)", "0.034", "Reversed"],
    ["LITCHI001700", "Fe(II)/2-OG-dependent dioxygenase 4", "+3.70", "0.033", "-2.00 (-3.24 to -0.77)", "0.035", "Reversed"],
]


REFERENCES = [
    "Benjamini Y., Hochberg Y. (1995): Controlling the false discovery rate: A practical and powerful approach to multiple testing. Journal of the Royal Statistical Society Series B, 57: 289-300.",
    "Berka M., Kopecká R., Berková V., Brzobohatý B., Černý M. (2022): Regulation of heat shock proteins 70 and their role in plant immunity. Journal of Experimental Botany, 73: 1894-1909.",
    "Buchfink B., Xie C., Huson D.H. (2015): Fast and sensitive protein alignment using DIAMOND. Nature Methods, 12: 59-60.",
    "Chen S., Zhou Y., Chen Y., Gu J. (2018): fastp: An ultra-fast all-in-one FASTQ preprocessor. Bioinformatics, 34: i884-i890.",
    "Dobin A., Davis C.A., Schlesinger F., Drenkow J., Zaleski C., Jha S. et al. (2013): STAR: Ultrafast universal RNA-seq aligner. Bioinformatics, 29: 15-21.",
    "Grant C.E., Bailey T.L., Noble W.S. (2011): FIMO: Scanning for occurrences of a given motif. Bioinformatics, 27: 1017-1018.",
    "Guo Z., Kuang Z., Wang Y., Zhao Y., Tao Y., Cheng C. et al. (2020): PmiREN: A comprehensive encyclopedia of plant miRNAs. Nucleic Acids Research, 48: D1114-D1121.",
    "Hu G., Feng J., Xiang X., Wang J., Salojärvi J., Liu C. et al. (2022): Two divergent haplotypes from a highly heterozygous lychee genome suggest independent domestication events for early and late-maturing cultivars. Nature Genetics, 54: 73-83.",
    "Jones P., Binns D., Chang H.Y., Fraser M., Li W., McAnulla C. et al. (2014): InterProScan 5: Genome-scale protein function classification. Bioinformatics, 30: 1236-1240.",
    "Li P., Li W., Zhou X., Situ J., Xie L., Xi P. et al. (2023): Peronophythora litchii RXLR effector PlAvh202 destabilizes a host ethylene biosynthesis enzyme. Plant Physiology, 193: 756-774.",
    "Liao Y., Smyth G.K., Shi W. (2014): featureCounts: An efficient general-purpose program for assigning sequence reads to genomic features. Bioinformatics, 30: 923-930.",
    "Love M.I., Huber W., Anders S. (2014): Moderated estimation of fold change and dispersion for RNA-seq data with DESeq2. Genome Biology, 15: 550.",
    "Mölder F., Jablonski K.P., Letcher B., Hall M.B., Tomkins-Tinch C.H., Sochat V. et al. (2021): Sustainable data analysis with Snakemake. F1000Research, 10: 33.",
    "Naithani S., Gupta P., Preece J., D'Eustachio P., Elser J., Garg P. et al. (2024): Plant Reactome knowledgebase: Empowering plant pathway exploration and OMICS data analysis. Nucleic Acids Research, 52: D1538-D1547.",
    "Nowicka M., Robinson M.D. (2016): DRIMSeq: A Dirichlet-multinomial framework for multivariate count outcomes in genomics. F1000Research, 5: 1356.",
    "Patro R., Duggal G., Love M.I., Irizarry R.A., Kingsford C. (2017): Salmon provides fast and bias-aware quantification of transcript expression. Nature Methods, 14: 417-419.",
    "Pockrandt C., Alzamel M., Iliopoulos C.S., Reinert K. (2020): GenMap: Ultra-fast computation of genome mappability. Bioinformatics, 36: 3687-3692.",
    "Rauluseviciute I., Riudavets-Puig R., Blanc-Mathieu R., Castro-Mondragon J.A., Ferenc K., Kumar V. et al. (2024): JASPAR 2024: 20th anniversary of the open-access database of transcription factor binding profiles. Nucleic Acids Research, 52: D174-D182.",
    "Robinson M.D., McCarthy D.J., Smyth G.K. (2010): edgeR: A Bioconductor package for differential expression analysis of digital gene expression data. Bioinformatics, 26: 139-140.",
    "Sun J., Cao L., Li H., Wang G., Wang S., Li F. et al. (2019): Early responses given distinct tactics to infection of Peronophythora litchii in susceptible and resistant litchi cultivar. Scientific Reports, 9: 2810.",
    "Sun J., Gao Z., Zhang X., Zou X., Cao L., Wang J. (2017): Transcriptome analysis of Phytophthora litchii reveals pathogenicity arsenals and confirms taxonomic status. PLoS ONE, 12: e0178245.",
    "Van den Berge K., Soneson C., Robinson M.D., Clement L. (2017): stageR: A general stage-wise method for controlling gene-level false discovery rate in differential expression and transcript usage. Genome Biology, 18: 151.",
    "Wu D., Smyth G.K. (2012): Camera: A competitive gene set test accounting for inter-gene correlation. Nucleic Acids Research, 40: e133.",
    "Zhu A., Ibrahim J.G., Love M.I. (2019): Heavy-tailed prior distributions for sequence count data: Removing the noise and preserving large differences. Bioinformatics, 35: 2084-2092.",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size: float = 12, bold=None, italic=None, color="000000") -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def write_rich_text(paragraph, text: str, size: float = 12, bold: bool = False) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    pattern = re.compile("(" + "|".join(re.escape(term) for term in ITALIC_TERMS) + ")")
    for part in pattern.split(text):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold, italic=part in ITALIC_TERMS)


def configure_paragraph(paragraph, *, line=1.5, before=0, after=6, justify=True) -> None:
    paragraph.style = paragraph.part.document.styles["Normal"]
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.widow_control = True
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def suppress_line_numbers(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:suppressLineNumbers")) is None:
        p_pr.append(OxmlElement("w:suppressLineNumbers"))


def iter_container_paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_container_paragraphs(cell)


def add_body(doc: Document, text: str, after: float = 6):
    p = doc.add_paragraph()
    configure_paragraph(p, line=1.5, after=after)
    write_rich_text(p, text)
    return p


def add_heading(doc: Document, text: str, major: bool = True):
    p = doc.add_paragraph()
    configure_paragraph(p, line=1.5, before=10 if major else 6, after=3, justify=False)
    p.paragraph_format.keep_with_next = True
    write_rich_text(p, text.upper() if major else text, bold=True)
    return p


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_table(doc: Document, rows, widths, font_size=9):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ridx, values in enumerate(rows):
        row = table.rows[ridx]
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if ridx == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for cidx, value in enumerate(values):
            cell = row.cells[cidx]
            cell.width = Inches(widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if ridx == 0:
                set_cell_shading(cell, "D9EAF7")
            elif ridx % 2 == 0:
                set_cell_shading(cell, "F4F8FB")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            write_rich_text(p, str(value), size=font_size, bold=(ridx == 0))
    return table


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, before=3, after=8)
    write_rich_text(p, text, size=10)
    return p


def set_inline_shape_alt(shape, title: str, description: str) -> None:
    doc_pr = shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_line_numbering(section) -> None:
    sectPr = section._sectPr
    old = sectPr.find(qn("w:lnNumType"))
    if old is not None:
        sectPr.remove(old)
    node = OxmlElement("w:lnNumType")
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:restart"), "continuous")
    node.set(qn("w:distance"), "360")
    sectPr.append(node)


def configure_doc_defaults(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        for paragraph in iter_container_paragraphs(section.header):
            suppress_line_numbers(paragraph)
            for doc_pr in paragraph._p.xpath(".//wp:docPr"):
                doc_pr.set("title", "CAAS logo")
                doc_pr.set("descr", "Czech Academy of Agricultural Sciences logo")
        for paragraph in iter_container_paragraphs(section.footer):
            suppress_line_numbers(paragraph)


def scrub_core_properties(doc: Document, blinded: bool) -> None:
    props = doc.core_properties
    props.author = "" if blinded else AUTHOR
    props.last_modified_by = "" if blinded else AUTHOR
    props.title = TITLE
    props.subject = "Plant Protection Science submission"
    props.keywords = KEYWORDS
    props.comments = ""
    props.category = "Original paper"


def scrub_blinded_package(path: Path) -> None:
    """Remove custom properties and revision-session identifiers from a DOCX."""
    temporary = path.with_name(path.stem + ".privacy-scrubbed.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            if info.filename == "docProps/custom.xml":
                continue
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"<dc:creator(?:\s[^>]*)?>.*?</dc:creator>|<dc:creator(?:\s[^>]*)?/>",
                    b"<dc:creator></dc:creator>",
                    data,
                    flags=re.DOTALL,
                )
                data = re.sub(
                    rb"<cp:lastModifiedBy(?:\s[^>]*)?>.*?</cp:lastModifiedBy>|<cp:lastModifiedBy(?:\s[^>]*)?/>",
                    b"<cp:lastModifiedBy></cp:lastModifiedBy>",
                    data,
                    flags=re.DOTALL,
                )
            elif info.filename == "_rels/.rels":
                data = re.sub(
                    rb'<Relationship\b[^>]*Type="[^"]*/custom-properties"[^>]*/>',
                    b"",
                    data,
                )
            elif info.filename == "[Content_Types].xml":
                data = re.sub(
                    rb'<Override\b[^>]*PartName="/docProps/custom.xml"[^>]*/>',
                    b"",
                    data,
                )
            elif info.filename.startswith("word/") and info.filename.endswith(".xml"):
                data = re.sub(rb'\s+w:rsid[A-Za-z0-9]+="[^"]*"', b"", data)
            target.writestr(info, data)
    temporary.replace(path)


def add_manuscript_front(doc: Document) -> None:
    p = doc.add_paragraph()
    configure_paragraph(p, line=1.5, after=6, justify=False)
    r = p.add_run("Manuscript type: ")
    set_run_font(r)
    r = p.add_run("original paper")
    set_run_font(r, bold=True)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.5, after=8, justify=False)
    p.paragraph_format.keep_with_next = True
    write_rich_text(p, TITLE, size=14, bold=True)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.5, after=6)
    r = p.add_run("Abstract: ")
    set_run_font(r, bold=True)
    for term_or_text in re.compile("(" + "|".join(re.escape(t) for t in ITALIC_TERMS) + ")").split(ABSTRACT):
        if not term_or_text:
            continue
        r = p.add_run(term_or_text)
        set_run_font(r, italic=term_or_text in ITALIC_TERMS)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.5, after=10)
    r = p.add_run("Keywords: ")
    set_run_font(r, bold=True)
    r = p.add_run(KEYWORDS)
    set_run_font(r)


def build_manuscript() -> int:
    doc = Document(MANUSCRIPT_TEMPLATE)
    clear_body(doc)
    configure_doc_defaults(doc)
    add_line_numbering(doc.sections[0])
    scrub_core_properties(doc, blinded=True)
    add_manuscript_front(doc)

    add_heading(doc, "Introduction")
    for text in INTRODUCTION:
        add_body(doc, text)

    add_heading(doc, "Material and Methods")
    for text in METHODS:
        label, rest = text.split(". ", 1)
        p = doc.add_paragraph()
        configure_paragraph(p, line=1.5, after=6)
        r = p.add_run(label + ". ")
        set_run_font(r, bold=True)
        for part in re.compile("(" + "|".join(re.escape(t) for t in ITALIC_TERMS) + ")").split(rest):
            if not part:
                continue
            r = p.add_run(part)
            set_run_font(r, italic=part in ITALIC_TERMS)

    add_heading(doc, "Results")
    add_body(doc, RESULTS[0])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / "Figure_1_study_design_and_QC.tif"), width=Inches(6.0))
    set_inline_shape_alt(shape, "Figure 1", FIGURE_ALT_TEXT[0])
    add_caption(doc, FIGURE_CAPTIONS[0])
    add_caption(doc, "Table 1. Fixed dataset roles and eligibility.")
    add_table(doc, TABLE1, [1.55, 2.85, 1.65], font_size=8.5)

    add_body(doc, RESULTS[1])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / "Figure_2_discovery_and_legacy_audit.tif"), width=Inches(6.0))
    set_inline_shape_alt(shape, "Figure 2", FIGURE_ALT_TEXT[1])
    add_caption(doc, FIGURE_CAPTIONS[1])

    add_body(doc, RESULTS[2])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / "Figure_3_robustness_and_external_evaluation.tif"), width=Inches(6.0))
    set_inline_shape_alt(shape, "Figure 3", FIGURE_ALT_TEXT[2])
    add_caption(doc, FIGURE_CAPTIONS[2])
    doc.add_page_break()
    add_caption(doc, "Table 2. Supported and opposite-direction genes in the primary external evaluation. Effects are cultivar-by-infection interaction log2 fold changes; external intervals are 95% confidence intervals.")
    add_table(doc, TABLE2, [0.8, 1.15, 0.65, 0.55, 1.3, 0.55, 0.65], font_size=7.5)

    add_body(doc, RESULTS[3])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / "Figure_4_pathway_and_signature_evaluation.tif"), width=Inches(6.0))
    set_inline_shape_alt(shape, "Figure 4", FIGURE_ALT_TEXT[3])
    add_caption(doc, FIGURE_CAPTIONS[3])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / "Figure_5_transcript_usage.tif"), width=Inches(6.0))
    set_inline_shape_alt(shape, "Figure 5", FIGURE_ALT_TEXT[4])
    add_caption(doc, FIGURE_CAPTIONS[4])

    add_body(doc, RESULTS[4])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIGURES / "Figure_6_orthogonal_evidence_and_tiers.tif"), width=Inches(6.0))
    set_inline_shape_alt(shape, "Figure 6", FIGURE_ALT_TEXT[5])
    add_caption(doc, FIGURE_CAPTIONS[5])

    add_heading(doc, "Discussion")
    for text in DISCUSSION:
        add_body(doc, text)

    add_heading(doc, "Conclusion")
    add_body(doc, CONCLUSION)

    add_heading(doc, "References")
    for ref in REFERENCES:
        p = doc.add_paragraph()
        configure_paragraph(p, line=1.0, after=4)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        write_rich_text(p, ref, size=10)

    doc.save(MANUSCRIPT_OUT)
    scrub_blinded_package(MANUSCRIPT_OUT)
    return manuscript_character_count(doc)


def manuscript_character_count(doc: Document) -> int:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return len("\n".join(parts))


def add_labelled_paragraph(doc, label: str, value: str, *, italic_value=False, after=6):
    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, after=after, justify=False)
    r = p.add_run(label)
    set_run_font(r, bold=True)
    r = p.add_run(value)
    set_run_font(r, italic=italic_value)
    return p


def build_title_page(character_count: int) -> None:
    doc = Document(TITLE_TEMPLATE)
    clear_body(doc)
    configure_doc_defaults(doc)
    scrub_core_properties(doc, blinded=False)

    add_labelled_paragraph(doc, "Manuscript type: ", "original paper", after=8)
    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, after=8, justify=False)
    write_rich_text(p, TITLE, size=14, bold=True)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, after=6, justify=False)
    r = p.add_run(AUTHOR + "*")
    set_run_font(r, bold=True)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, after=6, justify=False)
    r = p.add_run(AFFILIATION)
    set_run_font(r, italic=True)

    add_labelled_paragraph(doc, "Authors' data", "", after=3)
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    header_pr.append(header_flag)
    widths = [1.35, 1.25, 2.25, 1.5]
    values = [
        ["First name(s)", "Surname(s)", "E-mail", "ORCID"],
        ["Eric", "Zhuang", EMAIL, ORCID],
    ]
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 70, 70, 70, 70)
            set_cell_borders(cell)
            if ri == 0:
                set_cell_shading(cell, "D9EAF7")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            write_rich_text(p, values[ri][ci], size=9, bold=(ri == 0))

    add_labelled_paragraph(doc, "*Corresponding author: ", EMAIL, italic_value=True, after=6)
    add_labelled_paragraph(
        doc,
        "Citation (publication details to be assigned): ",
        f"Zhuang E.: {TITLE}. Plant Protection Science.",
        after=6,
    )

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, after=6)
    r = p.add_run("Abstract: ")
    set_run_font(r, bold=True)
    for part in re.compile("(" + "|".join(re.escape(t) for t in ITALIC_TERMS) + ")").split(ABSTRACT):
        if not part:
            continue
        r = p.add_run(part)
        set_run_font(r, size=10, italic=part in ITALIC_TERMS)

    add_labelled_paragraph(doc, "Number of characters including spaces: ", f"{character_count:,}", after=6)
    add_labelled_paragraph(
        doc,
        "Acknowledgement and AI disclosure: ",
        "OpenAI Codex (GPT-5; accessed 21 September 2026) assisted with language editing, manuscript condensation and structuring, and journal-specific document formatting. It did not generate or manipulate research data, results, interpretations, conclusions, figures, images or graphical elements. The author reviewed and takes responsibility for all content.",
        after=5,
    )
    add_labelled_paragraph(doc, "Funding acknowledgement statement: ", "This work received no external funding.", after=5)
    add_labelled_paragraph(doc, "Conflict of interest: ", "The author declares no conflict of interest.", after=5)
    add_labelled_paragraph(
        doc,
        "Ethics statement: ",
        "This reanalysis used public plant sequencing data and involved no new human, animal or field experiments; ethics approval was not required.",
        after=5,
    )
    add_labelled_paragraph(
        doc,
        "Data and code availability: ",
        "Data accessions, complete results, figure source data, the registered protocol and code are archived at https://zenodo.org/records/22436625 and supplied in the supplementary archive.",
        after=0,
    )
    doc.save(TITLE_OUT)


def replace_header_label(doc: Document, new_label: str) -> None:
    for section in doc.sections:
        for p in section.header.paragraphs:
            if "ACCOMPANYING LETTER" in p.text.upper():
                for run in p.runs:
                    upper = run.text.upper()
                    if "ACCOMPANYING LETTER" in upper:
                        run.text = new_label
                    elif "TEMPLATE" in upper:
                        run.text = ""


def build_cover_letter() -> None:
    doc = Document(LETTER_TEMPLATE)
    clear_body(doc)
    configure_doc_defaults(doc)
    replace_header_label(doc, "COVER LETTER")
    scrub_core_properties(doc, blinded=False)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, after=8, justify=False)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    write_rich_text(p, "21 September 2026")

    for line in (
        "Professor Aleš Lebeda",
        "Editor-in-Chief",
        "Plant Protection Science",
        "Czech Academy of Agricultural Sciences",
    ):
        p = doc.add_paragraph()
        configure_paragraph(p, line=1.0, after=0, justify=False)
        write_rich_text(p, line)

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, before=8, after=8, justify=False)
    write_rich_text(p, f"Re: Original paper - {TITLE}", bold=True)

    add_body(doc, "Dear Professor Lebeda,")
    add_body(
        doc,
        "Please consider this manuscript for publication as an original paper in Plant Protection Science. "
        "The study addresses a plant-protection problem within the journal's scope: cultivar-dependent responses "
        "of lychee to the downy-blight oomycete Peronophythora litchii.",
    )
    add_body(
        doc,
        "The main contribution is a prospectively registered, multi-cohort analysis that separates infection "
        "responses within cultivars from formal cultivar-by-infection effects. Across five public RNA-sequencing "
        "cohorts, it identifies 206 quality-controlled candidates, a stable 16-gene core, 12 internally robust "
        "Tier B genes and two genes supported across a different tissue and cultivar contrast. It also reports null "
        "and opposite-direction outcomes, revealing strong tissue and time dependence and preventing overstatement "
        "of single-cohort candidates. Cell-wall and carbohydrate remodelling emerge as focused hypotheses for validation.",
    )
    add_body(
        doc,
        "The manuscript is original, has not been published as a peer-reviewed article and is not under consideration "
        "elsewhere. The sole author approved the submitted version and accepts responsibility for the work. The study "
        "used public plant sequencing data and required no human or animal ethics approval. All data accessions, full "
        "result tables, figure source data, protocol, code and software specifications are available in the supplement "
        "and the cited Zenodo record. The manuscript is blinded, and the title page is separate.",
    )
    add_body(
        doc,
        "OpenAI Codex (GPT-5) assisted only with language editing, manuscript condensation and structuring, and "
        "journal-specific document formatting. It was "
        "not used to generate or manipulate research data, results, interpretations, conclusions or figures. The "
        "author reviewed all content and accepts full responsibility.",
    )
    add_body(doc, "Thank you for your consideration.")

    p = doc.add_paragraph()
    configure_paragraph(p, line=1.0, before=8, after=0, justify=False)
    write_rich_text(p, "Sincerely,")
    for line in (AUTHOR, AFFILIATION, EMAIL, f"ORCID: {ORCID}"):
        p = doc.add_paragraph()
        configure_paragraph(p, line=1.0, after=0, justify=False)
        write_rich_text(p, line)
    doc.save(LETTER_OUT)


def set_paragraph_text(paragraph, text: str, size=9, bold=False) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    write_rich_text(paragraph, text, size=size, bold=bold)


def build_declaration() -> None:
    doc = Document(DECLARATION_TEMPLATE)
    scrub_core_properties(doc, blinded=False)
    for section in doc.sections:
        for paragraph in iter_container_paragraphs(section.header):
            for doc_pr in paragraph._p.xpath(".//wp:docPr"):
                doc_pr.set("title", "CAAS logo")
                doc_pr.set("descr", "Czech Academy of Agricultural Sciences logo")
    paragraphs = doc.paragraphs
    set_paragraph_text(paragraphs[5], f"Manuscript title: {TITLE}", size=10)
    set_paragraph_text(paragraphs[6], f"Authors (all authors' full names): {AUTHOR}", size=10)
    set_paragraph_text(paragraphs[7], "Intended for publication in the journal: Plant Protection Science", size=10)
    set_paragraph_text(
        paragraphs[16],
        "any use of artificial intelligence (AI) tools in manuscript preparation is clearly disclosed; AI was used only for language editing, improving readability and structuring text, and not to generate or manipulate research data, results, interpretations or images;",
        size=9,
    )
    set_paragraph_text(paragraphs[30], f"Name of the corresponding author: {AUTHOR}\nEmail (of the corresponding author): {EMAIL}", size=10)
    set_paragraph_text(
        paragraphs[31],
        "\nDate: ____________________      Signature of the corresponding author: ______________________________\n"
        "                                                                          (hand-written)\n"
        "Signed and scanned, please submit with the manuscript.",
        size=10,
    )
    doc.save(DECLARATION_OUT)


def verify_inputs() -> None:
    expected = {
        MANUSCRIPT_TEMPLATE: "ef8487774cbc138ef4199ffd9ce85ae01188f4c549fa4903147e794d5f490768",
        TITLE_TEMPLATE: "c1b99e632553710254a0df1df2ad312ce7be02f9b04ef1c734d71a112f1d688c",
        LETTER_TEMPLATE: "ce87f6ec3f9e11b10f7a59c0a38600f5b187e910ed4a8dc936aad5cabb78f644",
        DECLARATION_TEMPLATE: "9a99391298b3431351845750fb3191f7a7898213c1d55b610ace6bb2562b3200",
    }
    for path, digest in expected.items():
        if not path.exists():
            raise FileNotFoundError(path)
        if sha256(path) != digest:
            raise RuntimeError(f"Template checksum mismatch: {path}")
    for name in FIGURE_FILES:
        if not (FIGURES / name).exists():
            raise FileNotFoundError(FIGURES / name)


def validate_docx(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"Corrupt part {bad} in {path}")


def main() -> None:
    verify_inputs()
    char_count = build_manuscript()
    build_title_page(char_count)
    build_cover_letter()
    build_declaration()
    for path in (MANUSCRIPT_OUT, TITLE_OUT, LETTER_OUT, DECLARATION_OUT):
        validate_docx(path)
        print(f"{path.name}\t{path.stat().st_size}\t{sha256(path)}")
    print(f"manuscript_characters_including_spaces\t{char_count}")
    print(f"abstract_words\t{len(ABSTRACT.split())}")
    print(f"title_characters\t{len(TITLE)}")
    print(f"references\t{len(REFERENCES)}")


if __name__ == "__main__":
    main()
