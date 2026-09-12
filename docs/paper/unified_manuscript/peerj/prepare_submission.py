#!/usr/bin/env python3
"""Build the PeerJ submission package from the unified manuscript sources."""

from __future__ import annotations

import hashlib
import importlib.util
import re
import shutil
import subprocess
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


OUTPUT_DIR = Path(__file__).resolve().parent
MANUSCRIPT_DIR = OUTPUT_DIR.parent
SOURCE_MD = MANUSCRIPT_DIR / "manuscript.md"
SOURCE_SUPPLEMENT = MANUSCRIPT_DIR / "lychee_unified_manuscript_supplement.zip"
RSOS_SCRIPT = MANUSCRIPT_DIR / "royal_society_open_science" / "prepare_submission.py"
PEERJ_TEMPLATE = OUTPUT_DIR / "PeerJ-research-manuscript-template.docx"

TITLE = (
    "Cultivar-dependent transcriptional responses of lychee to "
    "Peronophythora litchii: a registered genome-wide analysis"
)
AUTHOR = "Eric Zhuang"
AFFILIATION = "NYU Langone Health, New York, NY, USA"
# Replace this with the author's full street address and ZIP/postal code before
# submission if PeerJ requires every item shown in the template prompt.
CORRESPONDENCE_ADDRESS = AFFILIATION
EMAIL = "eric.zhuang@nyulangone.org"
ORCID = "0009-0001-9050-0214"
ZENODO_DOI = "https://zenodo.org/records/22436625"

KEYWORDS = (
    "Litchi chinensis; Peronophythora litchii; plant-pathogen interaction; "
    "RNA sequencing; cultivar-dependent response; preregistered analysis"
)

ABSTRACT_SECTIONS = [
    (
        "Background",
        "Litchi downy blight, caused by the oomycete *Peronophythora litchii*, "
        "is among the most damaging diseases of lychee (*Litchi chinensis* Sonn.). "
        "Public RNA-sequencing cohorts spanning cultivars, tissues, and infection "
        "time points offer a resource for identifying cultivar-dependent infection "
        "responses, but they are easily over-interpreted when the same cohort is used "
        "both to select candidates and to appear to confirm them.",
    ),
    (
        "Methods",
        "We unified an exploratory reanalysis with a prospectively registered "
        "confirmatory stage. The confirmatory protocol fixed dataset roles, statistical "
        "models, thresholds, and evidence vocabulary before external outcomes were "
        "examined. It combined a genome-wide negative-binomial cultivar-by-infection "
        "analysis with mapping and gene-model quality control, independent statistical "
        "and quantification checks, prespecified evaluation in independent public "
        "cohorts, and deterministic evidence integration.",
    ),
    (
        "Results",
        "Among 19,445 expressed genes, 262 met the genome-wide interaction threshold; "
        "206 passed mapping and gene-model quality control, 19 remained significant "
        "under an independent edgeR quasi-likelihood analysis, and 16 passed the complete "
        "registered robustness procedure. In a prespecified cross-context evaluation "
        "against an independent pericarp time course, two genes were supported and five "
        "were directionally contradictory. The internally robust and externally supported "
        "sets did not overlap, an outcome compatible with chance given the set sizes. None "
        "of 18 exploratory candidates met the subsequent genome-wide interaction criterion, "
        "and two failed uniform-mappability control. Evidence integration assigned no gene "
        "to the highest confidence tier, placed 12 internally robust genes in a middle tier, "
        "and retired 65 entities.",
    ),
    (
        "Conclusions",
        "Candidates selected for effect-size prominence in one cohort largely failed a "
        "formally specified genome-wide interaction test. The analysis defines the claims "
        "supported by the available public data and releases a reproducible candidate "
        "resource with explicit evidence boundaries. Experimental perturbation remains "
        "necessary for causal conclusions.",
    ),
]

FIGURES = [
    (1, "figure1_study_design_qc.png", "Figure_1.png"),
    (2, "figure2_discovery_legacy.png", "Figure_2.png"),
    (3, "figure3_robustness_external.png", "Figure_3.png"),
    (4, "figure4_pathways_signatures.png", "Figure_4.png"),
    (5, "figure5_transcript_usage.png", "Figure_5.png"),
    (6, "figure6_orthogonal_tiers.png", "Figure_6.png"),
]

HEADING_RE = re.compile(r"^(#{1,3})\s+(.+)$")
IMAGE_RE = re.compile(r"^!\[(?P<caption>.+)]\((?P<path>[^)]+)\)$")
REFERENCE_RE = re.compile(r"^\d+\.\s+")
TABLE_CAPTION_RE = re.compile(r"^\*\*Table\s+(\d+)\.\s+(.+)")


def load_shared_module():
    spec = importlib.util.spec_from_file_location("rsos_submission_helpers", RSOS_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load shared helpers from {RSOS_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


shared = load_shared_module()


def plain_markdown(value: str) -> str:
    return shared.plain_markdown(value)


def manuscript_text(value: str) -> str:
    """Apply PeerJ citation labels without altering scientific content."""
    value = re.sub(r"\bSupplementary S(\d+[a-z]?(?:[–-][a-z])?)", r"Table S\1", value)
    value = re.sub(r"(?<=\()Figure (?=\d)", "Fig. ", value)
    value = re.sub(r"(?<=; )Figure (?=\d)", "Fig. ", value)
    return value


def clear_template_placeholders(document: Document) -> None:
    """Remove the sample body while preserving template styles and section settings."""
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def template_paragraph(document: Document):
    paragraph = document.add_paragraph(style="normal")
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    return paragraph


def format_runs(paragraph, *, size: float = 12, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.font.name = "Times"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times")
        run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Times New Roman")
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold


def add_template_heading(document: Document, value: str, *, level: int = 1):
    paragraph = template_paragraph(document)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(0)
    shared.set_keep_with_next(paragraph)
    run = paragraph.add_run(value)
    run.bold = True
    format_runs(paragraph, size=14 if level == 1 else 12, bold=True)
    return paragraph


def add_author_cover_page(document: Document) -> None:
    title = template_paragraph(document)
    title.paragraph_format.space_after = Pt(0)
    shared.add_inline(
        title,
        "Cultivar-dependent transcriptional responses of lychee to "
        "*Peronophythora litchii*: a registered genome-wide analysis",
    )
    format_runs(title, size=18, bold=True)

    template_paragraph(document)
    author = template_paragraph(document)
    author.add_run(AUTHOR)
    affiliation_number = author.add_run("1")
    affiliation_number.font.superscript = True
    format_runs(author)

    template_paragraph(document)
    affiliation = template_paragraph(document)
    affiliation.add_run("1").font.superscript = True
    affiliation.add_run(f" {AFFILIATION}")
    format_runs(affiliation)

    template_paragraph(document)
    corresponding = template_paragraph(document)
    corresponding.add_run("Corresponding Author:").bold = True
    format_runs(corresponding)

    correspondent = template_paragraph(document)
    correspondent.add_run(AUTHOR)
    correspondent.add_run("1").font.superscript = True
    format_runs(correspondent)

    address = template_paragraph(document)
    address.add_run(CORRESPONDENCE_ADDRESS)
    format_runs(address)

    email = template_paragraph(document)
    email.add_run(f"Email address: {EMAIL}")
    format_runs(email)

    template_paragraph(document)

def add_structured_abstract(document: Document) -> None:
    add_template_heading(document, "Abstract")
    rendered = []
    for label, body in ABSTRACT_SECTIONS:
        paragraph = template_paragraph(document)
        paragraph.add_run(f"{label}. ").bold = True
        shared.add_inline(paragraph, body)
        format_runs(paragraph)
        rendered.append(f"{label}. {plain_markdown(body)}")

    abstract_text = "\n".join(rendered)
    word_count = len(re.findall(r"\b[\w'-]+\b", abstract_text))
    if word_count > 500 or len(abstract_text) > 3000:
        raise ValueError(
            f"Structured abstract exceeds PeerJ limit: {word_count} words, "
            f"{len(abstract_text)} characters"
        )

def strip_heading_number(value: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", value)


def section_bounds(lines: list[str], start: str, end: str) -> list[str]:
    start_index = lines.index(start) + 1
    end_index = lines.index(end)
    return lines[start_index:end_index]


def table_number_from_caption(value: str) -> int | None:
    match = re.match(r"^\*\*Table\s+(\d+)\.", value)
    return int(match.group(1)) if match else None


def add_body_lines(document: Document, lines: list[str]) -> None:
    index = 0
    pending_table: int | None = None
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            add_template_heading(
                document,
                plain_markdown(strip_heading_number(heading.group(2))),
                level=2,
            )
            index += 1
            continue

        image = IMAGE_RE.match(stripped)
        if image:
            index += 1
            continue

        number = table_number_from_caption(stripped)
        if number is not None:
            pending_table = number
            index += 1
            continue

        if stripped.startswith("|"):
            _, index = shared.parse_table(lines, index)
            if pending_table is None:
                raise ValueError("Found a Markdown table without a numbered caption")
            pending_table = None
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                HEADING_RE.match(candidate)
                or IMAGE_RE.match(candidate)
                or candidate.startswith("|")
                or table_number_from_caption(candidate) is not None
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = template_paragraph(document)
        shared.add_inline(paragraph, manuscript_text(" ".join(paragraph_lines)))
        format_runs(paragraph)


def add_section(document: Document, title: str, lines: list[str]) -> None:
    add_template_heading(document, title)
    add_body_lines(document, lines)


def extract_references(lines: list[str]) -> list[str]:
    start = lines.index("## References") + 1
    references: list[str] = []
    current: list[str] = []
    for line in lines[start:]:
        stripped = line.strip()
        if not stripped:
            continue
        if REFERENCE_RE.match(stripped):
            if current:
                references.append(" ".join(current))
            current = [stripped]
        elif current:
            current.append(stripped)
    if current:
        references.append(" ".join(current))
    return references


def add_acknowledgements(document: Document) -> None:
    add_template_heading(document, "Acknowledgements")
    acknowledgment = template_paragraph(document)
    shared.add_inline(
        acknowledgment,
        "The author has no personal acknowledgments to declare. OpenAI Codex "
        "(GPT-5; accessed September 6, 2026) was used to assist with language editing "
        "and journal-specific document formatting. It was not used to generate data, "
        "perform analyses, interpret results, or draw scientific conclusions. The "
        "author confirms that the originality and accuracy of the content were checked, "
        "that the applicable OpenAI terms of use were reviewed and judged suitable for "
        "publication, and that full responsibility is accepted for the integrity of the "
        "manuscript, including the accuracy of its references. The pre-edit and edited "
        "versions have been retained and can be supplied to the Editor on request.",
    )
    format_runs(acknowledgment)


def add_references(document: Document, references: list[str]) -> None:
    add_template_heading(document, "References")
    for reference in references:
        paragraph = template_paragraph(document)
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        shared.add_inline(paragraph, reference)
        format_runs(paragraph)


def build_manuscript() -> Path:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    if not PEERJ_TEMPLATE.is_file():
        raise FileNotFoundError(PEERJ_TEMPLATE)
    document = Document(PEERJ_TEMPLATE)
    clear_template_placeholders(document)

    properties = document.core_properties
    properties.title = TITLE
    properties.author = AUTHOR
    properties.subject = "Research Article submitted to PeerJ"
    properties.keywords = KEYWORDS

    add_author_cover_page(document)
    add_structured_abstract(document)
    add_section(
        document,
        "Introduction",
        section_bounds(lines, "## 1. Introduction", "## 2. Results"),
    )
    add_section(
        document,
        "Materials & Methods",
        section_bounds(lines, "## 5. Materials and Methods", "## Declarations"),
    )
    add_section(
        document,
        "Results",
        section_bounds(lines, "## 2. Results", "## 3. Discussion"),
    )
    add_section(
        document,
        "Discussion",
        section_bounds(lines, "## 3. Discussion", "## 4. Conclusions"),
    )
    add_section(
        document,
        "Conclusions",
        section_bounds(lines, "## 4. Conclusions", "## 5. Materials and Methods"),
    )
    add_acknowledgements(document)

    references = extract_references(lines)
    add_references(document, references)

    output = OUTPUT_DIR / "PeerJ_manuscript.docx"
    document.save(output)
    return output


def extract_tables(lines: list[str]) -> dict[int, tuple[str, list[list[str]]]]:
    tables: dict[int, tuple[str, list[list[str]]]] = {}
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        number = table_number_from_caption(stripped)
        if number is None:
            index += 1
            continue
        caption = plain_markdown(stripped)
        cursor = index + 1
        while cursor < len(lines) and not lines[cursor].strip():
            cursor += 1
        if cursor >= len(lines) or not lines[cursor].strip().startswith("|"):
            raise ValueError(f"Table {number} caption is not followed by a Markdown table")
        rows, cursor = shared.parse_table(lines, cursor)
        tables[number] = (caption, rows)
        index = cursor
    if sorted(tables) != [1, 2, 3, 4]:
        raise ValueError(f"Expected Tables 1–4, found {sorted(tables)}")
    return tables


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def build_table_file(number: int, caption: str, rows: list[list[str]]) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(8)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(0)

    title = document.add_paragraph()
    title.paragraph_format.line_spacing = 1
    title.paragraph_format.space_after = Pt(8)
    prefix = f"Table {number}."
    title.add_run(prefix).bold = True
    title.add_run(caption[len(prefix) :])

    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, source_row in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            shared.set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shared.set_cell_margins(cell, top=35, start=35, bottom=35, end=35)
            if row_index == 0:
                shared.shade_cell(cell, "D9EAF7")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(0)
            if column_index < len(source_row):
                shared.add_inline(paragraph, source_row[column_index])
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(7.5)
                if row_index == 0:
                    run.bold = True

    properties = document.core_properties
    properties.title = f"Table {number}: {TITLE}"
    properties.author = AUTHOR
    output = OUTPUT_DIR / "tables" / f"Table_{number}.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def build_tables() -> list[Path]:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    tables = extract_tables(lines)
    return [build_table_file(number, *tables[number]) for number in sorted(tables)]


def normalize_odt_page_layout(path: Path) -> None:
    """Retain the DOCX table page geometry after textutil conversion."""
    with ZipFile(path) as archive:
        members = [
            (member, archive.read(member.filename)) for member in archive.infolist()
        ]

    updated_members: list[tuple[object, bytes]] = []
    updated_layout = False
    for member, contents in members:
        if member.filename == "styles.xml":
            xml = contents.decode("utf-8")
            layout_match = re.search(
                r"<style:page-layout-properties\b[^>]*/>",
                xml,
            )
            if layout_match is None:
                raise RuntimeError(f"Could not find ODT page layout in {path}")
            layout = layout_match.group(0)
            replacements = {
                "fo:page-width": "11in",
                "fo:page-height": "8.5in",
                "fo:margin-top": "1in",
                "fo:margin-bottom": "1in",
                "fo:margin-left": "1in",
                "fo:margin-right": "1in",
            }
            for attribute, value in replacements.items():
                layout, count = re.subn(
                    rf'{re.escape(attribute)}="[^"]+"',
                    f'{attribute}="{value}"',
                    layout,
                )
                if count != 1:
                    raise RuntimeError(
                        f"Expected one {attribute} attribute in {path}, found {count}"
                    )
            xml = (
                xml[: layout_match.start()]
                + layout
                + xml[layout_match.end() :]
            )
            contents = xml.encode("utf-8")
            updated_layout = True
        updated_members.append((member, contents))

    if not updated_layout:
        raise RuntimeError(f"Did not update ODT page layout in {path}")

    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for member, contents in updated_members:
            archive.writestr(member, contents)
    path.write_bytes(buffer.getvalue())


def build_odt_tables(docx_tables: list[Path]) -> list[Path]:
    """Create PeerJ-compatible editable ODT alternatives with macOS textutil."""
    converter = shutil.which("textutil")
    if converter is None:
        raise RuntimeError(
            "ODT table generation requires macOS textutil; retain the DOCX tables "
            "or convert them with LibreOffice on another platform."
        )

    output_dir = OUTPUT_DIR / "tables_odt"
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for number, source in enumerate(docx_tables, start=1):
        target = output_dir / f"Table{number}.odt"
        subprocess.run(
            [
                converter,
                "-convert",
                "odt",
                "-output",
                str(target),
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        if not target.is_file():
            raise RuntimeError(f"ODT conversion did not create {target}")
        normalize_odt_page_layout(target)
        outputs.append(target)
    return outputs


def copy_figures() -> list[Path]:
    output_dir = OUTPUT_DIR / "figures"
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for _, source_name, target_name in FIGURES:
        source = MANUSCRIPT_DIR / "figures" / source_name
        target = output_dir / target_name
        shutil.copy2(source, target)
        with Image.open(target) as image:
            dpi = image.info.get("dpi", (0, 0))
            if min(image.size) < 900 or min(dpi) < 295:
                raise ValueError(
                    f"{target.name} failed PeerJ figure checks: "
                    f"pixels={image.size}, dpi={dpi}"
                )
        outputs.append(target)
    return outputs


def build_supplement() -> Path:
    payload: dict[str, bytes] = {}
    with ZipFile(SOURCE_SUPPLEMENT) as archive:
        source_root = archive.namelist()[0].split("/", 1)[0]
        for member in archive.infolist():
            if member.is_dir():
                continue
            relative = member.filename.removeprefix(source_root + "/")
            if relative in {"README.md", "MANIFEST.tsv"}:
                continue
            payload[relative] = archive.read(member)

    readme = f"""# PeerJ Supplemental Data S1

Associated article: {TITLE}
Author: {AUTHOR}
Package date: 2026-09-06

This archive contains the complete supplemental and reproducibility material:

- data/supplementary_tables: Supplementary Tables S1-S18 (S16 has parts a and b);
- data/supplementary_figures: Figures S1-S3 in PDF and 300-DPI PNG formats;
- data/figure_source_data: tab-separated source data for all main and
  supplementary analytical figures; and
- code/analysis: Snakemake workflows, Python and R scripts, configurations,
  metadata, synthetic fixtures, regression tests, and environment specifications.

Raw sequencing reads and large reference resources are not redistributed. Their
public accessions and identifiers are given in the manuscript. The supplementary
tables, figures, and figure source data are also archived at {ZENODO_DOI}; this
submission archive additionally contains the analysis code and workflows.

MANIFEST.tsv records the byte size and SHA-256 digest of every other file in this
archive. The materials are supplied under CC BY 4.0; third-party software retains
its own license.
""".encode()
    payload["README.md"] = readme

    manifest_lines = ["path\tbytes\tsha256"]
    for relative, contents in sorted(payload.items()):
        manifest_lines.append(
            f"{relative}\t{len(contents)}\t{hashlib.sha256(contents).hexdigest()}"
        )
    payload["MANIFEST.tsv"] = ("\n".join(manifest_lines) + "\n").encode()

    output = OUTPUT_DIR / "supplemental_information" / "PeerJ_supplemental_data_S1.zip"
    output.parent.mkdir(parents=True, exist_ok=True)
    root = "PeerJ_supplemental_data_S1"
    with ZipFile(output, "w", ZIP_DEFLATED, compresslevel=9) as archive:
        for relative, contents in sorted(payload.items()):
            archive.writestr(f"{root}/{relative}", contents)
    if output.stat().st_size > 30 * 1024 * 1024:
        raise ValueError(f"Supplement exceeds PeerJ's 30 MB individual-file limit: {output}")
    return output


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_manifest(paths: list[Path]) -> Path:
    output = OUTPUT_DIR / "UPLOAD_FILE_MANIFEST_SHA256.tsv"
    lines = ["sha256\tbytes\tfile"]
    for path in sorted(paths, key=lambda item: item.relative_to(OUTPUT_DIR).as_posix()):
        lines.append(
            f"{sha256(path)}\t{path.stat().st_size}\t"
            f"{path.relative_to(OUTPUT_DIR).as_posix()}"
        )
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return output


def main() -> None:
    manuscript = build_manuscript()
    docx_tables = build_tables()
    tables = build_odt_tables(docx_tables)
    figures = copy_figures()
    supplement = build_supplement()
    manifest = write_manifest([manuscript, *tables, *figures, supplement])
    print(f"Built PeerJ package in {OUTPUT_DIR}")
    print(f"Upload manifest: {manifest}")


if __name__ == "__main__":
    main()
