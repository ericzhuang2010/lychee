#!/usr/bin/env python3
"""Build and validate the Plant-Environment Interactions submission package."""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


OUTPUT_DIR = Path(__file__).resolve().parent
MANUSCRIPT_DIR = OUTPUT_DIR.parent
PROJECT_ROOT = OUTPUT_DIR.parents[3]
SOURCE_DIR = OUTPUT_DIR / "source"
FIGURE_DIR = OUTPUT_DIR / "figures"
SUPPORT_DIR = OUTPUT_DIR / "supporting_information"

MANUSCRIPT_SOURCE = SOURCE_DIR / "PEI_manuscript.md"
SUPPORT_SOURCE = SOURCE_DIR / "PEI_supporting_information.md"
COVER_SOURCE = SOURCE_DIR / "PEI_cover_letter.md"
SOURCE_SUPPLEMENT = MANUSCRIPT_DIR / "lychee_unified_manuscript_supplement.zip"
EXPECTED_SUPPLEMENT_MD5 = "a7dce65e381f77ca365820311704b736"

TITLE_MARKDOWN = (
    "A registered genome-wide analysis identifies robust, context-dependent "
    "transcriptional responses of lychee to *Peronophythora litchii*"
)
TITLE_PLAIN = (
    "A registered genome-wide analysis identifies robust, context-dependent "
    "transcriptional responses of lychee to Peronophythora litchii"
)
RUNNING_TITLE = "Registered analysis of lychee responses"
AUTHOR = "Eric Zhuang"
AFFILIATION = "NYU Langone Health, New York, New York, USA"
EMAIL = "eric.zhuang@nyulangone.org"
ORCID = "0009-0001-9050-0214"
ARTICLE_TYPE = "Research Article"
ZENODO_RECORD_URL = "https://zenodo.org/records/22436625"

MAIN_FIGURES = [
    (MANUSCRIPT_DIR / "figures/figure1_study_design_qc.pdf", "PEI_Figure_1.pdf"),
    (MANUSCRIPT_DIR / "figures/figure2_discovery_legacy.pdf", "PEI_Figure_2.pdf"),
    (MANUSCRIPT_DIR / "figures/figure3_robustness_external.pdf", "PEI_Figure_3.pdf"),
    (MANUSCRIPT_DIR / "figures/figure4_pathways_signatures.pdf", "PEI_Figure_4.pdf"),
    (MANUSCRIPT_DIR / "figures/figure6_orthogonal_tiers.pdf", "PEI_Figure_5.pdf"),
]
MAIN_FIGURE_SOURCES = {
    f"Figure {number}": source
    for number, (source, _) in enumerate(MAIN_FIGURES, start=1)
}

SUPPORTING_FIGURES = {
    "Figure S1": PROJECT_ROOT / "results/figures/FigureS1_replicate_level_counts.pdf",
    "Figure S2": PROJECT_ROOT / "results/figures/FigureS2_power_analysis.pdf",
    "Figure S3": MANUSCRIPT_DIR / "figures/figureS3_exploratory_signature.pdf",
    "Figure S4": MANUSCRIPT_DIR / "figures/figure5_transcript_usage.pdf",
}


@dataclass(frozen=True)
class Segment:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    url: str | None = None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def md5(path: Path) -> str:
    digest = hashlib.md5()  # noqa: S324 - record comparison, not security
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def word_count(text: str) -> int:
    text = re.sub(r"https?://\S+", " URL ", text)
    text = re.sub(r"[`*#|]", " ", text)
    return len(re.findall(r"\b[\w’'-]+\b", text))


def extract_markdown_section(text: str, heading: str, next_heading: str | None) -> str:
    start_marker = f"# {heading}\n"
    start = text.index(start_marker) + len(start_marker)
    if next_heading is None:
        return text[start:].strip()
    end = text.index(f"\n# {next_heading}", start)
    return text[start:end].strip()


def count_references(text: str) -> int:
    references = extract_markdown_section(text, "References", "Tables")
    return len(re.findall(r"^- ", references, flags=re.MULTILINE))


def inline_segments(text: str, *, bold: bool = False, italic: bool = False) -> list[Segment]:
    segments: list[Segment] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            segments.append(Segment("".join(buffer), bold=bold, italic=italic))
            buffer.clear()

    index = 0
    while index < len(text):
        if text.startswith("**", index):
            end = text.find("**", index + 2)
            if end >= 0:
                flush()
                segments.extend(inline_segments(text[index + 2 : end], bold=True, italic=italic))
                index = end + 2
                continue
        if text[index] == "*":
            end = text.find("*", index + 1)
            if end >= 0:
                flush()
                segments.extend(inline_segments(text[index + 1 : end], bold=bold, italic=True))
                index = end + 1
                continue
        if text[index] == "`":
            end = text.find("`", index + 1)
            if end >= 0:
                flush()
                segments.append(Segment(text[index + 1 : end], bold=bold, italic=italic, code=True))
                index = end + 1
                continue
        if text.startswith(("https://", "http://"), index):
            end = index
            while end < len(text) and not text[end].isspace():
                end += 1
            candidate = text[index:end]
            trailing = ""
            while candidate and candidate[-1] in ".,;":
                trailing = candidate[-1] + trailing
                candidate = candidate[:-1]
            flush()
            segments.append(Segment(candidate, bold=bold, italic=italic, url=candidate))
            if trailing:
                buffer.append(trailing)
            index = end
            continue
        buffer.append(text[index])
        index += 1
    flush()
    return segments


def add_hyperlink(paragraph, segment: Segment) -> None:
    relationship_id = paragraph.part.relate_to(
        segment.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    value = OxmlElement("w:t")
    value.text = segment.text
    run.extend([properties, value])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_segments(paragraph, segments: list[Segment], *, size: float | None = None) -> None:
    for segment in segments:
        if segment.url:
            add_hyperlink(paragraph, segment)
            continue
        run = paragraph.add_run(segment.text)
        run.bold = segment.bold
        run.italic = segment.italic
        if segment.code:
            run.font.name = "Courier New"
            run.font.size = Pt(10 if size is None else size)
        elif size is not None:
            run.font.size = Pt(size)


def add_markdown_paragraph(document: Document, text: str, *, style=None):
    paragraph = document.add_paragraph(style=style)
    add_segments(paragraph, inline_segments(text.strip()))
    return paragraph


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, display, end])


def configure_document(
    document: Document,
    *,
    header_text: str,
    line_numbers: bool,
    font_name: str = "Times New Roman",
    font_size: float = 12,
    line_spacing: float = 1.5,
) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(font_size)
    normal.paragraph_format.line_spacing = line_spacing
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        properties = style.element.get_or_add_pPr()
        borders = properties.find(qn("w:pBdr"))
        if borders is not None:
            properties.remove(borders)

    caption = document.styles["Caption"]
    caption.font.name = font_name
    caption.font.size = Pt(10)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.text = header_text
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = font_name
    header.runs[0].font.size = Pt(9)
    add_page_field(section.footer.paragraphs[0])

    if line_numbers:
        line_numbering = OxmlElement("w:lnNumType")
        line_numbering.set(qn("w:countBy"), "1")
        line_numbering.set(qn("w:distance"), "360")
        line_numbering.set(qn("w:restart"), "continuous")
        section._sectPr.append(line_numbering)


def set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    properties.append(no_split)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def parse_markdown_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line_number, line in enumerate(lines):
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if line_number == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("Malformed Markdown table")
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_header(table.rows[0])
    for row_number, values in enumerate(rows):
        prevent_row_split(table.rows[row_number])
        for column_number, value in enumerate(values):
            cell = table.cell(row_number, column_number)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            add_segments(
                paragraph,
                inline_segments(value),
                size=7.5 if len(rows[0]) >= 7 else 8,
            )
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                if row_number == 0:
                    run.bold = True
            if row_number == 0:
                shade_cell(cell, "D9E2F3")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def convert_pdf_to_png(source: Path, destination: Path) -> None:
    result = subprocess.run(
        ["sips", "-s", "format", "png", "-Z", "2600", str(source), "--out", str(destination)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(f"Could not rasterize {source}:\n{result.stdout}")


def add_embedded_figure(
    document: Document,
    source: Path,
    caption_title: str,
    caption_text: str,
    image_directory: Path,
) -> None:
    figure_key = re.match(r"Figure\s+(?:S)?\d+", caption_title).group(0)
    png = image_directory / (figure_key.replace(" ", "_") + ".png")
    convert_pdf_to_png(source, png)
    with Image.open(png) as image:
        width_px, height_px = image.size
    max_width = 6.6
    max_height = 7.1
    width = min(max_width, max_height * width_px / height_px)
    height = width * height_px / width_px
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(
        str(png), width=Inches(width), height=Inches(height)
    )
    shape._inline.docPr.set("descr", re.sub(r"[*`]", "", caption_text)[:1000])
    caption = document.add_paragraph(style="Caption")
    title_run = caption.add_run(caption_title + ". ")
    title_run.bold = True
    add_segments(caption, inline_segments(caption_text))


def render_markdown(
    document: Document,
    text: str,
    *,
    main_figures: dict[str, Path] | None = None,
    supporting_figures: dict[str, Path] | None = None,
    image_directory: Path | None = None,
) -> None:
    lines = text.splitlines()
    index = 0
    section = ""
    pending_caption: str | None = None

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue

        heading = re.match(r"^(#{1,2})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if re.match(r"^(?:Table|Figure)\s+(?:S)?\d+\.", title):
                pending_caption = title
            else:
                section = title
                document.add_heading(title, level=level)
            index += 1
            continue

        if line.startswith("|"):
            if pending_caption:
                caption = document.add_paragraph(style="Caption")
                caption.paragraph_format.page_break_before = pending_caption.startswith("Table S")
                title_run = caption.add_run(pending_caption + ".")
                title_run.bold = True
                pending_caption = None
            table_lines: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, parse_markdown_table(table_lines))
            continue

        if line.startswith("- "):
            value = line[2:].strip()
            paragraph = document.add_paragraph()
            if section == "References" or section == "References cited only in Supporting Information":
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(4)
                add_segments(paragraph, inline_segments(value), size=10)
            else:
                paragraph.style = "List Bullet"
                paragraph.paragraph_format.line_spacing = 1.15
                add_segments(paragraph, inline_segments(value))
            index += 1
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index].rstrip()
            if not candidate or candidate.startswith(("#", "|", "- ")):
                break
            paragraph_lines.append(candidate.strip())
            index += 1
        value = " ".join(paragraph_lines)

        if pending_caption:
            if (
                re.match(r"Figure \d+", pending_caption)
                and main_figures is not None
                and image_directory is not None
            ):
                key = re.match(r"Figure \d+", pending_caption).group(0)
                add_embedded_figure(
                    document,
                    main_figures[key],
                    pending_caption,
                    value,
                    image_directory,
                )
            elif (
                pending_caption.startswith("Figure S")
                and supporting_figures is not None
                and image_directory is not None
            ):
                key_match = re.match(r"Figure S\d+", pending_caption)
                key = key_match.group(0)
                add_embedded_figure(
                    document,
                    supporting_figures[key],
                    pending_caption,
                    value,
                    image_directory,
                )
            else:
                caption = document.add_paragraph(style="Caption")
                title_run = caption.add_run(pending_caption + ". ")
                title_run.bold = True
                add_segments(caption, inline_segments(value))
                if pending_caption.startswith("Table S"):
                    caption.paragraph_format.page_break_before = True
            pending_caption = None
        else:
            add_markdown_paragraph(document, value)


def build_manuscript_docx(
    source_text: str,
    output: Path,
    abstract_words: int,
    main_words: int,
    reference_count: int,
) -> None:
    document = Document()
    configure_document(
        document,
        header_text=RUNNING_TITLE,
        line_numbers=True,
        line_spacing=1.5,
    )
    document.core_properties.title = TITLE_PLAIN
    document.core_properties.subject = "Research Article submitted to Plant-Environment Interactions"
    document.core_properties.author = AUTHOR
    document.core_properties.keywords = (
        "downy blight; disease resistance; RNA sequencing; cultivar-by-infection "
        "interaction; transcript usage; reproducible bioinformatics"
    )

    article_type = document.add_paragraph()
    article_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    article_type.add_run(ARTICLE_TYPE).bold = True
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_segments(title, inline_segments(TITLE_MARKDOWN))
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run(AUTHOR).bold = True
    details = document.add_paragraph(
        f"{AFFILIATION}\nCorresponding author: {EMAIL}\nORCID: {ORCID}"
    )
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.paragraph_format.line_spacing = 1.0
    running = document.add_paragraph()
    running.alignment = WD_ALIGN_PARAGRAPH.CENTER
    running.add_run("Running title: ").bold = True
    running.add_run(RUNNING_TITLE)
    counts = document.add_paragraph(
        f"Abstract: {abstract_words} words | Main body (I/R/D): {main_words} words | "
        f"References: {reference_count} | Figures: 5 | Tables: 2"
    )
    counts.alignment = WD_ALIGN_PARAGRAPH.CENTER
    counts.paragraph_format.line_spacing = 1.0
    with TemporaryDirectory(prefix="pei_main_images_") as temp_name:
        render_markdown(
            document,
            source_text,
            main_figures=MAIN_FIGURE_SOURCES,
            image_directory=Path(temp_name),
        )
        document.save(output)


def build_support_docx(source_text: str, output: Path) -> None:
    document = Document()
    configure_document(
        document,
        header_text="Supporting Information",
        line_numbers=False,
        line_spacing=1.15,
    )
    document.core_properties.title = "Supporting Information: " + TITLE_PLAIN
    document.core_properties.subject = "Supporting Information for Plant-Environment Interactions"
    document.core_properties.author = AUTHOR

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Supporting Information")
    manuscript_title = document.add_paragraph()
    manuscript_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_segments(manuscript_title, inline_segments(TITLE_MARKDOWN))
    for run in manuscript_title.runs:
        run.bold = True
    author = document.add_paragraph(f"{AUTHOR}\n{AFFILIATION}")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.line_spacing = 1.0

    with TemporaryDirectory(prefix="pei_support_images_") as temp_name:
        render_markdown(
            document,
            source_text,
            supporting_figures=SUPPORTING_FIGURES,
            image_directory=Path(temp_name),
        )
        document.save(output)


def build_cover_docx(source_text: str, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    document.core_properties.title = "Cover letter for Plant-Environment Interactions"
    document.core_properties.author = AUTHOR

    paragraphs = re.split(r"\n\s*\n", source_text.strip())
    for number, block in enumerate(paragraphs):
        paragraph = document.add_paragraph()
        if number == len(paragraphs) - 1:
            paragraph.paragraph_format.space_before = Pt(2)
        lines = [line.rstrip().removesuffix("  ") for line in block.splitlines()]
        for line_number, line in enumerate(lines):
            if line_number:
                paragraph.add_run().add_break()
            add_segments(paragraph, inline_segments(line))
    document.save(output)


def build_metadata(
    abstract: str,
    why_statement: str,
    abstract_words: int,
    main_words: int,
    methods_words: int,
    conclusions_words: int,
    narrative_words: int,
    reference_count: int,
) -> str:
    return f"""# Plant-Environment Interactions submission metadata

## Article type

Research Article

## Title

{TITLE_MARKDOWN}

## Running title ({len(RUNNING_TITLE)} characters; limit <40)

{RUNNING_TITLE}

## Abstract ({abstract_words} words; limit 250)

{abstract}

## Keywords (six; requested range five to eight)

1. downy blight
2. disease resistance
3. RNA sequencing
4. cultivar-by-infection interaction
5. transcript usage
6. reproducible bioinformatics

## Why this research matters ({word_count(why_statement)} words; limit 200)

{why_statement}

## Author

- Full name: {AUTHOR}
- Affiliation: {AFFILIATION}
- Email: {EMAIL}
- ORCID: {ORCID}
- Corresponding and submitting author: Yes

## Author contributions

E.Z. conceived and designed the study, curated the data, developed the methodology and software, performed the analyses and validation, prepared all figures and tables, and wrote and revised the manuscript. E.Z. approved the final manuscript and accepts responsibility for all aspects of the work.

## Data availability

All analyzed sequencing data are publicly available under PRJNA830488/GSE201243, PRJNA450886, PRJNA922966/GSE222651, PRJNA922965/GSE222650 and PRJNA1090613/GSE262200. Complete result tables, figure source data, the registered protocol and amendment log, analysis code, workflows, tests and environment specifications are archived under CC BY 4.0 at <{ZENODO_RECORD_URL}>.

## Portal declarations

- Research data used or generated: Yes
- Third-party copyrighted figures, images or supplementary material: No
- Conflict of interest: The author declares no conflict of interest.
- Funding: This work received no external funding.
- Ethics approval: Not applicable; this is a reanalysis of public plant sequencing data.
- Permission to reproduce material: Not applicable.
- Article previously published: No
- Under consideration elsewhere: No
- Preprint: No manuscript preprint has been declared. The Zenodo record contains supporting materials, not a journal article.
- AI use: OpenAI Codex assisted with language editing, manuscript condensation and journal-specific document formatting. It did not generate data or perform analyses. The author verified all scientific statements and accepts responsibility for all content.

## Counts

- Main body (Introduction, Results and Discussion): {main_words} words (maximum 3,000)
- Materials and Methods: {methods_words} words (reported separately from the main-body limit)
- Conclusions: {conclusions_words} words (reported separately from the main-body limit)
- Total narrative, Introduction through Conclusions: {narrative_words} words
- Abstract: {abstract_words} words (maximum 250)
- References: {reference_count} (journal optimum 30)
- Main illustrations: seven total (five figures and two tables; maximum 10)

## Publishing model and charge

- Fully open access; the journal states that CC BY is required.
- Current APC: USD 2,940 / GBP 2,180 / EUR 2,520 if accepted, before taxes.
- Check whether NYU Langone has a Wiley Open Access Account or other institutional coverage before submission.
"""


def build_readme(
    abstract_words: int,
    main_words: int,
    methods_words: int,
    conclusions_words: int,
    narrative_words: int,
    reference_count: int,
) -> str:
    return f"""# Plant-Environment Interactions submission package

Prepared from `../manuscript_v2.tex` and journal-specific condensed source files in `source/`. The original manuscript and other journal packages were not modified.

## Upload these files

1. `PEI_manuscript.docx` — **Main Document**. Editable text with title page, continuous line and page numbering, five embedded figures, two editable tables and complete legends.
2. `PEI_cover_letter.docx` — **Cover Letter**.
3. `figures/PEI_Figure_1.pdf` through `figures/PEI_Figure_5.pdf` — original-quality vector copies. Upload individually as **Figure 1** through **Figure 5** if the portal provides or requires separate figure-file slots; the figures also remain embedded in the Main Document for review.
4. `supporting_information/PEI_Supporting_Information.docx` — **Supporting Information**. Contains detailed methods/results, Tables S19–S20 and Figures S1–S4.
5. `supporting_information/PEI_Supporting_Data_and_Code.zip` — **Supporting Information/Data**. Contains machine-readable Tables S1–S18, all figure source data, protocol, code, workflows, tests, metadata and environments.

Do not upload the Markdown source files or validation files. Tables 1–2 are already editable in the manuscript. Supporting Figures S1–S4 are embedded in the supporting-information DOCX and do not need separate uploads unless the portal requests them.

## Portal selections

- Journal: Plant-Environment Interactions
- Article type: Research Article
- Special issue/collection: Regular submission; the relevant Plant-Biotic Interactions call closed on 30 June 2026.
- Research data used or generated: Yes
- Third-party copyrighted material: No
- Licence: CC BY (journal requirement)

Copy the abstract, keywords, plain-language significance statement, author information and declarations from `submission_metadata.md`.

## Automated guideline checks

- Abstract: {abstract_words} words (maximum 250)
- Main body (Introduction, Results and Discussion): {main_words} words (maximum 3,000)
- Materials and Methods: {methods_words} words (reported separately from the main-body limit)
- Conclusions: {conclusions_words} words (reported separately from the main-body limit)
- Total narrative, Introduction through Conclusions: {narrative_words} words
- Running title: {len(RUNNING_TITLE)} characters (must be under 40)
- References: {reference_count} (journal optimum 30)
- Main illustrations: five figures plus two tables (maximum 10)
- Spacing: 1.5 lines in the manuscript
- Line numbering: continuous
- Page numbering: included
- Figures: embedded in the Main Document, with separate original-quality vector PDFs retained
- Tables: editable Word tables

## Author checks before upload

- Confirm whether `{AFFILIATION}` needs a department/division and postal code.
- Confirm `{EMAIL}` and ORCID `{ORCID}`.
- Confirm the manuscript is not under consideration elsewhere.
- Review the AI-use acknowledgement and revise it if it does not accurately describe the author's use.
- Open the DOCX files in Microsoft Word and inspect all tables, page breaks, mathematical symbols and supporting figures.
- Inspect the portal-generated proof before approving submission.
- The Zenodo record retains an earlier manuscript title. Update its metadata if possible, or confirm that retaining that historical title is acceptable.
- The journal's current APC is USD 2,940 / GBP 2,180 / EUR 2,520 if accepted. Check Wiley institutional coverage or waiver eligibility before submission.

## Rebuild

Run `python3 prepare_submission.py` from this directory. The builder regenerates the DOCX files, copies the figure PDFs and data/code archive, and repeats all structural checks and checksums.
"""


def validate_docx(
    path: Path,
    *,
    expected_images: int,
    expected_tables: int,
    required_text: list[str],
) -> dict[str, int]:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError(f"Corrupt DOCX archive: {path.name}")
        xml = archive.read("word/document.xml")
    document = Document(path)
    combined = "\n".join(paragraph.text for paragraph in document.paragraphs)
    combined += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    if len(document.inline_shapes) != expected_images:
        raise ValueError(f"{path.name}: unexpected image count")
    if len(document.tables) != expected_tables:
        raise ValueError(f"{path.name}: unexpected table count")
    for value in required_text:
        if value not in combined:
            raise ValueError(f"{path.name}: missing required text {value!r}")
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "images": len(document.inline_shapes),
        "hyperlinks": xml.count(b"<w:hyperlink"),
        "line_numbering": xml.count(b"<w:lnNumType"),
    }


def main() -> None:
    required = [
        MANUSCRIPT_SOURCE,
        SUPPORT_SOURCE,
        COVER_SOURCE,
        SOURCE_SUPPLEMENT,
        *(source for source, _ in MAIN_FIGURES),
        *SUPPORTING_FIGURES.values(),
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required files:\n" + "\n".join(missing))
    if md5(SOURCE_SUPPLEMENT) != EXPECTED_SUPPLEMENT_MD5:
        raise ValueError("Local supporting archive does not match the verified Zenodo MD5")

    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    SUPPORT_DIR.mkdir(parents=True, exist_ok=True)
    manuscript_text = MANUSCRIPT_SOURCE.read_text(encoding="utf-8")
    support_text = SUPPORT_SOURCE.read_text(encoding="utf-8")
    cover_text = COVER_SOURCE.read_text(encoding="utf-8")

    abstract = extract_markdown_section(manuscript_text, "Abstract", "Why this research matters")
    abstract = abstract.split("**Keywords:**", maxsplit=1)[0].strip()
    why_statement = extract_markdown_section(
        manuscript_text, "Why this research matters", "1 Introduction"
    )
    introduction_text = extract_markdown_section(
        manuscript_text, "1 Introduction", "2 Materials and Methods"
    )
    methods_text = extract_markdown_section(
        manuscript_text, "2 Materials and Methods", "3 Results"
    )
    results_text = extract_markdown_section(manuscript_text, "3 Results", "4 Discussion")
    discussion_text = extract_markdown_section(
        manuscript_text, "4 Discussion", "5 Conclusions"
    )
    conclusions_text = extract_markdown_section(
        manuscript_text, "5 Conclusions", "Acknowledgements"
    )
    abstract_words = word_count(abstract)
    main_words = sum(
        word_count(section)
        for section in (introduction_text, results_text, discussion_text)
    )
    methods_words = word_count(methods_text)
    conclusions_words = word_count(conclusions_text)
    narrative_words = main_words + methods_words + conclusions_words
    reference_count = count_references(manuscript_text)

    if abstract_words > 250:
        raise ValueError(f"Abstract has {abstract_words} words (maximum 250)")
    if main_words > 3000:
        raise ValueError(
            "Main body (Introduction, Results and Discussion) has "
            f"{main_words} words (maximum 3,000)"
        )
    if len(RUNNING_TITLE) >= 40:
        raise ValueError("Running title is not under 40 characters")
    if not 5 <= manuscript_text.count(";", manuscript_text.index("**Keywords:**"), manuscript_text.index("# Why")) + 1 <= 8:
        raise ValueError("Keyword count is outside the requested range")
    if reference_count > 30:
        raise ValueError("Reference count exceeds the journal optimum of 30")

    manuscript_docx = OUTPUT_DIR / "PEI_manuscript.docx"
    cover_docx = OUTPUT_DIR / "PEI_cover_letter.docx"
    support_docx = SUPPORT_DIR / "PEI_Supporting_Information.docx"
    data_archive = SUPPORT_DIR / "PEI_Supporting_Data_and_Code.zip"

    build_manuscript_docx(
        manuscript_text,
        manuscript_docx,
        abstract_words,
        main_words,
        reference_count,
    )
    build_cover_docx(cover_text, cover_docx)
    build_support_docx(support_text, support_docx)
    shutil.copy2(SOURCE_SUPPLEMENT, data_archive)
    for source, filename in MAIN_FIGURES:
        shutil.copy2(source, FIGURE_DIR / filename)

    why_path = OUTPUT_DIR / "PEI_why_this_research_matters.txt"
    why_path.write_text(why_statement + "\n", encoding="utf-8")
    (OUTPUT_DIR / "submission_metadata.md").write_text(
        build_metadata(
            abstract,
            why_statement,
            abstract_words,
            main_words,
            methods_words,
            conclusions_words,
            narrative_words,
            reference_count,
        ),
        encoding="utf-8",
    )
    (OUTPUT_DIR / "README_submission.md").write_text(
        build_readme(
            abstract_words,
            main_words,
            methods_words,
            conclusions_words,
            narrative_words,
            reference_count,
        ),
        encoding="utf-8",
    )

    manuscript_validation = validate_docx(
        manuscript_docx,
        expected_images=5,
        expected_tables=2,
        required_text=[
            "1 Introduction",
            "2 Materials and Methods",
            "3 Results",
            "4 Discussion",
            "5 Conclusions",
            "Conflict of Interest",
            "Author Contributions",
            "Data Availability Statement",
            "Figure Legends",
            "Figure 5.",
        ],
    )
    support_validation = validate_docx(
        support_docx,
        expected_images=4,
        expected_tables=2,
        required_text=["Methods S1", "Table S19", "Table S20", "Figure S4"],
    )
    cover_validation = validate_docx(
        cover_docx,
        expected_images=0,
        expected_tables=0,
        required_text=["Dear Dr Ingram", "Research Article", ZENODO_RECORD_URL],
    )

    upload_files = [
        manuscript_docx,
        cover_docx,
        *(FIGURE_DIR / filename for _, filename in MAIN_FIGURES),
        support_docx,
        data_archive,
    ]
    manifest_lines = ["sha256\tbytes\tfile"]
    for path in upload_files:
        manifest_lines.append(
            f"{sha256(path)}\t{path.stat().st_size}\t{path.relative_to(OUTPUT_DIR)}"
        )
    (OUTPUT_DIR / "UPLOAD_FILE_MANIFEST_SHA256.tsv").write_text(
        "\n".join(manifest_lines) + "\n", encoding="utf-8"
    )

    validation = [
        "Plant-Environment Interactions submission validation",
        f"Build date: {date.today().isoformat()}",
        f"Article type: {ARTICLE_TYPE}",
        f"Abstract: PASS ({abstract_words} words; maximum 250)",
        (
            "Main body (Introduction + Results + Discussion): PASS "
            f"({main_words} words; maximum 3,000)"
        ),
        f"Materials and Methods: INFO ({methods_words} words; counted separately)",
        f"Conclusions: INFO ({conclusions_words} words; counted separately)",
        (
            "Total narrative (Introduction through Conclusions): INFO "
            f"({narrative_words} words)"
        ),
        f"Running title: PASS ({len(RUNNING_TITLE)} characters; required <40)",
        "Keywords: PASS (6; required 5–8)",
        f"References: PASS ({reference_count}; journal optimum 30)",
        "Main illustrations: PASS (5 figures + 2 tables = 7; maximum 10)",
        (
            "Main DOCX structure: PASS "
            f"({manuscript_validation['paragraphs']} paragraphs; 5 embedded figures; "
            "2 editable tables; "
            f"continuous line numbering; page-number field)"
        ),
        (
            "Supporting DOCX structure: PASS "
            f"({support_validation['paragraphs']} paragraphs; 2 editable tables; 4 figures)"
        ),
        f"Cover-letter DOCX structure: PASS ({cover_validation['paragraphs']} paragraphs)",
        "Main figures embedded and retained at original vector quality: PASS (5 embedded images + 5 PDFs)",
        f"Supporting data/code archive MD5 matches verified Zenodo copy: PASS ({EXPECTED_SUPPLEMENT_MD5})",
        f"Recommended upload files checksummed: PASS ({len(upload_files)})",
        "Manual inspection in Microsoft Word and of the portal-generated proof remains required.",
    ]
    (OUTPUT_DIR / "VALIDATION_REPORT.txt").write_text(
        "\n".join(validation) + "\n", encoding="utf-8"
    )

    print("Built Plant-Environment Interactions submission package")
    print(f"Abstract: {abstract_words} words")
    print(f"Main body (Introduction + Results + Discussion): {main_words} words")
    print(f"Materials and Methods: {methods_words} words")
    print(f"Conclusions: {conclusions_words} words")
    print(f"Total narrative (Introduction through Conclusions): {narrative_words} words")
    print(f"References: {reference_count}")
    for path in upload_files:
        print(f"{path.relative_to(OUTPUT_DIR)}\t{path.stat().st_size} bytes")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
