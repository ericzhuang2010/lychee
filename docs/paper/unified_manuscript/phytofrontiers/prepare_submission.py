#!/usr/bin/env python3
"""Build a PhytoFrontiers submission package from the unified manuscript."""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


OUTPUT_DIR = Path(__file__).resolve().parent
MANUSCRIPT_DIR = OUTPUT_DIR.parent
PROJECT_ROOT = OUTPUT_DIR.parents[3]
SOURCE_MD = MANUSCRIPT_DIR / "manuscript.md"

# Reuse the tested Markdown-to-Word helpers used by the Plant Direct package.
sys.path.insert(0, str(MANUSCRIPT_DIR))
from plant_direct import prepare_submission as common  # noqa: E402


TITLE = (
    "Cultivar-dependent transcriptional responses of lychee to "
    "Peronophythora litchii: a registered genome-wide analysis"
)
RUNNING_TITLE = "Cultivar-dependent lychee responses"
AUTHOR = "Eric Zhuang"
AFFILIATION = "NYU Langone Health, New York, NY, USA"
EMAIL = "eric.zhuang@nyulangone.org"
ORCID = "0009-0001-9050-0214"
ZENODO_DOI = "https://doi.org/10.5281/zenodo.22436625"

ABSTRACT = (
    "Litchi downy blight, caused by the oomycete *Peronophythora litchii*, is a "
    "major threat to lychee production. Public RNA-sequencing cohorts offer a "
    "resource for identifying cultivar-dependent infection responses but are "
    "readily overinterpreted when candidate selection and confirmation use the "
    "same data. We combined an exploratory reanalysis with a prospectively "
    "registered genome-wide discovery and external-evaluation workflow. A "
    "negative-binomial cultivar-by-infection analysis of 19,445 expressed genes "
    "yielded 262 statistical candidates; 206 passed mapping and gene-model quality "
    "control, 19 remained significant under an independent statistical framework, "
    "and 16 passed all registered robustness checks. In a prespecified cross-context "
    "evaluation using an independent pericarp time course, two genes were supported "
    "and five showed significant opposite-direction effects; the internally robust "
    "and externally supported sets did not overlap. None of 18 exploratory candidates "
    "met the subsequent genome-wide interaction criterion. Deterministic evidence "
    "integration produced no highest-confidence candidates, retained 12 internally "
    "robust genes in a middle tier, and retired 65 entities. These results show that "
    "prominent single-cohort candidates often fail formal interaction testing and "
    "cross-context evaluation. The openly archived workflow, complete results, and "
    "explicit null and contradictory findings provide a reproducible baseline for "
    "experimental validation."
)

KEYWORDS = (
    "downy blight; plant-oomycete interaction; RNA sequencing; transcriptomics; "
    "cross-context evaluation; preregistration"
)

FIGURES = [
    (1, "figure1_study_design_qc.pdf", "Figure_1_study_design_and_QC.pdf"),
    (2, "figure2_discovery_legacy.pdf", "Figure_2_discovery_and_legacy_audit.pdf"),
    (3, "figure3_robustness_external.pdf", "Figure_3_robustness_and_external_evaluation.pdf"),
    (4, "figure4_pathways_signatures.pdf", "Figure_4_pathway_and_signature_evaluation.pdf"),
    (5, "figure5_transcript_usage.pdf", "Figure_5_transcript_usage.pdf"),
    (6, "figure6_orthogonal_tiers.pdf", "Figure_6_orthogonal_evidence_and_tiers.pdf"),
]

SUPPLEMENTARY_FIGURES = [
    (
        "Figure S1",
        PROJECT_ROOT / "results/figures/FigureS1_replicate_level_counts.png",
        "Replicate-level normalized counts for the two cross-context-supported genes, "
        "twelve Tier B genes, and two legacy highlights. Points are individual "
        "deposited libraries; bars show cultivar-treatment medians.",
    ),
    (
        "Figure S2",
        PROJECT_ROOT / "results/figures/FigureS2_power_analysis.png",
        "Parametric detection probability for cultivar-by-infection effects under "
        "genome-wide discovery and candidate-family external adjustment. Curves show "
        "the overall result and mean-expression quartiles; the dashed line marks 80% power.",
    ),
    (
        "Figure S3",
        MANUSCRIPT_DIR / "figures/figureS3_exploratory_signature.png",
        "Quarantined PRJNA1090613 signed-signature estimate, retained as exploratory "
        "rather than confirmatory evidence.",
    ),
]

SUPPLEMENTAL_ITEMS = [
    "Table S1. Biological-unit registry.",
    "Table S2. Per-library quality control.",
    "Table S3. Genome-wide discovery statistics.",
    "Table S4. Robustness file manifest.",
    "Table S5. All external frozen tests.",
    "Table S6. Pathway and signature tests.",
    "Table S7. Conditional differential-transcript-usage results.",
    "Table S8. Candidate annotation and orthology evidence.",
    "Table S9. Small-RNA reference-gate outcome.",
    "Table S10. Motif-background tests and sensitivity results.",
    "Table S11. Accession-aware evidence registry.",
    "Table S12. Scripts, environments, and command inventory.",
    "Table S13. Protocol amendment and deviation log.",
    "Table S14. Legacy within-cultivar audit.",
    "Table S15. Controlled promoter-motif background comparison.",
    "Table S16a. Reconstructed dataset-search queries.",
    "Table S16b. Dataset eligibility decisions.",
    "Table S17. Exact software versions.",
    "Table S18. Power-simulation minimum detectable effects.",
    "Figure S1. Replicate-level normalized counts.",
    "Figure S2. Conditional parametric power curves.",
    "Figure S3. Exploratory signed-signature estimate.",
    "Data S1. Tab-separated source data for all analytical figures.",
    "Code S1. Analysis code, workflows, configurations, tests, metadata, and environments.",
]


# APS uses author-year citations. Database accessions 8 and 14 are stated in the text
# rather than included in Literature Cited, following the APS accession-number guidance.
CITE_LABELS = {
    1: "Morton 1987",
    2: "Sun et al. 2017",
    3: "Yi et al. 2010",
    4: "Hu et al. 2022",
    5: "Sun et al. 2019",
    6: "Liu et al. 2023",
    7: "Li et al. 2023",
    8: None,
    9: "Berka et al. 2022",
    10: "Andrews 2010",
    11: "Bolger et al. 2014",
    12: "Kim et al. 2013",
    13: "Anders et al. 2015",
    14: None,
    15: "Love et al. 2014",
    16: "Chen et al. 2020",
    17: "Bailey and Elkan 1994",
    18: "Lescot et al. 2002",
    19: "Benjamini and Hochberg 1995",
    20: "Chen et al. 2018",
    21: "Dobin et al. 2013",
    22: "Liao et al. 2014",
    23: "Patro et al. 2017",
    24: "Pockrandt et al. 2020",
    25: "Zhu et al. 2019",
    26: "Gupta et al. 2024",
    27: "Buchfink et al. 2015",
    28: "Wu and Smyth 2012",
    29: "Wu et al. 2010",
    30: "Korotkevich et al. 2021",
    31: "Nowicka and Robinson 2016",
    32: "Anders et al. 2012",
    33: "Van den Berge et al. 2017",
    34: "Robinson et al. 2010",
    35: "UniProt Consortium 2023",
    36: "Jones et al. 2014",
    37: "Rauluseviciute et al. 2024",
    38: "McLeay and Bailey 2010",
    39: "Grant et al. 2011",
    40: "Guo et al. 2020",
    41: "Mölder et al. 2021",
}

# The abbreviated in-text labels obscure the second-author ordering of the two Wu
# papers. These keys preserve APS's author-by-author alphabetization within a citation.
CITE_SORT_KEYS = {
    28: "wu|smyth|2012",
    29: "wu|lim|2010",
}

CITATION_RE = re.compile(r"\[(\d+(?:[–-]\d+)?(?:,\d+(?:[–-]\d+)?)*)\]")
IMAGE_RE = re.compile(r"^!\[(?P<caption>.+)]\((?P<path>[^)]+)\)$")
HEADING_RE = re.compile(r"^(#{2,3})\s+(.+)$")
TABLE_CAPTION_RE = re.compile(r"^\*\*Table ([1-4])\.\s*(.+?)\*\*(.*)$")


def citation_numbers(expression: str) -> list[int]:
    numbers: list[int] = []
    for part in expression.replace("–", "-").split(","):
        if "-" in part:
            start, end = (int(value) for value in part.split("-", 1))
            numbers.extend(range(start, end + 1))
        else:
            numbers.append(int(part))
    return numbers


def aps_citations(value: str) -> str:
    """Convert the source manuscript's numeric citations to APS author-year style."""
    narrative = {
        (
            "Sun et al. compared a susceptible and a resistant cultivar across a pericarp "
            "infection time course and reported that the two genotypes deploy distinct early "
            "response programs [5]"
        ): (
            "Sun et al. (2019) compared a susceptible and a resistant cultivar across a "
            "pericarp infection time course and reported that the two genotypes deploy "
            "distinct early response programs"
        ),
        (
            "Liu et al. described calcium-dependent protein kinase (LcCDPK) family members "
            "induced by *P. litchii* inoculation [6]"
        ): (
            "Liu et al. (2023) described calcium-dependent protein kinase (LcCDPK) family "
            "members induced by *P. litchii* inoculation"
        ),
        (
            "Li et al. showed that the pathogen RXLR effector PlAvh202 promotes virulence by "
            "destabilizing a host ethylene-biosynthesis enzyme [7]"
        ): (
            "Li et al. (2023) showed that the pathogen RXLR effector PlAvh202 promotes "
            "virulence by destabilizing a host ethylene-biosynthesis enzyme"
        ),
    }
    for source, target in narrative.items():
        value = value.replace(source, target)
    value = value.replace("the GSE201243 leaf series [8]", "the GSE201243 leaf series")
    value = value.replace(
        "(historically also treated as *Phytophthora litchii* [2])",
        "(historically also treated as *Phytophthora litchii*; Sun et al. 2017)",
    )

    def replacement(match: re.Match[str]) -> str:
        pairs = [
            (number, CITE_LABELS[number])
            for number in citation_numbers(match.group(1))
            if CITE_LABELS[number]
        ]
        pairs.sort(
            key=lambda item: CITE_SORT_KEYS.get(item[0], item[1].casefold())
        )
        labels = [label for _, label in pairs]
        return f"({'; '.join(labels)})" if labels else ""

    return CITATION_RE.sub(replacement, value)


def clean_heading(value: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\s+", "", value)


def section_bounds(lines: list[str], start_heading: str, end_heading: str) -> tuple[int, int]:
    return lines.index(start_heading) + 1, lines.index(end_heading)


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "000000")
        borders.append(element)


def set_header_bottom_border(row) -> None:
    for cell in row.cells:
        properties = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        borders.append(bottom)
        properties.append(borders)


def add_aps_table(document: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, source_row in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            common.set_repeat_table_header(row)
            set_header_bottom_border(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            common.set_cell_margins(cell, top=50, start=45, bottom=50, end=45)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            paragraph.paragraph_format.space_after = Pt(0)
            if column_index < len(source_row):
                common.add_inline(paragraph, aps_citations(source_row[column_index]))
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(12)
                if row_index == 0:
                    run.bold = True


def extract_tables(lines: list[str]) -> list[tuple[int, str, list[list[str]]]]:
    tables: list[tuple[int, str, list[list[str]]]] = []
    index = 0
    while index < len(lines):
        match = TABLE_CAPTION_RE.match(lines[index].strip())
        if not match:
            index += 1
            continue
        table_number = int(match.group(1))
        caption = f"Table {table_number}. {match.group(2)}{match.group(3)}"
        table_start = index + 1
        while table_start < len(lines) and not lines[table_start].strip():
            table_start += 1
        rows, index = common.parse_table(lines, table_start)
        tables.append((table_number, aps_citations(caption), rows))
    if [number for number, _, _ in tables] != [1, 2, 3, 4]:
        raise ValueError(f"Expected Tables 1-4, found {[item[0] for item in tables]}")
    return tables


def add_body_range(document: Document, lines: list[str], start: int, end: int) -> None:
    index = start
    while index < end:
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            paragraph = document.add_paragraph(style="Heading 2")
            common.add_inline(paragraph, clean_heading(heading.group(2)))
            index += 1
            continue

        image = IMAGE_RE.match(stripped)
        if image:
            number = re.match(r"Figure ([1-6])\.", image.group("caption"))
            if number:
                document.add_paragraph(
                    f"[Insert Figure {number.group(1)} near here]", style="Figure Callout"
                )
            index += 1
            continue

        table_caption = TABLE_CAPTION_RE.match(stripped)
        if table_caption:
            document.add_paragraph(
                f"[Insert Table {table_caption.group(1)} near here]", style="Figure Callout"
            )
            index += 1
            while index < end and not lines[index].strip():
                index += 1
            if index < end and lines[index].strip().startswith("|"):
                _, index = common.parse_table(lines, index)
            continue

        if stripped.startswith("|"):
            _, index = common.parse_table(lines, index)
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < end:
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                HEADING_RE.match(candidate)
                or IMAGE_RE.match(candidate)
                or TABLE_CAPTION_RE.match(candidate)
                or candidate.startswith("|")
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        common.add_inline(paragraph, aps_citations(" ".join(paragraph_lines)))


def add_title_page(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    common.add_inline(
        paragraph,
        "Cultivar-dependent transcriptional responses of lychee to "
        "*Peronophythora litchii*: a registered genome-wide analysis",
    )
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
    paragraph.paragraph_format.space_after = Pt(24)

    for line, bold in (
        (AUTHOR, True),
        (AFFILIATION, False),
        ("Corresponding author: Eric Zhuang", False),
        (f"Email: {EMAIL}", False),
        (f"ORCID iD: {ORCID}", False),
        ("Article type: Research", False),
        (f"Running title: {RUNNING_TITLE}", False),
    ):
        item = document.add_paragraph()
        item.alignment = WD_ALIGN_PARAGRAPH.CENTER
        item.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        item.add_run(line).bold = bold


def add_declarations(document: Document) -> None:
    sections = [
        (
            "Acknowledgments",
            "OpenAI Codex was used to assist with language editing and journal-specific "
            "document formatting. It was not used to generate data, conduct analyses, or "
            "generate scientific concepts, results, or interpretations. The author reviewed "
            "and accepts responsibility for all content.",
        ),
        ("Funding", "This work received no external funding."),
        (
            "Author Contributions",
            "E.Z. conceived the study, curated the data, developed the methodology and "
            "software, performed the analyses and validation, prepared the visualizations, "
            "and wrote and revised the manuscript.",
        ),
        ("Conflict of Interest", "The author declares no conflicts of interest."),
        (
            "Ethics Statement",
            "This study reanalyzed publicly available sequencing datasets and involved no "
            "new experiments with humans, human tissue, animals, or field sampling; ethical "
            "approval was not required.",
        ),
        (
            "Data and Code Availability",
            "All analyzed sequencing data are public under PRJNA830488/GSE201243, "
            "PRJNA450886, PRJNA922966/GSE222651, PRJNA922965/GSE222650, and "
            "PRJNA1090613/GSE262200. The registered protocol, complete result tables, "
            "supplementary figures, figure source data, analysis code, workflows, "
            f"configurations, tests, and environment specifications are archived at {ZENODO_DOI}.",
        ),
    ]
    for heading, body in sections:
        document.add_paragraph(heading, style="Heading 1")
        paragraph = document.add_paragraph()
        common.add_inline(paragraph, body)

    document.add_paragraph("Supplemental Materials", style="Heading 1")
    intro = document.add_paragraph()
    common.add_inline(
        intro,
        "The e-Xtra PDF contains Figures S1 to S3. The following complete, machine-readable "
        f"items are permanently archived at {ZENODO_DOI}:",
    )
    for item in SUPPLEMENTAL_ITEMS:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        paragraph.add_run(item)


LITERATURE_CITED = [
    "Anders, S., Pyl, P. T., and Huber, W. 2015. HTSeq—a Python framework to work with high-throughput sequencing data. Bioinformatics 31:166-169. https://doi.org/10.1093/bioinformatics/btu638.",
    "Anders, S., Reyes, A., and Huber, W. 2012. Detecting differential usage of exons from RNA-seq data. Genome Res. 22:2008-2017. https://doi.org/10.1101/gr.133744.111.",
    "Andrews, S. 2010. FastQC: A quality control tool for high throughput sequence data. Babraham Bioinformatics, Cambridge, United Kingdom. https://www.bioinformatics.babraham.ac.uk/projects/fastqc/.",
    "Bailey, T. L., and Elkan, C. 1994. Fitting a mixture model by expectation maximization to discover motifs in biopolymers. Pages 28-36 in: Proceedings of the Second International Conference on Intelligent Systems for Molecular Biology. AAAI Press, Menlo Park, CA.",
    "Benjamini, Y., and Hochberg, Y. 1995. Controlling the false discovery rate: A practical and powerful approach to multiple testing. J. R. Stat. Soc. Ser. B 57:289-300.",
    "Berka, M., Kopecká, R., Berková, V., Brzobohatý, B., and Černý, M. 2022. Regulation of heat shock proteins 70 and their role in plant immunity. J. Exp. Bot. 73:1894-1909. https://doi.org/10.1093/jxb/erab549.",
    "Bolger, A. M., Lohse, M., and Usadel, B. 2014. Trimmomatic: A flexible trimmer for Illumina sequence data. Bioinformatics 30:2114-2120. https://doi.org/10.1093/bioinformatics/btu170.",
    "Buchfink, B., Xie, C., and Huson, D. H. 2015. Fast and sensitive protein alignment using DIAMOND. Nat. Methods 12:59-60. https://doi.org/10.1038/nmeth.3176.",
    "Chen, C., Chen, H., Zhang, Y., Thomas, H. R., Frank, M. H., He, Y., and Xia, R. 2020. TBtools: An integrative toolkit developed for interactive analyses of big biological data. Mol. Plant 13:1194-1202. https://doi.org/10.1016/j.molp.2020.06.009.",
    "Chen, S., Zhou, Y., Chen, Y., and Gu, J. 2018. fastp: An ultra-fast all-in-one FASTQ preprocessor. Bioinformatics 34:i884-i890. https://doi.org/10.1093/bioinformatics/bty560.",
    "Dobin, A., Davis, C. A., Schlesinger, F., Drenkow, J., Zaleski, C., Jha, S., Batut, P., Chaisson, M., and Gingeras, T. R. 2013. STAR: Ultrafast universal RNA-seq aligner. Bioinformatics 29:15-21. https://doi.org/10.1093/bioinformatics/bts635.",
    "Grant, C. E., Bailey, T. L., and Noble, W. S. 2011. FIMO: Scanning for occurrences of a given motif. Bioinformatics 27:1017-1018. https://doi.org/10.1093/bioinformatics/btr064.",
    "Guo, Z., Kuang, Z., Wang, Y., Zhao, Y., Tao, Y., Cheng, C., Yang, J., Lu, X., Hao, C., Wang, T., Cao, X., Wei, J., Li, L., and Yang, X. 2020. PmiREN: A comprehensive encyclopedia of plant miRNAs. Nucleic Acids Res. 48:D1114-D1121. https://doi.org/10.1093/nar/gkz894.",
    "Gupta, P., Elser, J., Hooks, E., D'Eustachio, P., Jaiswal, P., and Naithani, S. 2024. Plant Reactome Knowledgebase: Empowering plant pathway exploration and OMICS data analysis. Nucleic Acids Res. 52:D1538-D1547. https://doi.org/10.1093/nar/gkad1052.",
    "Hu, G., Feng, J., Xiang, X., Wang, J., Salojärvi, J., Liu, C., Wu, Z., Zhang, J., Liang, X., Jiang, Z., Liu, W., Ou, L., Li, J., Fan, G., Mai, Y., Chen, C., Zhang, X., Zheng, J., Zhang, Y., Peng, H., Yao, L., Wai, C. M., Luo, X., Fu, J., Tang, H., Lan, T., Lai, B., Sun, J., Wei, Y., Li, H., Chen, J., Huang, X., Yan, Q., Liu, X., McHale, L. K., Rolling, W., Guyot, R., Sankoff, D., Zheng, C., Albert, V. A., Ming, R., Chen, H., Xia, R., and Li, J. 2022. Two divergent haplotypes from a highly heterozygous lychee genome suggest independent domestication events for early and late-maturing cultivars. Nat. Genet. 54:73-83. https://doi.org/10.1038/s41588-021-00971-3.",
    "Jones, P., Binns, D., Chang, H.-Y., Fraser, M., Li, W., McAnulla, C., McWilliam, H., Maslen, J., Mitchell, A., Nuka, G., Pesseat, S., Quinn, A. F., Sangrador-Vegas, A., Scheremetjew, M., Yong, S.-Y., Lopez, R., and Hunter, S. 2014. InterProScan 5: Genome-scale protein function classification. Bioinformatics 30:1236-1240. https://doi.org/10.1093/bioinformatics/btu031.",
    "Kim, D., Pertea, G., Trapnell, C., Pimentel, H., Kelley, R., and Salzberg, S. L. 2013. TopHat2: Accurate alignment of transcriptomes in the presence of insertions, deletions and gene fusions. Genome Biol. 14:R36. https://doi.org/10.1186/gb-2013-14-4-r36.",
    "Korotkevich, G., Sukhov, V., Budin, N., Shpak, B., Artyomov, M. N., and Sergushichev, A. 2021. Fast gene set enrichment analysis. bioRxiv 060012. https://doi.org/10.1101/060012.",
    "Lescot, M., Déhais, P., Thijs, G., Marchal, K., Moreau, Y., Van de Peer, Y., Rouzé, P., and Rombauts, S. 2002. PlantCARE, a database of plant cis-acting regulatory elements and a portal to tools for in silico analysis of promoter sequences. Nucleic Acids Res. 30:325-327. https://doi.org/10.1093/nar/30.1.325.",
    "Li, P., Li, W., Zhou, X., Situ, J., Xie, L., Xi, P., Yang, B., Kong, G., and Jiang, Z. 2023. *Peronophythora litchii* RXLR effector PlAvh202 destabilizes a host ethylene biosynthesis enzyme. Plant Physiol. 193:756-774. https://doi.org/10.1093/plphys/kiad311.",
    "Liao, Y., Smyth, G. K., and Shi, W. 2014. featureCounts: An efficient general purpose program for assigning sequence reads to genomic features. Bioinformatics 30:923-930. https://doi.org/10.1093/bioinformatics/btt656.",
    "Liu, H., Yan, Q., Jiang, Y., Shi, F., Chen, J., Cai, C., and Ou, L. 2023. Identification of LcCDPKs and analysis of their expression patterns in response to downy mildew stresses in lychee. J. Fruit Sci. 40:442-456. https://doi.org/10.13925/j.cnki.gsxb.20220307.",
    "Love, M. I., Huber, W., and Anders, S. 2014. Moderated estimation of fold change and dispersion for RNA-seq data with DESeq2. Genome Biol. 15:550. https://doi.org/10.1186/s13059-014-0550-8.",
    "McLeay, R. C., and Bailey, T. L. 2010. Motif enrichment analysis: A unified framework and an evaluation on ChIP data. BMC Bioinformatics 11:165. https://doi.org/10.1186/1471-2105-11-165.",
    "Mölder, F., Jablonski, K. P., Letcher, B., Hall, M. B., Tomkins-Tinch, C. H., Sochat, V., Forster, J., Lee, S., Twardziok, S. O., Kanitz, A., Wilm, A., Holtgrewe, M., Rahmann, S., Nahnsen, S., and Köster, J. 2021. Sustainable data analysis with Snakemake. F1000Research 10:33. https://doi.org/10.12688/f1000research.29032.2.",
    "Morton, J. F. 1987. Lychee. Pages 249-259 in: Fruits of Warm Climates. Creative Resource Systems, Winterville, NC.",
    "Nowicka, M., and Robinson, M. D. 2016. DRIMSeq: A Dirichlet-multinomial framework for multivariate count outcomes in genomics. F1000Research 5:1356. https://doi.org/10.12688/f1000research.8900.2.",
    "Patro, R., Duggal, G., Love, M. I., Irizarry, R. A., and Kingsford, C. 2017. Salmon provides fast and bias-aware quantification of transcript expression. Nat. Methods 14:417-419. https://doi.org/10.1038/nmeth.4197.",
    "Pockrandt, C., Alzamel, M., Iliopoulos, C. S., and Reinert, K. 2020. GenMap: Ultra-fast computation of genome mappability. Bioinformatics 36:3687-3692. https://doi.org/10.1093/bioinformatics/btaa222.",
    "Rauluseviciute, I., Riudavets-Puig, R., Blanc-Mathieu, R., Castro-Mondragon, J. A., Ferenc, K., Kumar, V., Lemma, R. B., Lucas, J., Chèneby, J., Baranasic, D., Khan, A., Fornes, O., Gundersen, S., Johansen, M., Hovig, E., Lenhard, B., Sandelin, A., Wasserman, W. W., Parcy, F., and Mathelier, A. 2024. JASPAR 2024: 20th anniversary of the open-access database of transcription factor binding profiles. Nucleic Acids Res. 52:D174-D182. https://doi.org/10.1093/nar/gkad1059.",
    "Robinson, M. D., McCarthy, D. J., and Smyth, G. K. 2010. edgeR: A Bioconductor package for differential expression analysis of digital gene expression data. Bioinformatics 26:139-140. https://doi.org/10.1093/bioinformatics/btp616.",
    "Sun, J., Cao, L., Li, H., Wang, G., Wang, S., Li, F., Zou, X., and Wang, J. 2019. Early responses given distinct tactics to infection of *Peronophythora litchii* in susceptible and resistant litchi cultivar. Sci. Rep. 9:2810. https://doi.org/10.1038/s41598-019-39100-w.",
    "Sun, J., Gao, Z., Zhang, X., Zou, X., Cao, L., and Wang, J. 2017. Transcriptome analysis of *Phytophthora litchii* reveals pathogenicity arsenals and confirms taxonomic status. PLoS One 12:e0178245. https://doi.org/10.1371/journal.pone.0178245.",
    "UniProt Consortium. 2023. UniProt: The Universal Protein Knowledgebase in 2023. Nucleic Acids Res. 51:D523-D531. https://doi.org/10.1093/nar/gkac1052.",
    "Van den Berge, K., Soneson, C., Robinson, M. D., and Clement, L. 2017. stageR: A general stage-wise method for controlling the gene-level false discovery rate in differential expression and differential transcript usage. Genome Biol. 18:151. https://doi.org/10.1186/s13059-017-1277-0.",
    "Wu, D., Lim, E., Vaillant, F., Asselin-Labat, M. L., Visvader, J. E., and Smyth, G. K. 2010. ROAST: Rotation gene set tests for complex microarray experiments. Bioinformatics 26:2176-2182. https://doi.org/10.1093/bioinformatics/btq401.",
    "Wu, D., and Smyth, G. K. 2012. Camera: A competitive gene set test accounting for inter-gene correlation. Nucleic Acids Res. 40:e133. https://doi.org/10.1093/nar/gks461.",
    "Yi, C., Jiang, Y., Shi, J., Qu, H., Xue, S., Duan, X., Shi, J., and Prasad, N. K. 2010. ATP-regulation of antioxidant properties and phenolics in litchi fruit during browning and pathogen infection process. Food Chem. 118:42-47. https://doi.org/10.1016/j.foodchem.2009.04.074.",
    "Zhu, A., Ibrahim, J. G., and Love, M. I. 2019. Heavy-tailed prior distributions for sequence count data: Removing the noise and preserving large differences. Bioinformatics 35:2084-2092. https://doi.org/10.1093/bioinformatics/bty895.",
]


def add_literature_cited(document: Document) -> None:
    heading = document.add_paragraph("Literature Cited", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    for reference in LITERATURE_CITED:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        paragraph.paragraph_format.space_after = Pt(0)
        common.add_inline(paragraph, reference)


def add_tables(document: Document, tables: list[tuple[int, str, list[list[str]]]]) -> None:
    for number, caption, rows in tables:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        prefix = f"Table {number}."
        paragraph.add_run(prefix).bold = True
        common.add_inline(paragraph, caption[len(prefix) :])
        common.set_keep_with_next(paragraph)
        add_aps_table(document, readable_review_table(number, rows))


def set_keep_together(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    properties.append(OxmlElement("w:keepLines"))


def add_embedded_figures(document: Document, legends: dict[int, str]) -> None:
    for number in range(1, 7):
        source_pdf = next(source for item, source, _ in FIGURES if item == number)
        source_png = MANUSCRIPT_DIR / "figures" / source_pdf.replace(".pdf", ".png")
        if not source_png.is_file():
            raise FileNotFoundError(source_png)
        with Image.open(source_png) as image:
            width, height = image.size
        scale = min(6.35 / width, 5.1 / height)
        image_paragraph = document.add_paragraph()
        image_paragraph.paragraph_format.page_break_before = True
        image_paragraph.paragraph_format.keep_with_next = True
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(
            str(source_png), width=Inches(width * scale), height=Inches(height * scale)
        )

        caption = aps_citations(legends[number])
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        set_keep_together(paragraph)
        prefix = f"Figure {number}."
        paragraph.add_run(prefix).bold = True
        common.add_inline(paragraph, caption[len(prefix) :])


def configure_aps_review_document(document: Document) -> None:
    common.configure_document(document, line_numbers=True)
    for style_name in ("Normal", "Heading 1", "Heading 2", "Figure Callout", "List Bullet"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(12)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    section = document.sections[0]
    section.footer.paragraphs[0].clear()
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.line_spacing = 1
    run = header.add_run("Zhuang | PhytoFrontiers | Page ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    page_run = header.add_run()
    page_run.font.name = "Arial"
    page_run.font.size = Pt(9)
    common.add_field(page_run, " PAGE ")


def validate_source_citations(lines: list[str]) -> None:
    body = "\n".join(lines[: lines.index("## References")])
    used: set[int] = set()
    for match in CITATION_RE.finditer(body):
        used.update(citation_numbers(match.group(1)))
    unknown = used.difference(CITE_LABELS)
    if unknown:
        raise ValueError(f"Unmapped source citations: {sorted(unknown)}")


def build_manuscript() -> Path:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    validate_source_citations(lines)
    if len(TITLE) > 150:
        raise ValueError(f"Title has {len(TITLE)} characters; APS limit is 150")
    word_count = len(re.findall(r"\b[\w'-]+\b", common.plain_markdown(ABSTRACT)))
    if word_count > 250:
        raise ValueError(f"Abstract has {word_count} words; APS limit is 250")

    legends = common.extract_figure_legends(lines)
    tables = extract_tables(lines)
    intro = section_bounds(lines, "## 1. Introduction", "## 2. Results")
    results = section_bounds(lines, "## 2. Results", "## 3. Discussion")
    discussion = section_bounds(lines, "## 3. Discussion", "## 4. Conclusions")
    conclusions = section_bounds(lines, "## 4. Conclusions", "## 5. Materials and Methods")
    methods = section_bounds(lines, "## 5. Materials and Methods", "## Declarations")

    document = Document()
    configure_aps_review_document(document)
    properties = document.core_properties
    properties.title = TITLE
    properties.author = AUTHOR
    properties.subject = "Research manuscript submitted to PhytoFrontiers"
    properties.keywords = KEYWORDS

    add_title_page(document)
    heading = document.add_paragraph("Abstract", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    common.add_inline(document.add_paragraph(), ABSTRACT)
    keywords = document.add_paragraph()
    keywords.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    keywords.add_run("Keywords: ").bold = True
    keywords.add_run(KEYWORDS)

    for heading_text, bounds in (
        ("Introduction", intro),
        ("Materials and Methods", methods),
        ("Results", results),
        ("Discussion", discussion),
    ):
        document.add_paragraph(heading_text, style="Heading 1")
        add_body_range(document, lines, *bounds)

    document.add_paragraph("Conclusions", style="Heading 2")
    add_body_range(document, lines, *conclusions)
    add_declarations(document)
    add_literature_cited(document)
    add_tables(document, tables)
    add_embedded_figures(document, legends)

    output = OUTPUT_DIR / "PhytoFrontiers_manuscript.docx"
    document.save(output)
    return output


LATEX_INLINE_TOKEN = re.compile(
    r"(× 10[⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|\[[^]]+\]\([^)]+\)|https?://\S+)"
)

SUPERSCRIPT_DIGITS = str.maketrans("⁻⁰¹²³⁴⁵⁶⁷⁸⁹", "-0123456789")


def latex_escape(value: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(character, character) for character in value)


def latex_inline(value: str) -> str:
    """Convert the small Markdown subset used by the manuscript to XeLaTeX."""
    output: list[str] = []
    cursor = 0
    for match in LATEX_INLINE_TOKEN.finditer(value):
        if match.start() > cursor:
            output.append(latex_escape(value[cursor : match.start()]))
        token = match.group(0)
        if token.startswith("× 10"):
            exponent = token[4:].translate(SUPERSCRIPT_DIGITS)
            output.append("× 10" + r"\textsuperscript{" + exponent + "}")
        elif token.startswith("**"):
            output.append(r"\textbf{" + latex_escape(token[2:-2]) + "}")
        elif token.startswith("*"):
            output.append(r"\textit{" + latex_escape(token[1:-1]) + "}")
        elif token.startswith("`"):
            output.append(r"\texttt{" + latex_escape(token[1:-1]) + "}")
        elif token.startswith("["):
            link = re.match(r"\[([^]]+)]\(([^)]+)\)", token)
            if not link:
                raise ValueError(f"Malformed Markdown link: {token}")
            output.append(r"\href{" + link.group(2) + "}{" + latex_escape(link.group(1)) + "}")
        else:
            trailing = "." if token.endswith(".") else ""
            url = token[:-1] if trailing else token
            output.append(r"\url{" + url + "}" + trailing)
        cursor = match.end()
    if cursor < len(value):
        output.append(latex_escape(value[cursor:]))
    return "".join(output)


def latex_body_range(lines: list[str], start: int, end: int) -> str:
    output: list[str] = []
    index = start
    while index < end:
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            output.append(r"\subsection*{" + latex_inline(clean_heading(heading.group(2))) + "}")
            index += 1
            continue

        if IMAGE_RE.match(stripped):
            index += 1
            continue

        if TABLE_CAPTION_RE.match(stripped):
            index += 1
            while index < end and not lines[index].strip():
                index += 1
            if index < end and lines[index].strip().startswith("|"):
                _, index = common.parse_table(lines, index)
            continue

        if stripped.startswith("|"):
            _, index = common.parse_table(lines, index)
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < end:
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                HEADING_RE.match(candidate)
                or IMAGE_RE.match(candidate)
                or TABLE_CAPTION_RE.match(candidate)
                or candidate.startswith("|")
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        output.append(latex_inline(aps_citations(" ".join(paragraph_lines))))
    return "\n\n".join(output)


TABLE_WIDTHS = {
    1: (0.17, 0.31, 0.22, 0.14),
    2: (0.16, 0.34, 0.34),
    3: (0.22, 0.17, 0.25, 0.20),
    4: (0.35, 0.23, 0.26),
}


def readable_review_table(number: int, rows: list[list[str]]) -> list[list[str]]:
    """Combine tightly related fields so review tables remain legible at 12 pt."""
    if number == 1:
        return rows
    if number == 2:
        output = [["Gene", "Exploratory annotation", "Re-estimated interaction"]]
        output.extend(
            [
                f"**{row[0]}**",
                row[1],
                f"log2FC (SE): {row[2]}; genome-wide *q*: {row[3]}; "
                f"uniform mappability: {row[4]}",
            ]
            for row in rows[1:]
        )
        return output
    if number == 3:
        output = [["Gene and annotation", "Discovery interaction", "External interaction", "Outcome"]]
        output.extend(
            [
                f"**{row[0]}**; {row[1]}",
                f"log2FC: {row[2]}; genome-wide *q*: {row[3]}",
                f"log2FC (95% CI): {row[4]}; family-adjusted *q*: {row[5]}",
                row[6],
            ]
            for row in rows[1:]
        )
        return output
    if number == 4:
        output = [["Gene and family-level annotation", "Discovery interaction", "Primary external evidence"]]
        output.extend(
            [
                f"**{row[0]}**; {row[1]}",
                f"log2FC: {row[2]}; genome-wide *q*: {row[3]}",
                f"log2FC (95% CI): {row[4]}; family-adjusted *q*: {row[5]}",
            ]
            for row in rows[1:]
        )
        return output
    raise ValueError(f"Unexpected table number: {number}")


def latex_table(number: int, caption: str, rows: list[list[str]]) -> str:
    rows = readable_review_table(number, rows)
    widths = TABLE_WIDTHS[number]
    if len(widths) != len(rows[0]):
        raise ValueError(f"Table {number} has {len(rows[0])} columns, expected {len(widths)}")
    columns = "".join(
        rf">{{\raggedright\arraybackslash}}p{{{width:.2f}\textwidth}}" for width in widths
    )
    prefix = f"Table {number}."
    if not caption.startswith(prefix):
        raise ValueError(f"Unexpected caption for Table {number}: {caption}")
    caption_body = caption[len(prefix) :].strip()
    header = " & ".join(r"\textbf{" + latex_inline(cell) + "}" for cell in rows[0])
    body = "\n".join(" & ".join(latex_inline(cell) for cell in row) + r" \\" for row in rows[1:])
    output = [r"\clearpage"]
    output.extend(
        [
            r"\setlength{\tabcolsep}{5pt}",
            r"\begin{longtable}{@{}" + columns + r"@{}}",
            r"\caption{" + latex_inline(caption_body) + r"} \\" ,
            r"\toprule",
            header + r" \\",
            r"\midrule",
            r"\endfirsthead",
            rf"\multicolumn{{{len(widths)}}}{{l}}{{\textit{{Table {number} continued}}}} \\",
            r"\toprule",
            header + r" \\",
            r"\midrule",
            r"\endhead",
            rf"\midrule \multicolumn{{{len(widths)}}}{{r}}{{\textit{{Continued on next page}}}} \\",
            r"\endfoot",
            r"\bottomrule",
            r"\endlastfoot",
            body,
            r"\end{longtable}",
        ]
    )
    return "\n".join(output)


def latex_figure_pages(legends: dict[int, str]) -> str:
    output: list[str] = []
    for number in range(1, 7):
        source_name = next(source for item, source, _ in FIGURES if item == number)
        source = (MANUSCRIPT_DIR / "figures" / source_name).resolve()
        caption = aps_citations(legends[number])
        prefix = f"Figure {number}."
        output.extend(
            [
                r"\clearpage",
                r"\nolinenumbers",
                r"\begin{center}",
                r"\includegraphics[width=\textwidth,height=4.9in,keepaspectratio]{\detokenize{"
                + str(source)
                + "}}",
                r"\end{center}",
                r"\linenumbers",
                r"\noindent\textbf{" + latex_escape(prefix) + "} "
                + latex_inline(caption[len(prefix) :]),
            ]
        )
    return "\n".join(output)


def build_initial_submission_tex() -> Path:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    legends = common.extract_figure_legends(lines)
    tables = extract_tables(lines)
    intro = section_bounds(lines, "## 1. Introduction", "## 2. Results")
    results = section_bounds(lines, "## 2. Results", "## 3. Discussion")
    discussion = section_bounds(lines, "## 3. Discussion", "## 4. Conclusions")
    conclusions = section_bounds(lines, "## 4. Conclusions", "## 5. Materials and Methods")
    methods = section_bounds(lines, "## 5. Materials and Methods", "## Declarations")

    preamble = r"""\documentclass[12pt,letterpaper]{article}
\usepackage[margin=1in,headheight=15pt]{geometry}
\usepackage{fontspec}
\setmainfont{Arial}
\setsansfont{Arial}
\setmonofont{Arial}
\usepackage{graphicx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{longtable}
\usepackage{caption}
\usepackage{setspace}
\usepackage{titlesec}
\usepackage[switch]{lineno}
\usepackage{fancyhdr}
\usepackage[hidelinks]{hyperref}
\usepackage{microtype}
\pagestyle{fancy}
\fancyhf{}
\lhead{\small Zhuang}
\chead{\small PhytoFrontiers}
\rhead{\small Page \thepage}
\renewcommand{\headrulewidth}{0pt}
\captionsetup[table]{labelsep=period,justification=raggedright,singlelinecheck=false}
\titleformat{\section}{\normalfont\normalsize\bfseries}{}{0pt}{}
\titleformat{\subsection}{\normalfont\normalsize\bfseries}{}{0pt}{}
\titlespacing*{\section}{0pt}{2.5ex plus 1ex minus .2ex}{1.3ex}
\titlespacing*{\subsection}{0pt}{2.25ex plus 1ex minus .2ex}{1ex}
\setlength{\parindent}{1.2em}
\setlength{\parskip}{0pt}
\emergencystretch=3em
\raggedbottom
\sloppy
\doublespacing
\modulolinenumbers[1]
\begin{document}
\linenumbers
\thispagestyle{fancy}
"""

    sections = [
        preamble,
        r"\begin{center}",
        r"{\normalsize\bfseries "
        + latex_inline(
            "Cultivar-dependent transcriptional responses of lychee to "
            "*Peronophythora litchii*: a registered genome-wide analysis"
        )
        + r"\par}",
        r"\vspace{1.5em}",
        r"\textbf{" + latex_escape(AUTHOR) + r"}\\",
        latex_escape(AFFILIATION) + r"\\",
        r"Corresponding author: " + latex_escape(AUTHOR) + r"\\",
        r"Email: \href{mailto:" + EMAIL + "}{" + latex_escape(EMAIL) + r"}\\",
        r"ORCID iD: " + latex_escape(ORCID) + r"\\",
        r"Article type: Research\\",
        r"Running title: " + latex_escape(RUNNING_TITLE),
        r"\end{center}",
        r"\section*{Abstract}",
        latex_inline(ABSTRACT),
        r"\noindent\textbf{Keywords:} " + latex_escape(KEYWORDS),
        r"\section*{Introduction}",
        latex_body_range(lines, *intro),
        r"\section*{Materials and Methods}",
        latex_body_range(lines, *methods),
        r"\section*{Results}",
        latex_body_range(lines, *results),
        r"\section*{Discussion}",
        latex_body_range(lines, *discussion),
        r"\subsection*{Conclusions}",
        latex_body_range(lines, *conclusions),
        r"\section*{Acknowledgments}",
        latex_inline(
            "OpenAI Codex was used to assist with language editing and journal-specific "
            "document formatting. It was not used to generate data, conduct analyses, or "
            "generate scientific concepts, results, or interpretations. The author reviewed "
            "and accepts responsibility for all content."
        ),
        r"\section*{Funding}",
        "This work received no external funding.",
        r"\section*{Author Contributions}",
        "E.Z. conceived the study, curated the data, developed the methodology and software, "
        "performed the analyses and validation, prepared the visualizations, and wrote and "
        "revised the manuscript.",
        r"\section*{Conflict of Interest}",
        "The author declares no conflicts of interest.",
        r"\section*{Ethics Statement}",
        "This study reanalyzed publicly available sequencing datasets and involved no new "
        "experiments with humans, human tissue, animals, or field sampling; ethical approval "
        "was not required.",
        r"\section*{Data and Code Availability}",
        latex_inline(
            "All analyzed sequencing data are public under PRJNA830488/GSE201243, "
            "PRJNA450886, PRJNA922966/GSE222651, PRJNA922965/GSE222650, and "
            "PRJNA1090613/GSE262200. The registered protocol, complete result tables, "
            "supplementary figures, figure source data, analysis code, workflows, "
            f"configurations, tests, and environment specifications are archived at {ZENODO_DOI}."
        ),
        r"\section*{Supplemental Materials}",
        latex_inline(
            "The e-Xtra PDF contains Figures S1 to S3. Complete machine-readable Tables S1 "
            f"to S18, figure source data, the registered protocol, and analysis code are "
            f"permanently archived at {ZENODO_DOI}."
        ),
        r"\clearpage\section*{Literature Cited}",
    ]

    for reference in LITERATURE_CITED:
        sections.append(r"\noindent\hangindent=1.5em " + latex_inline(reference) + r"\par")
    for number, caption, rows in tables:
        sections.append(latex_table(number, caption, rows))
    sections.append(latex_figure_pages(legends))
    sections.append(r"\end{document}")

    output = OUTPUT_DIR / "PhytoFrontiers_initial_submission.tex"
    output.write_text("\n\n".join(sections) + "\n", encoding="utf-8")
    return output


def build_initial_submission_pdf() -> Path:
    tex_path = build_initial_submission_tex()
    output = OUTPUT_DIR / "PhytoFrontiers_initial_submission.pdf"
    with TemporaryDirectory(prefix="phytofrontiers_tex_") as temporary:
        command = [
            "xelatex",
            "-interaction=nonstopmode",
            "-halt-on-error",
            f"-output-directory={temporary}",
            str(tex_path),
        ]
        for _ in range(2):
            result = subprocess.run(
                command,
                cwd=OUTPUT_DIR,
                check=False,
                capture_output=True,
                text=True,
            )
            if result.returncode != 0:
                tail = "\n".join((result.stdout + result.stderr).splitlines()[-80:])
                raise RuntimeError(f"XeLaTeX failed:\n{tail}")
        built = Path(temporary) / "PhytoFrontiers_initial_submission.pdf"
        if not built.is_file():
            raise FileNotFoundError(built)
        shutil.copy2(built, output)
    return output


def build_cover_letter() -> Path:
    document = Document()
    common.configure_document(document, line_numbers=False)
    properties = document.core_properties
    properties.title = f"Cover letter: {TITLE}"
    properties.author = AUTHOR
    properties.subject = "Submission to PhytoFrontiers"

    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for line in (
        date.today().strftime("%B %-d, %Y"),
        "Natalia Peres, Editor-in-Chief",
        "PhytoFrontiers",
        "The American Phytopathological Society",
    ):
        document.add_paragraph(line)
    document.add_paragraph("Dear Dr. Peres,")

    sentences = [
        "Please consider my manuscript, “Cultivar-dependent transcriptional responses of "
        "lychee to *Peronophythora litchii*: a registered genome-wide analysis,” as a "
        "Research article in PhytoFrontiers.",
        "The study uses a prospectively registered genome-wide workflow, independent "
        "robustness checks, and prespecified cross-context evaluation to distinguish "
        "cultivar-dependent infection responses from candidates selected within a single "
        "small RNA-sequencing cohort.",
        "Its transparent reporting of supported, null, and contradictory outcomes fits the "
        "journal's emphasis on rigorous data science relevant to plant health and on sound "
        "research regardless of perceived impact.",
        f"All underlying sequencing data are publicly accessioned, and the complete results, "
        f"figure source data, protocol, and analysis code are archived at {ZENODO_DOI}.",
        "The manuscript is not under consideration elsewhere and has not been published as "
        "a peer-reviewed article; I am the sole author, received no external funding, declare "
        "no conflicts of interest, and confirm that the reanalysis of public data required no "
        "human- or animal-research approval.",
    ]
    paragraph = document.add_paragraph()
    common.add_inline(paragraph, " ".join(sentences))

    for line in ("Sincerely,", AUTHOR, AFFILIATION, EMAIL, f"ORCID iD: {ORCID}"):
        document.add_paragraph(line)

    output = OUTPUT_DIR / "PhytoFrontiers_cover_letter.docx"
    document.save(output)
    return output


def fit_image(image: Image.Image, maximum_width: int, maximum_height: int) -> Image.Image:
    fitted = image.copy()
    fitted.thumbnail((maximum_width, maximum_height), Image.Resampling.LANCZOS)
    if fitted.mode != "RGB":
        background = Image.new("RGB", fitted.size, "white")
        if "A" in fitted.getbands():
            background.paste(fitted, mask=fitted.getchannel("A"))
        else:
            background.paste(fitted)
        fitted = background
    return fitted


def wrapped_lines(draw: ImageDraw.ImageDraw, text: str, font, maximum_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        candidate = " ".join(current + [word])
        if current and draw.textbbox((0, 0), candidate, font=font)[2] > maximum_width:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return lines


def add_wrapped_text(draw, text, font, x, y, width, line_height, fill="black") -> int:
    for line in wrapped_lines(draw, text, font, width):
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height
    return y


def build_extras_pdf() -> Path:
    page_width, page_height = 2550, 3300
    margin = 225
    title_font = ImageFont.truetype(
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf", 54
    )
    body_font = ImageFont.truetype(
        "/System/Library/Fonts/Supplemental/Arial.ttf", 42
    )
    pages: list[Image.Image] = []

    cover = Image.new("RGB", (page_width, page_height), "white")
    draw = ImageDraw.Draw(cover)
    y = margin
    draw.text((margin, y), "PhytoFrontiers e-Xtra", font=title_font, fill="black")
    y += 110
    y = add_wrapped_text(draw, TITLE, title_font, margin, y, page_width - 2 * margin, 68)
    y += 100
    y = add_wrapped_text(
        draw,
        "This single supplemental PDF contains Figures S1 to S3 with captions. Complete "
        "machine-readable Tables S1 to S18, figure source data, the registered protocol, "
        "and analysis code are openly archived at:",
        body_font,
        margin,
        y,
        page_width - 2 * margin,
        54,
    )
    y += 45
    draw.text((margin, y), ZENODO_DOI, font=body_font, fill="black")
    y += 125
    for label, _, caption in SUPPLEMENTARY_FIGURES:
        y = add_wrapped_text(
            draw,
            f"{label}. {caption}",
            body_font,
            margin,
            y,
            page_width - 2 * margin,
            54,
        )
        y += 45
    pages.append(cover)

    for label, path, caption in SUPPLEMENTARY_FIGURES:
        if not path.is_file():
            raise FileNotFoundError(path)
        page = Image.new("RGB", (page_width, page_height), "white")
        draw = ImageDraw.Draw(page)
        draw.text((margin, margin), label, font=title_font, fill="black")
        full_caption = f"{label}. {caption}"
        lines = wrapped_lines(draw, full_caption, body_font, page_width - 2 * margin)
        caption_height = 54 * len(lines)
        image_top = margin + 95
        image_height = page_height - image_top - margin - caption_height - 80
        with Image.open(path) as source:
            fitted = fit_image(source, page_width - 2 * margin, image_height)
        image_x = (page_width - fitted.width) // 2
        page.paste(fitted, (image_x, image_top))
        caption_y = image_top + fitted.height + 55
        for line in lines:
            draw.text((margin, caption_y), line, font=body_font, fill="black")
            caption_y += 54
        pages.append(page)

    output = OUTPUT_DIR / "supporting_information/PhytoFrontiers_eXtra_figures.pdf"
    output.parent.mkdir(parents=True, exist_ok=True)
    pages[0].save(
        output,
        "PDF",
        resolution=300.0,
        save_all=True,
        append_images=pages[1:],
        quality=95,
    )
    return output


def prepare_production_figure_files() -> list[Path]:
    figure_dir = OUTPUT_DIR / "figures"
    figure_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for _, source_name, target_name in FIGURES:
        source_pdf = MANUSCRIPT_DIR / "figures" / source_name
        if not source_pdf.is_file() or source_pdf.read_bytes()[:4] != b"%PDF":
            raise ValueError(f"Missing or invalid vector PDF: {source_pdf}")
        target_pdf = figure_dir / target_name
        shutil.copy2(source_pdf, target_pdf)
        outputs.append(target_pdf)

        source_png = source_pdf.with_suffix(".png")
        if not source_png.is_file():
            raise FileNotFoundError(source_png)
        target_tiff = target_pdf.with_suffix(".tif")
        with Image.open(source_png) as source_image:
            if source_image.mode == "RGB":
                production_image = source_image.copy()
            else:
                production_image = Image.new("RGB", source_image.size, "white")
                if "A" in source_image.getbands():
                    production_image.paste(source_image, mask=source_image.getchannel("A"))
                else:
                    production_image.paste(source_image)
            production_image.save(
                target_tiff,
                format="TIFF",
                compression="tiff_lzw",
                dpi=(300, 300),
            )
        with Image.open(target_tiff) as check:
            if check.format != "TIFF" or check.width < 2100:
                raise ValueError(
                    f"Production TIFF failed validation: {target_tiff} ({check.size})"
                )
        outputs.append(target_tiff)
    return outputs


def write_figure_captions() -> Path:
    legends = common.extract_figure_legends(SOURCE_MD.read_text(encoding="utf-8").splitlines())
    output = OUTPUT_DIR / "FIGURE_CAPTIONS.txt"
    blocks = [aps_citations(legends[number]) for number in range(1, 7)]
    output.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
    return output


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_manifest(paths: list[Path]) -> Path:
    output = OUTPUT_DIR / "UPLOAD_FILE_MANIFEST_SHA256.tsv"
    lines = ["sha256\tbytes\tfile"]
    for path in sorted(paths, key=lambda item: str(item.relative_to(OUTPUT_DIR))):
        lines.append(f"{sha256(path)}\t{path.stat().st_size}\t{path.relative_to(OUTPUT_DIR)}")
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return output


def write_validation_report(upload_paths: list[Path], retained_paths: list[Path]) -> Path:
    abstract_words = len(re.findall(r"\b[\w'-]+\b", common.plain_markdown(ABSTRACT)))
    report = [
        "PhytoFrontiers package validation",
        f"Build date: {date.today().isoformat()}",
        f"Title characters: {len(TITLE)} (limit: 150)",
        f"Abstract words: {abstract_words} (limit: 250)",
        "Initial Main Document: single PDF with 12-point Arial, double spacing, line numbers, and page headers",
        "Main sections: Introduction; Materials and Methods; Results; Discussion",
        "Main tables: 4 tables embedded after Literature Cited",
        "Main figures: 6 figures embedded after the tables",
        "Figure captions: each placed on the same page as its figure",
        "Editable backup: Word manuscript with embedded tables, figures, and captions",
        "Production figures retained: six vector PDFs and six 300-dpi LZW TIFFs",
        "Supplement: 1 four-page e-Xtra PDF (cover plus Figures S1-S3)",
        f"Literature Cited entries: {len(LITERATURE_CITED)} (author-year; alphabetized)",
        "Database accessions: stated in text rather than Literature Cited",
        "AI-use disclosure: present in Acknowledgments",
        "Cover-letter significance paragraph: 5 sentences",
        "Recommended initial-upload files:",
    ]
    report.extend(
        f"- {path.relative_to(OUTPUT_DIR)} ({path.stat().st_size} bytes)"
        for path in upload_paths
    )
    report.append("Retained revision/production aids (do not upload with the streamlined PDF):")
    report.extend(
        f"- {path.relative_to(OUTPUT_DIR)} ({path.stat().st_size} bytes)"
        for path in retained_paths
    )
    output = OUTPUT_DIR / "VALIDATION_REPORT.txt"
    output.write_text("\n".join(report) + "\n", encoding="utf-8")
    return output


def main() -> None:
    manuscript = build_manuscript()
    initial_pdf = build_initial_submission_pdf()
    cover_letter = build_cover_letter()
    extra = build_extras_pdf()
    figures = prepare_production_figure_files()
    captions = write_figure_captions()
    upload_files = [initial_pdf, cover_letter, extra]
    retained_files = [manuscript, captions, *figures]
    write_validation_report(upload_files, retained_files)
    write_manifest(upload_files)
    print(f"Built PhytoFrontiers package in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
