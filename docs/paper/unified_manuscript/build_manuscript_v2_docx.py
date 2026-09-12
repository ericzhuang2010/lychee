#!/usr/bin/env python3
"""Create an editable Word counterpart of a unified-manuscript PDF from LaTeX."""

from __future__ import annotations

import hashlib
import re
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


HERE = Path(__file__).resolve().parent
PROJECT_ROOT = HERE.parents[2]
DOCUMENT_STEM = sys.argv[1] if len(sys.argv) > 1 else "manuscript_v2"
if not re.fullmatch(r"manuscript(?:_v2)?", DOCUMENT_STEM):
    raise ValueError("Document stem must be 'manuscript' or 'manuscript_v2'")
SOURCE = HERE / f"{DOCUMENT_STEM}.tex"
SOURCE_PDF = HERE / f"{DOCUMENT_STEM}.pdf"
OUTPUT = HERE / f"{DOCUMENT_STEM}.docx"

REFERENCE_LABELS = {
    "fig:design": "1",
    "tab:roles": "1",
    "fig:discovery": "2",
    "tab:legacy": "2",
    "fig:robust": "3",
    "tab:external": "3",
    "fig:pathways": "4",
    "fig:dtu": "5",
    "fig:ortho": "6",
    "tab:tierb": "4",
    "fig:supp-counts": "S1",
    "fig:supp-power": "S2",
    "fig:supp-exploratory": "S3",
}

ACCENTS = {
    r"Saloj\"arvi": "Salojärvi",
    r"M\"older": "Mölder",
    r"Kopeck\'a": "Kopecká",
    r"Berkov\'a": "Berková",
    r"Brzobohat\'y": "Brzobohatý",
    r"\v{C}ern\'y": "Černý",
    r"D\'eustachio": "D’Eustachio",
    r"D'Eustachio": "D’Eustachio",
    r"D\'ehais": "Déhais",
}

MATH_REPLACEMENTS = {
    r"\mathrm{log2FC}": "log2FC",
    r"\log_2": "log₂",
    r"\beta": "β",
    r"\rho": "ρ",
    r"\times": "×",
    r"\ge": "≥",
    r"\le": "≤",
    r"\sim": "~",
    r"\q": "q",
    r"\ldots": "…",
}

SUPERSCRIPT = str.maketrans("-+0123456789", "⁻⁺⁰¹²³⁴⁵⁶⁷⁸⁹")


@dataclass
class Segment:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    url: str | None = None


def balanced_argument(value: str, brace_index: int) -> tuple[str, int]:
    if brace_index >= len(value) or value[brace_index] != "{":
        raise ValueError(f"Expected braced argument in {value!r}")
    depth = 0
    for index in range(brace_index, len(value)):
        if value[index] == "{":
            depth += 1
        elif value[index] == "}":
            depth -= 1
            if depth == 0:
                return value[brace_index + 1 : index], index + 1
    raise ValueError(f"Unbalanced braces in {value!r}")


def command_argument(value: str, command: str) -> str:
    start = value.index(command) + len(command)
    while start < len(value) and value[start].isspace():
        start += 1
    result, _ = balanced_argument(value, start)
    return result


def normalize_math(value: str) -> str:
    for source, replacement in MATH_REPLACEMENTS.items():
        value = value.replace(source, replacement)
    value = re.sub(
        r"10\^\{([+-]?\d+)\}",
        lambda match: "10" + match.group(1).translate(SUPERSCRIPT),
        value,
    )
    value = re.sub(r"\^\{([+-]?\d+)\}", lambda match: match.group(1).translate(SUPERSCRIPT), value)
    value = value.replace(r"\%", "%").replace(r"\_", "_")
    value = value.replace("{", "").replace("}", "")
    return value.strip()


def merge_segments(segments: list[Segment]) -> list[Segment]:
    merged: list[Segment] = []
    for segment in segments:
        if not segment.text:
            continue
        if (
            merged
            and merged[-1].bold == segment.bold
            and merged[-1].italic == segment.italic
            and merged[-1].code == segment.code
            and merged[-1].url == segment.url
        ):
            merged[-1].text += segment.text
        else:
            merged.append(segment)
    return merged


def inline_segments(
    value: str,
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
) -> list[Segment]:
    for source, replacement in ACCENTS.items():
        value = value.replace(source, replacement)
    segments: list[Segment] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            segments.append(Segment("".join(buffer), bold=bold, italic=italic, code=code))
            buffer.clear()

    index = 0
    while index < len(value):
        if value[index] == "$":
            end = value.find("$", index + 1)
            if end < 0:
                buffer.append(value[index])
                index += 1
                continue
            flush()
            segments.append(
                Segment(normalize_math(value[index + 1 : end]), bold=bold, italic=True, code=code)
            )
            index = end + 1
            continue

        if value[index] != "\\":
            char = value[index]
            if value.startswith("--", index):
                buffer.append("–")
                index += 2
            elif char == "~":
                buffer.append(" ")
                index += 1
            else:
                buffer.append(char)
                index += 1
            continue

        if index + 1 < len(value) and value[index + 1] in "%_&#":
            buffer.append(value[index + 1])
            index += 2
            continue
        if value.startswith(r"\ ", index):
            buffer.append(" ")
            index += 2
            continue

        command_match = re.match(r"\\([A-Za-z]+)", value[index:])
        if not command_match:
            buffer.append(value[index])
            index += 1
            continue
        command = command_match.group(1)
        after = index + len(command_match.group(0))

        if command in {"textit", "emph", "textbf", "texttt", "mathrm"} and after < len(value) and value[after] == "{":
            argument, next_index = balanced_argument(value, after)
            flush()
            segments.extend(
                inline_segments(
                    argument,
                    bold=bold or command == "textbf",
                    italic=italic or command in {"textit", "emph"},
                    code=code or command == "texttt",
                )
            )
            index = next_index
            continue

        if command == "href" and after < len(value) and value[after] == "{":
            url, after_url = balanced_argument(value, after)
            label, next_index = balanced_argument(value, after_url)
            flush()
            label_segments = inline_segments(label, bold=bold, italic=italic, code=code)
            for segment in label_segments:
                segment.url = url
            segments.extend(label_segments)
            index = next_index
            continue

        if command == "url" and after < len(value) and value[after] == "{":
            url, next_index = balanced_argument(value, after)
            flush()
            segments.append(Segment(url, bold=bold, italic=italic, code=code, url=url))
            index = next_index
            continue

        if command == "ref" and after < len(value) and value[after] == "{":
            label, next_index = balanced_argument(value, after)
            buffer.append(REFERENCE_LABELS.get(label, label))
            index = next_index
            continue

        if command == "Pl":
            flush()
            segments.append(Segment("P. litchii", bold=bold, italic=True, code=code))
            index = after + 2 if value[after : after + 2] == "{}" else after
            continue

        if command == "q":
            flush()
            segments.append(Segment("q", bold=bold, italic=True, code=code))
            index = after + 2 if value[after : after + 2] == "{}" else after
            continue

        replacements = {
            "ldots": "…",
            "times": "×",
            "ge": "≥",
            "le": "≤",
            "sim": "~",
            "rho": "ρ",
            "beta": "β",
            "noindent": "",
        }
        if command in replacements:
            buffer.append(replacements[command])
            index = after
            continue

        # Retain unknown narrative commands visibly so validation can detect them.
        buffer.append("\\" + command)
        index = after

    flush()
    return merge_segments(segments)


def add_hyperlink(paragraph, segment: Segment) -> None:
    relationship_id = paragraph.part.relate_to(
        segment.url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    if segment.italic:
        properties.append(OxmlElement("w:i"))
    if segment.bold:
        properties.append(OxmlElement("w:b"))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = segment.text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_segments(paragraph, segments: list[Segment], *, size: float | None = None) -> None:
    for segment in segments:
        if segment.url:
            add_hyperlink(paragraph, segment)
            continue
        run = paragraph.add_run(segment.text)
        run.bold = segment.bold
        run.italic = segment.italic
        run.font.name = "Courier New" if segment.code else "Times New Roman"
        if size is not None:
            run.font.size = Pt(size)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    add_page_field(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.first_line_indent = Inches(0.25)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        properties = style.element.get_or_add_pPr()
        borders = properties.find(qn("w:pBdr"))
        if borders is not None:
            properties.remove(borders)

    caption = document.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_together = True


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_table(
    document: Document,
    block: str,
    caption_text: str,
    page_break: bool,
    table_number: int,
) -> None:
    start = block.index(r"\begin{tabular}")
    start = block.index("\n", start) + 1
    end = block.index(r"\end{tabular}", start)
    rows: list[list[str]] = []
    for raw_line in block[start:end].splitlines():
        line = raw_line.strip()
        if not line or line in {r"\toprule", r"\midrule", r"\bottomrule"}:
            continue
        if line.endswith(r"\\"):
            line = line[:-2].strip()
        rows.append([cell.strip() for cell in line.split("&")])
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("Malformed LaTeX table")

    caption = document.add_paragraph(style="Caption")
    caption.paragraph_format.page_break_before = page_break
    caption.paragraph_format.keep_with_next = True
    caption.add_run(f"Table {table_number}. ").bold = True
    add_segments(caption, inline_segments(caption_text), size=9)

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    repeat_header(table.rows[0])
    font_size = 7 if len(rows[0]) >= 6 else 7.5
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, "D9E2F3")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            add_segments(paragraph, inline_segments(value), size=font_size)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def resolve_figure(filename: str) -> Path:
    candidates = [HERE / "figures" / filename, PROJECT_ROOT / "results" / "figures" / filename]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Could not resolve figure {filename}")


def rasterize_pdf(source: Path, destination: Path) -> None:
    result = subprocess.run(
        ["sips", "-s", "format", "png", "-Z", "2600", str(source), "--out", str(destination)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(f"Could not rasterize {source}:\n{result.stdout}")


def add_figure(
    document: Document,
    filename: str,
    caption_text: str,
    temp_directory: Path,
    figure_label: str,
) -> None:
    source = resolve_figure(filename)
    png = temp_directory / (source.stem + ".png")
    rasterize_pdf(source, png)
    with Image.open(png) as image:
        width_px, height_px = image.size
    max_width = 6.75
    max_height = 7.0
    width = min(max_width, max_height * width_px / height_px)
    height = width * height_px / width_px

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(png), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", re.sub(r"\\[A-Za-z]+|[{}$]", "", caption_text)[:1000])

    caption = document.add_paragraph(style="Caption")
    caption.add_run(f"Figure {figure_label}. ").bold = True
    add_segments(caption, inline_segments(caption_text), size=9)


def add_body_paragraph(document: Document, value: str, *, page_break: bool = False) -> None:
    value = value.strip()
    if not value:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.page_break_before = page_break
    add_segments(paragraph, inline_segments(value))


def add_title(document: Document, source_text: str) -> None:
    title_text = command_argument(source_text, r"\title")
    title_text = title_text.replace(r"\bfseries ", "").replace(r"\bfseries", "")
    author_text = command_argument(source_text, r"\author")
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Inches(0)
    add_segments(title, inline_segments(title_text))
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.first_line_indent = Inches(0)
    author.add_run(author_text).bold = True


def build() -> None:
    source_text = SOURCE.read_text(encoding="utf-8")
    document = Document()
    configure_document(document)
    document.core_properties.title = re.sub(
        r"\\textit\{([^}]+)\}|\\bfseries\s*", r"\1", command_argument(source_text, r"\title")
    )
    document.core_properties.author = command_argument(source_text, r"\author")
    document.core_properties.subject = f"Editable Word version of {SOURCE_PDF.name}"
    add_title(document, source_text)

    body = source_text.split(r"\begin{document}", 1)[1].split(r"\end{document}", 1)[0]
    lines = body.splitlines()
    index = 0
    figure_count = 0
    table_count = 0
    reference_count = 0
    in_abstract = False
    in_references = False
    page_break_next = False
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal page_break_next
        if paragraph_buffer:
            add_body_paragraph(document, " ".join(paragraph_buffer), page_break=page_break_next)
            paragraph_buffer.clear()
            page_break_next = False

    with TemporaryDirectory(prefix="manuscript_v2_docx_") as temp_name:
        temp_directory = Path(temp_name)
        while index < len(lines):
            raw = lines[index]
            line = raw.strip()

            if not line:
                flush_paragraph()
                index += 1
                continue
            if line in {r"\maketitle", r"\noindent"} or line.startswith(r"\vspace"):
                index += 1
                continue

            if line == r"\begin{abstract}":
                flush_paragraph()
                heading = document.add_paragraph("Abstract", style="Heading 1")
                heading.paragraph_format.space_before = Pt(14)
                in_abstract = True
                index += 1
                continue
            if line == r"\end{abstract}":
                flush_paragraph()
                in_abstract = False
                index += 1
                continue

            section = re.match(r"\\section(\*)?\{(.+)\}$", line)
            if section:
                flush_paragraph()
                heading = document.add_paragraph(style="Heading 1")
                heading.paragraph_format.page_break_before = page_break_next
                page_break_next = False
                add_segments(heading, inline_segments(section.group(2)))
                in_references = section.group(2) == "References"
                index += 1
                continue

            subsection = re.match(r"\\subsection\{(.+)\}$", line)
            if subsection:
                flush_paragraph()
                heading = document.add_paragraph(style="Heading 2")
                heading.paragraph_format.page_break_before = page_break_next
                page_break_next = False
                add_segments(heading, inline_segments(subsection.group(1)))
                index += 1
                continue

            if line.startswith(r"\begin{figure}"):
                flush_paragraph()
                end = index
                while end < len(lines) and lines[end].strip() != r"\end{figure}":
                    end += 1
                block = "\n".join(lines[index : end + 1])
                # includegraphics has an optional width before the filename.
                filename = re.search(r"\\includegraphics(?:\[[^]]+\])?\{([^}]+)\}", block).group(1)
                caption_text = command_argument(block, r"\caption")
                figure_label = str(figure_count + 1) if figure_count < 6 else f"S{figure_count - 5}"
                add_figure(document, filename, caption_text, temp_directory, figure_label)
                figure_count += 1
                index = end + 1
                continue

            if line.startswith(r"\begin{table}"):
                flush_paragraph()
                end = index
                while end < len(lines) and lines[end].strip() != r"\end{table}":
                    end += 1
                block = "\n".join(lines[index : end + 1])
                caption_text = command_argument(block, r"\caption")
                add_table(document, block, caption_text, "[p]" in line, table_count + 1)
                table_count += 1
                index = end + 1
                continue

            if line == r"\clearpage":
                flush_paragraph()
                page_break_next = True
                index += 1
                continue

            if line in {
                r"\begin{enumerate}[itemsep=1pt,leftmargin=1.6em]",
                r"\end{enumerate}",
            }:
                flush_paragraph()
                index += 1
                continue

            if in_references and line.startswith(r"\item "):
                flush_paragraph()
                reference_count += 1
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.28)
                paragraph.paragraph_format.first_line_indent = Inches(-0.28)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.add_run(f"{reference_count}. ").bold = False
                add_segments(paragraph, inline_segments(line[len(r"\item ") :]), size=9.5)
                index += 1
                continue

            if line.startswith(
                (
                    r"\centering",
                    r"\label",
                    r"\small",
                    r"\footnotesize",
                    r"\setlength",
                    r"\setcounter",
                    r"\renewcommand",
                )
            ):
                index += 1
                continue

            paragraph_buffer.append(line)
            index += 1

        flush_paragraph()
        document.save(OUTPUT)

    with zipfile.ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise ValueError("Generated DOCX is corrupt")
        xml = archive.read("word/document.xml")
    reopened = Document(OUTPUT)
    combined = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    combined += "\n" + "\n".join(
        cell.text for table in reopened.tables for row in table.rows for cell in row.cells
    )
    if len(reopened.inline_shapes) != 9 or figure_count != 9:
        raise ValueError(f"Expected 9 embedded figures, found {len(reopened.inline_shapes)}")
    if len(reopened.tables) != 4 or table_count != 4:
        raise ValueError(f"Expected 4 editable tables, found {len(reopened.tables)}")
    if reference_count != 41:
        raise ValueError(f"Expected 41 references, found {reference_count}")
    for required in ("Abstract", "Introduction", "Results", "Discussion", "Conclusions", "Materials and Methods", "References"):
        if required not in combined:
            raise ValueError(f"Missing required section: {required}")
    leftovers = sorted(set(re.findall(r"\\\S+", combined)))
    if leftovers:
        raise ValueError(f"Unconverted LaTeX commands remain: {leftovers}")
    if xml.count(b"<w:drawing") != 9:
        raise ValueError("Unexpected Word drawing count")

    digest = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    print(f"Created {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size} bytes")
    print(f"Figures: {len(reopened.inline_shapes)}; tables: {len(reopened.tables)}; references: {reference_count}")
    print(f"SHA-256: {digest}")


if __name__ == "__main__":
    if not SOURCE.exists() or not SOURCE_PDF.exists():
        raise FileNotFoundError(f"{SOURCE.name} and {SOURCE_PDF.name} are required")
    build()
