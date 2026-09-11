#!/usr/bin/env python3
"""Create editable Word versions of the Discover Plants submission files.

The converter is deliberately scoped to the LaTeX constructs used by this
manuscript.  Building from the journal-adapted TeX source keeps the Word and
LaTeX submissions synchronized while allowing proper Word tables, captions,
headings, hyperlinks, and embedded high-resolution figures.
"""

from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


@dataclass(frozen=True)
class Segment:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    math: bool = False
    url: str | None = None


def extract_braced(text: str, opening_brace: int) -> tuple[str, int]:
    """Return balanced-brace contents and the index after the closing brace."""
    if opening_brace >= len(text) or text[opening_brace] != "{":
        raise ValueError("Expected an opening brace")
    depth = 0
    for index in range(opening_brace, len(text)):
        if text[index] == "{" and (index == 0 or text[index - 1] != "\\"):
            depth += 1
        elif text[index] == "}" and (index == 0 or text[index - 1] != "\\"):
            depth -= 1
            if depth == 0:
                return text[opening_brace + 1 : index], index + 1
    raise ValueError("Unbalanced braces in LaTeX source")


def command_argument(text: str, command: str) -> str:
    match = re.search(rf"\\{re.escape(command)}\s*\{{", text)
    if not match:
        raise ValueError(f"Command \\{command} not found")
    content, _ = extract_braced(text, match.end() - 1)
    return content


def clean_math(text: str) -> str:
    superscripts = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")
    subscripts = str.maketrans("0123456789+-", "₀₁₂₃₄₅₆₇₈₉₊₋")

    text = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\text\{([^{}]+)\}", r"\1", text)
    text = text.replace(r"\beta", "β").replace(r"\rho", "ρ")
    text = text.replace(r"\times", "×").replace(r"\ge", "≥")
    text = text.replace(r"\le", "≤").replace(r"\sim", "∼")
    text = text.replace(r"\ldots", "…").replace(r"\q", "q")
    text = re.sub(
        r"\^\{([0-9+\-]+)\}", lambda m: m.group(1).translate(superscripts), text
    )
    text = re.sub(
        r"_\{?([0-9+\-]+)\}?", lambda m: m.group(1).translate(subscripts), text
    )
    text = text.replace(r"\log", "log")
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


def _merge_segments(segments: list[Segment]) -> list[Segment]:
    merged: list[Segment] = []
    for segment in segments:
        if not segment.text:
            continue
        formatting = (
            segment.bold,
            segment.italic,
            segment.code,
            segment.math,
            segment.url,
        )
        previous_formatting = (
            merged[-1].bold,
            merged[-1].italic,
            merged[-1].code,
            merged[-1].math,
            merged[-1].url,
        ) if merged else None
        if merged and previous_formatting == formatting:
            previous = merged[-1]
            merged[-1] = Segment(
                previous.text + segment.text,
                previous.bold,
                previous.italic,
                previous.code,
                previous.math,
                previous.url,
            )
        else:
            merged.append(segment)
    return merged


def tex_segments(
    text: str,
    labels: dict[str, str],
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
) -> list[Segment]:
    """Convert the manuscript's inline LaTeX to styled Word text segments."""
    segments: list[Segment] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            value = "".join(buffer)
            value = value.replace("---", "—").replace("--", "–")
            value = value.replace("``", "“").replace("''", "”")
            segments.append(Segment(value, bold, italic, code))
            buffer.clear()

    index = 0
    while index < len(text):
        char = text[index]
        if char == "$":
            end = text.find("$", index + 1)
            if end < 0:
                buffer.append(char)
                index += 1
                continue
            flush()
            segments.append(Segment(clean_math(text[index + 1 : end]), bold, italic, math=True))
            index = end + 1
            continue

        if char == "\\":
            escaped = {"%": "%", "&": "&", "_": "_", "#": "#", "$": "$"}
            if index + 1 < len(text) and text[index + 1] in escaped:
                buffer.append(escaped[text[index + 1]])
                index += 2
                continue

            match = re.match(r"\\([A-Za-z]+\*?)", text[index:])
            if not match:
                index += 1
                continue
            command = match.group(1)
            command_end = index + match.end()
            if command_end < len(text) and text[command_end] == "{":
                argument, next_index = extract_braced(text, command_end)
                if command in {"textit", "emph"}:
                    flush()
                    segments.extend(
                        tex_segments(argument, labels, bold=bold, italic=True, code=code)
                    )
                elif command == "textbf":
                    flush()
                    segments.extend(
                        tex_segments(argument, labels, bold=True, italic=italic, code=code)
                    )
                elif command == "texttt":
                    flush()
                    segments.extend(
                        tex_segments(argument, labels, bold=bold, italic=italic, code=True)
                    )
                elif command == "url":
                    flush()
                    segments.append(Segment(argument, bold, italic, code, url=argument))
                elif command == "ref":
                    flush()
                    segments.append(Segment(labels.get(argument, argument), bold, italic, code))
                elif command == "label":
                    pass
                elif command == "Pl":
                    flush()
                    segments.append(Segment("P. litchii", bold, True, code))
                elif command == "q":
                    flush()
                    segments.append(Segment("q", bold, True, code))
                else:
                    # Preserve the human-readable argument of an unfamiliar
                    # formatting command instead of leaking LaTeX into the DOCX.
                    flush()
                    segments.extend(
                        tex_segments(argument, labels, bold=bold, italic=italic, code=code)
                    )
                index = next_index
                continue

            if command in {"Pl", "q"}:
                flush()
                replacement = "P. litchii" if command == "Pl" else "q"
                segments.append(Segment(replacement, bold, True, code))
                index = command_end
                continue
            if command == "ldots":
                buffer.append("…")
            elif command in {"backmatter", "centering", "small", "footnotesize"}:
                pass
            else:
                # Inline symbols outside math are rare in this source.  Keep a
                # readable form for any that occur.
                symbols = {"times": "×", "ge": "≥", "le": "≤", "sim": "∼"}
                buffer.append(symbols.get(command, command))
            index = command_end
            if text[index : index + 2] == "{}":
                index += 2
            continue

        if char == "~":
            buffer.append("\u00a0")
        elif char in "{}":
            pass
        elif char == "\n":
            buffer.append(" ")
        else:
            buffer.append(char)
        index += 1

    flush()
    normalized: list[Segment] = []
    for segment in segments:
        value = re.sub(r"[ \t\r\f\v]+", " ", segment.text)
        normalized.append(
            Segment(value, segment.bold, segment.italic, segment.code, segment.math, segment.url)
        )
    return _merge_segments(normalized)


def add_hyperlink(paragraph, segment: Segment) -> None:
    relationship_id = paragraph.part.relate_to(
        segment.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = segment.text
    run.extend([properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_segments(paragraph, segments: list[Segment]) -> None:
    for segment in segments:
        if segment.url:
            add_hyperlink(paragraph, segment)
            continue
        run = paragraph.add_run(segment.text)
        run.bold = segment.bold
        run.italic = segment.italic
        if segment.code:
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif segment.math:
            run.font.name = "Cambria Math"


def add_tex_paragraph(document: Document, text: str, labels: dict[str, str], style=None):
    paragraph = document.add_paragraph(style=style)
    add_segments(paragraph, tex_segments(text.strip(), labels))
    return paragraph


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, display, end])


def configure_document(document: Document, *, line_numbers: bool) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        # Remove the blue bottom rule carried by python-docx's default Title
        # style so the submission has neutral journal formatting.
        paragraph_properties = style.element.get_or_add_pPr()
        borders = paragraph_properties.find(qn("w:pBdr"))
        if borders is not None:
            paragraph_properties.remove(borders)

    caption = document.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.text = "Discover Plants — Research"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Times New Roman"
    header.runs[0].font.size = Pt(9)
    add_page_field(section.footer.paragraphs[0])

    if line_numbers:
        line_numbering = OxmlElement("w:lnNumType")
        line_numbering.set(qn("w:countBy"), "5")
        line_numbering.set(qn("w:distance"), "360")
        line_numbering.set(qn("w:restart"), "continuous")
        section._sectPr.append(line_numbering)


def labels_from_tex(text: str) -> tuple[dict[str, str], dict[str, int], dict[str, int]]:
    labels: dict[str, str] = {}
    figure_numbers: dict[str, int] = {}
    table_numbers: dict[str, int] = {}
    for environment, number_map in (("figure", figure_numbers), ("table", table_numbers)):
        pattern = re.compile(
            rf"\\begin\{{{environment}\}}.*?\\end\{{{environment}\}}", re.DOTALL
        )
        for number, match in enumerate(pattern.finditer(text), start=1):
            label_match = re.search(r"\\label\{([^{}]+)\}", match.group(0))
            if label_match:
                label = label_match.group(1)
                labels[label] = str(number)
                number_map[label] = number
    return labels, figure_numbers, table_numbers


def split_paragraphs(text: str) -> list[str]:
    text = re.sub(r"\\label\{[^{}]+\}", "", text)
    text = re.sub(r"\\(?:backmatter|centering|small|footnotesize)\b", "", text)
    return [re.sub(r"\s+", " ", item).strip() for item in re.split(r"\n\s*\n", text) if item.strip()]


def extract_environment(text: str, start: int, environment: str) -> tuple[str, int]:
    end_marker = rf"\end{{{environment}}}"
    end = text.find(end_marker, start)
    if end < 0:
        raise ValueError(f"Unclosed {environment} environment")
    end += len(end_marker)
    return text[start:end], end


def convert_pdf_to_png(source: Path, destination: Path) -> None:
    result = subprocess.run(
        ["sips", "-s", "format", "png", "-Z", "2600", str(source), "--out", str(destination)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(f"Could not rasterize {source}:\n{result.stdout}")


def add_picture(
    document: Document,
    png: Path,
    caption_plain: str,
    *,
    supplement: bool = False,
    page_break_before: bool = False,
) -> None:
    with Image.open(png) as image:
        width_px, height_px = image.size
    max_width = 6.55
    max_height = 7.3 if supplement else 6.7
    width = min(max_width, max_height * width_px / height_px)
    height = width * height_px / width_px
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    shape = paragraph.add_run().add_picture(
        str(png), width=Inches(width), height=Inches(height)
    )
    shape._inline.docPr.set("descr", caption_plain[:1000])


def add_caption(document: Document, prefix: str, caption_tex: str, labels: dict[str, str]) -> None:
    paragraph = document.add_paragraph(style="Caption")
    run = paragraph.add_run(prefix)
    run.bold = True
    add_segments(paragraph, tex_segments(caption_tex, labels))


def table_rows(block: str) -> list[list[str]]:
    begin = block.find(r"\begin{tabular}")
    if begin < 0:
        raise ValueError("Table has no tabular environment")
    body_start = block.find("\n", begin)
    body_end = block.find(r"\end{tabular}", body_start)
    body = block[body_start:body_end]
    body = re.sub(r"\\(?:toprule|midrule|bottomrule)\s*", "", body)
    rows: list[list[str]] = []
    for raw_row in re.split(r"\\\\", body):
        raw_row = re.sub(r"\\(?:small|footnotesize)\b", "", raw_row).strip()
        if not raw_row:
            continue
        cells = [cell.strip() for cell in raw_row.split("&")]
        if len(cells) > 1:
            rows.append(cells)
    return rows


def add_table(document: Document, block: str, number: int, labels: dict[str, str]) -> None:
    caption = command_argument(block, "caption")
    add_caption(document, f"Table {number}. ", caption, labels)
    rows = table_rows(block)
    if not rows:
        raise ValueError(f"Table {number} has no rows")
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError(f"Table {number} has inconsistent column counts")

    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            add_segments(paragraph, tex_segments(value, labels))
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8 if column_count >= 6 else 8.5)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                shade_cell(cell, "D9E2F3")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_reference_list(document: Document, block: str, labels: dict[str, str]) -> None:
    inner = re.sub(r"^.*?\n", "", block, count=1, flags=re.DOTALL)
    inner = re.sub(r"\\end\{enumerate\}\s*$", "", inner, flags=re.DOTALL)
    items = re.split(r"\\item\s+", inner)[1:]
    for number, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.32)
        paragraph.paragraph_format.first_line_indent = Inches(-0.32)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(4)
        number_run = paragraph.add_run(f"[{number}] ")
        number_run.bold = True
        add_segments(paragraph, tex_segments(re.sub(r"\s+", " ", item).strip(), labels))
        for run in paragraph.runs:
            run.font.size = Pt(10)


def plain_tex(text: str, labels: dict[str, str]) -> str:
    return "".join(segment.text for segment in tex_segments(text, labels)).strip()


def render_manuscript_body(
    document: Document,
    body: str,
    labels: dict[str, str],
    figure_sources: dict[str, Path],
    image_directory: Path,
) -> None:
    event_pattern = re.compile(
        r"\\(?P<heading>section\*?|subsection\*?|bmhead)\{(?P<title>[^{}]+)\}"
        r"|\\backmatter"
        r"|\\begin\{(?P<environment>figure|table|enumerate)\}(?:\[[^]]*\])?"
    )
    position = 0
    section_number = 0
    subsection_number = 0
    figure_number = 0
    table_number = 0

    def add_text(chunk: str) -> None:
        for item in split_paragraphs(chunk):
            add_tex_paragraph(document, item, labels)

    while True:
        match = event_pattern.search(body, position)
        if not match:
            add_text(body[position:])
            break
        add_text(body[position : match.start()])

        heading = match.group("heading")
        environment = match.group("environment")
        if heading:
            title = plain_tex(match.group("title"), labels)
            if heading == "section":
                section_number += 1
                subsection_number = 0
                document.add_heading(f"{section_number} {title}", level=1)
            elif heading == "subsection":
                subsection_number += 1
                document.add_heading(f"{section_number}.{subsection_number} {title}", level=2)
            elif heading in {"section*", "bmhead"}:
                document.add_heading(title, level=1)
            else:
                document.add_heading(title, level=2)
            position = match.end()
            continue

        if match.group(0).startswith(r"\backmatter"):
            position = match.end()
            continue

        block, position = extract_environment(body, match.start(), environment)
        if environment == "figure":
            figure_number += 1
            filename_match = re.search(r"\\includegraphics(?:\[[^]]*\])?\{([^{}]+)\}", block)
            if not filename_match:
                raise ValueError(f"Figure {figure_number} has no image")
            filename = filename_match.group(1)
            caption = command_argument(block, "caption")
            png = image_directory / f"main_{figure_number}.png"
            convert_pdf_to_png(figure_sources[filename], png)
            add_picture(document, png, plain_tex(caption, labels))
            add_caption(document, f"Figure {figure_number}. ", caption, labels)
        elif environment == "table":
            table_number += 1
            add_table(document, block, table_number, labels)
        else:
            add_reference_list(document, block, labels)


def build_manuscript_docx(
    tex_path: Path,
    output_path: Path,
    figure_sources: dict[str, Path],
    title_tex: str,
    author: str,
    affiliation: str,
    email: str,
    orcid: str,
    image_directory: Path,
) -> None:
    text = tex_path.read_text(encoding="ascii")
    labels, _, _ = labels_from_tex(text)
    abstract = command_argument(text, "abstract")
    keywords = command_argument(text, "keywords")
    body_start = text.index(r"\maketitle") + len(r"\maketitle")
    body_end = text.rindex(r"\end{document}")
    body = text[body_start:body_end]

    document = Document()
    configure_document(document, line_numbers=True)
    document.core_properties.title = plain_tex(title_tex, labels)
    document.core_properties.subject = "Research article submitted to Discover Plants"
    document.core_properties.author = author
    document.core_properties.keywords = plain_tex(keywords, labels)
    document.core_properties.comments = "Editable Word submission generated from the journal-adapted LaTeX source."

    article_type = document.add_paragraph()
    article_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    article_type.add_run("Article type: Research").bold = True

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_segments(title, tex_segments(title_tex, labels))

    author_paragraph = document.add_paragraph()
    author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_paragraph.add_run(author).bold = True
    affiliation_paragraph = document.add_paragraph(
        f"{affiliation}\nCorresponding author: {email}\nORCID: {orcid}"
    )
    affiliation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for paragraph in (author_paragraph, affiliation_paragraph):
        paragraph.paragraph_format.line_spacing = 1.0

    document.add_heading("Abstract", level=1)
    abstract_paragraph = add_tex_paragraph(document, abstract, labels)
    abstract_paragraph.paragraph_format.line_spacing = 1.5
    keyword_paragraph = document.add_paragraph()
    keyword_paragraph.add_run("Keywords: ").bold = True
    add_segments(keyword_paragraph, tex_segments(keywords, labels))
    keyword_paragraph.paragraph_format.line_spacing = 1.5

    render_manuscript_body(document, body, labels, figure_sources, image_directory)
    document.save(output_path)


def build_supplement_docx(
    output_path: Path,
    supplementary_figures: list[tuple[Path, str, str]],
    title_tex: str,
    author: str,
    affiliation: str,
    zenodo_doi: str,
    image_directory: Path,
) -> None:
    labels: dict[str, str] = {}
    document = Document()
    configure_document(document, line_numbers=False)
    document.core_properties.title = "Supplementary Information 1"
    document.core_properties.subject = plain_tex(title_tex, labels)
    document.core_properties.author = author

    heading = document.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("Supplementary Information 1")
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_segments(title, tex_segments(title_tex, labels))
    for run in title.runs:
        run.bold = True
    byline = document.add_paragraph(f"{author}\n{affiliation}")
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.paragraph_format.line_spacing = 1.0

    document.add_heading("Contents", level=1)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(
        "This file contains Figures S1–S3 and their captions. Supplementary Data 2, "
        "supplied as Discover_Plants_Data_and_Code.zip, contains machine-readable "
        "Tables S1–S18, all figure source data, the registered protocol and amendment "
        "log, analysis code, workflows, configurations, tests, metadata, and "
        "environment specifications. The complete archive is also permanently "
        "available at "
    )
    add_hyperlink(paragraph, Segment(zenodo_doi, url=zenodo_doi))
    paragraph.add_run(".")
    for number, (_, _, caption_tex) in enumerate(supplementary_figures, start=1):
        item = document.add_paragraph(style="List Bullet")
        item.paragraph_format.line_spacing = 1.0
        short_caption = plain_tex(caption_tex, labels).split(".", maxsplit=1)[0]
        item.add_run(f"Figure S{number}: {short_caption}.")

    for number, (source, _, caption_tex) in enumerate(supplementary_figures, start=1):
        png = image_directory / f"supplement_{number}.png"
        convert_pdf_to_png(source, png)
        add_picture(
            document,
            png,
            plain_tex(caption_tex, labels),
            supplement=True,
            page_break_before=True,
        )
        add_caption(document, f"Figure S{number}. ", caption_tex, labels)

    document.save(output_path)


def validate_docx(path: Path, *, expected_images: int, expected_tables: int) -> dict[str, int]:
    document = Document(path)
    text_parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    joined = "\n".join(text_parts)
    if len(document.inline_shapes) != expected_images:
        raise ValueError(
            f"{path.name}: expected {expected_images} images, found {len(document.inline_shapes)}"
        )
    if len(document.tables) != expected_tables:
        raise ValueError(
            f"{path.name}: expected {expected_tables} tables, found {len(document.tables)}"
        )
    if "\\text" in joined or "\\begin" in joined or "\\ref" in joined:
        raise ValueError(f"{path.name}: unconverted LaTeX found")
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "images": len(document.inline_shapes),
        "characters": len(joined),
    }


def build_word_files(
    *,
    manuscript_tex: Path,
    manuscript_docx: Path,
    supplement_docx: Path,
    main_figures: list[tuple[Path, str]],
    supplementary_figures: list[tuple[Path, str, str]],
    title_tex: str,
    author: str,
    affiliation: str,
    email: str,
    orcid: str,
    zenodo_doi: str,
) -> dict[str, dict[str, int]]:
    with TemporaryDirectory(prefix="discover_plants_word_") as temp_name:
        image_directory = Path(temp_name)
        figure_sources = {filename: source for source, filename in main_figures}
        build_manuscript_docx(
            manuscript_tex,
            manuscript_docx,
            figure_sources,
            title_tex,
            author,
            affiliation,
            email,
            orcid,
            image_directory,
        )
        build_supplement_docx(
            supplement_docx,
            supplementary_figures,
            title_tex,
            author,
            affiliation,
            zenodo_doi,
            image_directory,
        )

    return {
        "manuscript": validate_docx(manuscript_docx, expected_images=6, expected_tables=4),
        "supplement": validate_docx(supplement_docx, expected_images=3, expected_tables=0),
    }
