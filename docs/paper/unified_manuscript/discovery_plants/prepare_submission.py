#!/usr/bin/env python3
"""Build the Discover Plants Word and LaTeX submission packages.

The recommended upload set uses editable DOCX files.  A flat, self-contained
LaTeX source archive and matching PDFs are retained as alternate formats and
for visual comparison.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUTPUT_DIR = Path(__file__).resolve().parent
MANUSCRIPT_DIR = OUTPUT_DIR.parent
PROJECT_ROOT = OUTPUT_DIR.parents[3]
SOURCE_TEX = MANUSCRIPT_DIR / "manuscript_v2.tex"
TEMPLATE_DIR = OUTPUT_DIR / "springer_template"
LATEX_DIR = OUTPUT_DIR / "latex_source"
SUPPLEMENT_DIR = OUTPUT_DIR / "supplementary_information"

TITLE_TEX = (
    "A registered genome-wide analysis identifies robust, context-dependent "
    "transcriptional responses of lychee to \\textit{Peronophythora litchii}"
)
TITLE_PLAIN = (
    "A registered genome-wide analysis identifies robust, context-dependent "
    "transcriptional responses of lychee to Peronophythora litchii"
)
SHORT_TITLE = "Registered analysis of lychee responses"
AUTHOR = "Eric Zhuang"
AFFILIATION = "NYU Langone Health, New York, New York, USA"
EMAIL = "eric.zhuang@nyulangone.org"
ORCID = "0009-0001-9050-0214"
ZENODO_DOI = "https://doi.org/10.5281/zenodo.22436625"
ARTICLE_TYPE = "Research"

KEYWORDS_TEX = (
    "\\textit{Litchi chinensis}, \\textit{Peronophythora litchii}, downy blight, "
    "RNA sequencing, cultivar-by-infection interaction, preregistered analysis, "
    "cross-context evaluation"
)

MAIN_FIGURES = [
    (MANUSCRIPT_DIR / "figures/figure1_study_design_qc.pdf", "Fig1.pdf"),
    (MANUSCRIPT_DIR / "figures/figure2_discovery_legacy.pdf", "Fig2.pdf"),
    (MANUSCRIPT_DIR / "figures/figure3_robustness_external.pdf", "Fig3.pdf"),
    (MANUSCRIPT_DIR / "figures/figure4_pathways_signatures.pdf", "Fig4.pdf"),
    (MANUSCRIPT_DIR / "figures/figure5_transcript_usage.pdf", "Fig5.pdf"),
    (MANUSCRIPT_DIR / "figures/figure6_orthogonal_tiers.pdf", "Fig6.pdf"),
]

SUPPLEMENTARY_FIGURES = [
    (
        PROJECT_ROOT / "results/figures/FigureS1_replicate_level_counts.pdf",
        "FigS1.pdf",
        "Replicate-level normalized counts for the two cross-context-supported genes, "
        "twelve Tier B genes, and two legacy highlights. Points are individual "
        "deposited libraries; bars show cultivar-treatment medians.",
    ),
    (
        PROJECT_ROOT / "results/figures/FigureS2_power_analysis.pdf",
        "FigS2.pdf",
        "Parametric detection probability for cultivar-by-infection effects under "
        "genome-wide discovery and candidate-family external adjustment. Curves show "
        "the overall result and mean-expression quartiles; the dashed line marks "
        "80\\% power.",
    ),
    (
        MANUSCRIPT_DIR / "figures/figureS3_exploratory_signature.pdf",
        "FigS3.pdf",
        "Exploratory PRJNA1090613 signed-signature estimate, shown separately because "
        "cohort metadata did not support its inclusion in the confirmatory evaluation.",
    ),
]

SOURCE_SUPPLEMENT = MANUSCRIPT_DIR / "lychee_unified_manuscript_supplement.zip"
EXPECTED_SUPPLEMENT_MD5 = "a7dce65e381f77ca365820311704b736"

DOIS = {
    2: "10.1371/journal.pone.0178245",
    3: "10.1016/j.foodchem.2009.04.074",
    4: "10.1038/s41588-021-00971-3",
    5: "10.1038/s41598-019-39100-w",
    6: "10.13925/j.cnki.gsxb.20220307",
    7: "10.1093/plphys/kiad311",
    9: "10.1093/jxb/erab549",
    11: "10.1093/bioinformatics/btu170",
    12: "10.1186/gb-2013-14-4-r36",
    13: "10.1093/bioinformatics/btu638",
    15: "10.1186/s13059-014-0550-8",
    16: "10.1016/j.molp.2020.06.009",
    18: "10.1093/nar/30.1.325",
    20: "10.1093/bioinformatics/bty560",
    21: "10.1093/bioinformatics/bts635",
    22: "10.1093/bioinformatics/btt656",
    23: "10.1038/nmeth.4197",
    24: "10.1093/bioinformatics/btaa222",
    25: "10.1093/bioinformatics/bty895",
    26: "10.1093/nar/gkad1052",
    27: "10.1038/nmeth.3176",
    28: "10.1093/nar/gks461",
    29: "10.1093/bioinformatics/btq401",
    30: "10.1101/060012",
    31: "10.12688/f1000research.8900.2",
    32: "10.1101/gr.133744.111",
    33: "10.1186/s13059-017-1277-0",
    34: "10.1093/bioinformatics/btp616",
    35: "10.1093/nar/gkac1052",
    36: "10.1093/bioinformatics/btu031",
    37: "10.1093/nar/gkad1059",
    38: "10.1186/1471-2105-11-165",
    39: "10.1093/bioinformatics/btr064",
    40: "10.1093/nar/gkz894",
    41: "10.12688/f1000research.29032.2",
}


def run(command: list[str], cwd: Path) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        command,
        cwd=cwd,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    if result.returncode:
        print(result.stdout)
        raise RuntimeError(f"Command failed ({result.returncode}): {' '.join(command)}")
    return result


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def md5(path: Path) -> str:
    digest = hashlib.md5()  # noqa: S324 - checksum comparison, not security
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_section(text: str, heading: str, next_heading: str) -> str:
    pattern = (
        rf"\\section\{{{re.escape(heading)}\}}\s*"
        rf"(.*?)"
        rf"(?=\\section\{{{re.escape(next_heading)}\}})"
    )
    match = re.search(pattern, text, flags=re.DOTALL)
    if not match:
        raise ValueError(f"Could not extract section {heading!r}")
    return match.group(1).strip()


def extract_abstract(text: str) -> str:
    match = re.search(
        r"\\begin\{abstract\}(.*?)\\end\{abstract\}", text, flags=re.DOTALL
    )
    if not match:
        raise ValueError("Abstract not found")
    abstract = match.group(1)
    abstract = abstract.split(r"\vspace{0.6em}", maxsplit=1)[0]
    abstract = abstract.replace(r"\noindent", "").strip()
    return re.sub(r"\s+", " ", abstract)


def abstract_word_count(abstract: str) -> int:
    plain = abstract.replace(r"\Pl{}", "P. litchii")
    plain = re.sub(r"\\textit\{([^{}]+)\}", r"\1", plain)
    plain = re.sub(r"\\[A-Za-z]+", " ", plain)
    plain = plain.replace("{", " ").replace("}", " ")
    return len(plain.split())


def extract_references(text: str) -> list[str]:
    match = re.search(
        r"\\section\*\{References\}\s*"
        r"\\begin\{enumerate\}.*?\n"
        r"(.*?)"
        r"\\end\{enumerate\}",
        text,
        flags=re.DOTALL,
    )
    if not match:
        raise ValueError("Reference list not found")
    return [
        item.strip()
        for item in re.findall(
            r"\\item\s+(.*?)(?=\n\\item|\Z)", match.group(1), flags=re.DOTALL
        )
    ]


def add_doi_links(items: list[str]) -> list[str]:
    enriched: list[str] = []
    for number, item in enumerate(items, start=1):
        item = re.sub(r"\s+", " ", item).strip()
        if number in DOIS:
            item = item.rstrip(".") + f". \\url{{https://doi.org/{DOIS[number]}}}."
        enriched.append(item)

    enriched.extend(
        [
            "National Center for Biotechnology Information. BioProject accession "
            "PRJNA450886. \\url{https://www.ncbi.nlm.nih.gov/bioproject/PRJNA450886}.",
            "National Center for Biotechnology Information. GEO accession GSE222651. "
            "\\url{https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE222651}.",
            "National Center for Biotechnology Information. GEO accession GSE222650. "
            "\\url{https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE222650}.",
            "National Center for Biotechnology Information. GEO accession GSE262200. "
            "\\url{https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE262200}.",
            "Zhuang E. Supplemental material for ``Cultivar-dependent transcriptional "
            "responses of lychee to \\textit{Peronophythora litchii}: a registered "
            "genome-wide analysis.'' Zenodo. 2026. "
            "\\url{https://doi.org/10.5281/zenodo.22436625}.",
        ]
    )
    return enriched


def adapt_body(text: str) -> str:
    replacements = {
        "figure1_study_design_qc.pdf": "Fig1.pdf",
        "figure2_discovery_legacy.pdf": "Fig2.pdf",
        "figure3_robustness_external.pdf": "Fig3.pdf",
        "figure4_pathways_signatures.pdf": "Fig4.pdf",
        "figure5_transcript_usage.pdf": "Fig5.pdf",
        "figure6_orthogonal_tiers.pdf": "Fig6.pdf",
        "Section~2.9": "the differential-transcript-usage analysis",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    # Avoid a long float queue in the referee-layout PDF while retaining each
    # figure and table at its logical location in the source.
    text = re.sub(r"\\begin\{figure\}\[[^]]+\]", r"\\begin{figure}[htbp]", text)
    text = re.sub(r"\\begin\{table\}\[[^]]+\]", r"\\begin{table}[htbp]", text)
    return text


def build_manuscript_tex() -> tuple[str, int]:
    source = SOURCE_TEX.read_text(encoding="utf-8")
    abstract = extract_abstract(source)
    introduction = adapt_body(extract_section(source, "Introduction", "Results"))
    results = adapt_body(extract_section(source, "Results", "Discussion"))
    discussion = adapt_body(extract_section(source, "Discussion", "Conclusions"))
    conclusions = adapt_body(
        extract_section(source, "Conclusions", "Materials and Methods")
    )

    methods_match = re.search(
        r"\\section\{Materials and Methods\}\s*"
        r"(.*?)"
        r"(?=\\section\*\{Declarations\})",
        source,
        flags=re.DOTALL,
    )
    if not methods_match:
        raise ValueError("Materials and Methods section not found")
    methods = adapt_body(methods_match.group(1).strip())

    references = add_doi_links(extract_references(source))
    references_tex = "\n\n".join(f"\\item {item}" for item in references)

    tex = rf"""% Springer Nature journal article template, version 3.1 (December 2024)
% Target journal: Discover Plants; article type: Research
% Compile with: pdflatex Discover_Plants_manuscript.tex (twice)
\documentclass[referee,pdflatex,sn-basic,Numbered]{{sn-jnl}}

\usepackage{{graphicx}}
\usepackage{{booktabs}}
\usepackage{{array}}
\usepackage{{amsmath}}
\usepackage{{enumitem}}
\usepackage{{microtype}}

\newcommand{{\Pl}}{{\textit{{P.~litchii}}}}
\newcommand{{\q}}{{\textit{{q}}}}
\setlength{{\emergencystretch}}{{3em}}
\raggedbottom

\begin{{document}}

\title[{SHORT_TITLE}]{{{TITLE_TEX}}}

\author*[1]{{\fnm{{Eric}} \sur{{Zhuang}}}}\email{{{EMAIL}}}
\affil*[1]{{\orgname{{NYU Langone Health}}, \orgaddress{{\city{{New York}}, \state{{New York}}, \country{{USA}}}}}}

\abstract{{{abstract}}}

\keywords{{{KEYWORDS_TEX}}}

\maketitle

\section{{Introduction}}\label{{sec:introduction}}

{introduction}

\section{{Materials and methods}}\label{{sec:methods}}

{methods}

\section{{Results}}\label{{sec:results}}

{results}

\section{{Discussion}}\label{{sec:discussion}}

{discussion}

\section{{Conclusions}}\label{{sec:conclusions}}

{conclusions}

\backmatter

\bmhead{{Supplementary information}}

Supplementary Information 1 provides an inventory of the supporting files and Figures S1--S3. Supplementary Data 2 provides machine-readable Tables S1--S18, figure source data, the registered protocol and amendment log, analysis code, workflows, configurations, tests, metadata, and environment specifications. The same complete archive is permanently available on Zenodo [46].

\bmhead{{Acknowledgements}}

OpenAI Codex was used to assist with language editing and journal-specific document formatting. It was not used to generate data, perform analyses, or determine scientific interpretations or conclusions. The author reviewed and accepts responsibility for all content.

\section*{{Declarations}}

\subsection*{{Funding}}

This work received no external funding.

\subsection*{{Competing interests}}

The author declares no competing interests.

\subsection*{{Ethics approval and consent to participate}}

Not applicable. This study reanalyzed publicly available plant sequencing datasets and involved no new experiments with humans, human tissue, animals, or field sampling.

\subsection*{{Consent for publication}}

Not applicable.

\subsection*{{Data availability}}

All analyzed sequencing data are publicly available under PRJNA830488/GSE201243, PRJNA450886, PRJNA922966/GSE222651, PRJNA922965/GSE222650, and PRJNA1090613/GSE262200 [8,42--45]. Complete result tables, supplementary figures, and tab-separated figure source data are archived under CC BY 4.0 on Zenodo [46].

\subsection*{{Code availability}}

The registered protocol, amendment log, analysis source code, Snakemake workflows, configuration files, tests, metadata, and environment specifications are archived under CC BY 4.0 at \url{{{ZENODO_DOI}}} [46].

\subsection*{{Materials availability}}

Not applicable. No new biological materials were generated in this study.

\subsection*{{Author contributions}}

Eric Zhuang: conceptualization, data curation, formal analysis, investigation, methodology, project administration, software, validation, visualization, writing---original draft, and writing---review and editing. The author approved the submitted version and accepts responsibility for the integrity of the work.

\section*{{References}}

\begin{{enumerate}}[label={{[\arabic*]}},itemsep=2pt,leftmargin=2.6em]
{references_tex}
\end{{enumerate}}

\end{{document}}
"""
    abstract_words = abstract_word_count(abstract)
    return tex, abstract_words


def build_supplement_tex() -> str:
    figures = []
    for _, filename, caption in SUPPLEMENTARY_FIGURES:
        figures.append(
            rf"""\begin{{figure}}[p]
\centering
\includegraphics[width=0.96\textwidth,height=0.72\textheight,keepaspectratio]{{{filename}}}
\caption{{{caption}}}
\end{{figure}}
\clearpage"""
        )
    figure_block = "\n\n".join(figures)
    return rf"""\documentclass[12pt]{{article}}
\usepackage[margin=1in]{{geometry}}
\usepackage{{graphicx}}
\usepackage[hidelinks]{{hyperref}}
\usepackage{{microtype}}
\setlength{{\parindent}}{{0pt}}
\setlength{{\parskip}}{{0.6em}}
\begin{{document}}

\begin{{center}}
{{\Large\bfseries Supplementary Information 1}}\\[1.2em]
{{\large {TITLE_TEX}}}\\[1.2em]
{AUTHOR}\\
{AFFILIATION}
\end{{center}}

\section*{{Contents}}

This file contains Figures S1--S3 and their captions. Supplementary Data 2, supplied as \texttt{{Discover\_Plants\_Data\_and\_Code.zip}}, contains machine-readable Tables S1--S18, all figure source data, the registered protocol and amendment log, analysis code, workflows, configurations, tests, metadata, and environment specifications. The complete archive is also permanently available at \url{{{ZENODO_DOI}}}.

\begin{{itemize}}
\item Figure S1: replicate-level normalized counts.
\item Figure S2: conditional parametric power curves.
\item Figure S3: exploratory signed-signature estimate.
\end{{itemize}}

\clearpage
\setcounter{{figure}}{{0}}
\renewcommand{{\thefigure}}{{S\arabic{{figure}}}}

{figure_block}

\end{{document}}
"""


def build_cover_letter_tex() -> str:
    display_date = date.today().strftime("%B %d, %Y").replace(" 0", " ")
    return rf"""\documentclass[11pt]{{letter}}
\usepackage[margin=1in]{{geometry}}
\usepackage[hidelinks]{{hyperref}}
\setlength{{\parskip}}{{0.55em}}
\signature{{Eric Zhuang\\NYU Langone Health\\\href{{mailto:{EMAIL}}}{{{EMAIL}}}}}
\address{{Eric Zhuang\\NYU Langone Health\\New York, New York, USA\\{EMAIL}}}
\begin{{document}}
\begin{{letter}}{{Editor-in-Chief\\Discover Plants}}
\opening{{Dear Editor-in-Chief,}}

Please consider the manuscript ``{TITLE_TEX}'' for publication in \textit{{Discover Plants}} as a \textbf{{Research}} article.

This study integrates five public RNA-sequencing cohorts to identify cultivar-dependent transcriptional responses of lychee to the oomycete pathogen \textit{{Peronophythora litchii}}. A prospectively registered analysis separated genome-wide discovery from cross-context evaluation and applied independent quantification, an alternative statistical framework, filter sensitivity, leave-one-library-out refitting, mapping controls, pathway analysis, and transcript-usage analysis. The study identifies 206 quality-controlled candidates, a stable core of 16 genes, 12 prioritized Tier B genes, and two genes supported across an independent tissue and cultivar contrast.

The manuscript fits the journal's coverage of plant pathology, plant genomics, plant molecular biology, and plant--microbe interactions. Its contribution is a new biological analysis and openly reproducible candidate resource rather than the generation of a new sequencing cohort. All analyzed data are public, and the registered protocol, complete results, figure source data, code, workflows, tests, and computational environments are archived at \url{{{ZENODO_DOI}}}.

The word ``registered'' in the title refers to the prospectively time-stamped analytical protocol governing the confirmatory analyses. The manuscript is submitted as a Research article, not under the journal's Registered Report format. No new wet-laboratory experiments were conducted, and no causal or functional claims are made without experimental support.

This manuscript has not been published as a peer-reviewed article and is not under consideration elsewhere. The author approved the submitted version and accepts responsibility for the work. The study received no external funding, the author declares no competing interests, and no human- or animal-research approval was required.

Thank you for considering this work for \textit{{Discover Plants}}.

\closing{{Sincerely,}}
\end{{letter}}
\end{{document}}
"""


def add_mixed_italic_paragraph(document: Document, prefix: str, title: str, suffix: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(prefix)
    run = paragraph.add_run(title)
    run.italic = True
    paragraph.add_run(suffix)


def build_cover_letter_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(8)

    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(f"{AUTHOR}\n{AFFILIATION}\n{EMAIL}\nORCID: {ORCID}")

    document.add_paragraph(date.today().strftime("%B %d, %Y").replace(" 0", " "))
    document.add_paragraph("Editor-in-Chief\nDiscover Plants")
    document.add_paragraph("Dear Editor-in-Chief,")

    add_mixed_italic_paragraph(
        document,
        "Please consider the manuscript \"",
        TITLE_PLAIN,
        '\" for publication in Discover Plants as a Research article.',
    )

    document.add_paragraph(
        "This study integrates five public RNA-sequencing cohorts to identify "
        "cultivar-dependent transcriptional responses of lychee to the oomycete "
        "pathogen Peronophythora litchii. A prospectively registered analysis "
        "separated genome-wide discovery from cross-context evaluation and applied "
        "independent quantification, an alternative statistical framework, filter "
        "sensitivity, leave-one-library-out refitting, mapping controls, pathway "
        "analysis, and transcript-usage analysis. The study identifies 206 "
        "quality-controlled candidates, a stable core of 16 genes, 12 prioritized "
        "Tier B genes, and two genes supported across an independent tissue and "
        "cultivar contrast."
    )
    document.add_paragraph(
        "The manuscript fits the journal's coverage of plant pathology, plant "
        "genomics, plant molecular biology, and plant-microbe interactions. Its "
        "contribution is a new biological analysis and openly reproducible candidate "
        "resource rather than the generation of a new sequencing cohort. All analyzed "
        "data are public, and the registered protocol, complete results, figure source "
        f"data, code, workflows, tests, and computational environments are archived at {ZENODO_DOI}."
    )
    document.add_paragraph(
        'The word "registered" in the title refers to the prospectively time-stamped '
        "analytical protocol governing the confirmatory analyses. The manuscript is "
        "submitted as a Research article, not under the journal's Registered Report "
        "format. No new wet-laboratory experiments were conducted, and no causal or "
        "functional claims are made without experimental support."
    )
    document.add_paragraph(
        "This manuscript has not been published as a peer-reviewed article and is not "
        "under consideration elsewhere. The author approved the submitted version and "
        "accepts responsibility for the work. The study received no external funding, "
        "the author declares no competing interests, and no human- or animal-research "
        "approval was required."
    )
    document.add_paragraph("Thank you for considering this work for Discover Plants.")
    document.add_paragraph(f"Sincerely,\n\n{AUTHOR}\n{EMAIL}")
    document.save(path)


def build_metadata(abstract: str, abstract_words: int) -> str:
    plain_abstract = abstract.replace(r"\Pl{}", "*P. litchii*")
    plain_abstract = re.sub(r"\\textit\{([^{}]+)\}", r"*\1*", plain_abstract)
    plain_abstract = plain_abstract.replace("--", "–")
    return f"""# Discover Plants submission metadata

## Article type

Research

The word “registered” describes the time-stamped analysis protocol. This is not a Registered Report submission.

## Title

{TITLE_PLAIN}

## Short title

{SHORT_TITLE}

## Abstract (approximately {abstract_words} words)

{plain_abstract}

## Keywords

1. *Litchi chinensis*
2. *Peronophythora litchii*
3. downy blight
4. RNA sequencing
5. cultivar-by-infection interaction
6. preregistered analysis
7. cross-context evaluation

## Author and corresponding author

- Full name: {AUTHOR}
- Affiliation: {AFFILIATION}
- Email: {EMAIL}
- ORCID iD: {ORCID}
- Corresponding author: Yes
- Submitting author: Yes

## Closest scope areas

- Plant pathology
- Plant genomics and evolution
- Plant molecular biology
- Plant–microbe interactions
- Crop science

## Author contributions

Eric Zhuang: conceptualization, data curation, formal analysis, investigation, methodology, project administration, software, validation, visualization, writing—original draft, and writing—review and editing.

## Funding

This work received no external funding.

## Competing interests

The author declares no competing interests.

## Ethics approval and consent to participate

Not applicable. This study reanalyzed publicly available plant sequencing datasets and involved no new experiments with humans, human tissue, animals, or field sampling.

## Consent for publication

Not applicable.

## Materials availability

Not applicable. No new biological materials were generated.

## Data availability

All analyzed sequencing data are publicly available under PRJNA830488/GSE201243, PRJNA450886, PRJNA922966/GSE222651, PRJNA922965/GSE222650, and PRJNA1090613/GSE262200. Complete result tables, supplementary figures, and figure source data are archived under CC BY 4.0 at <{ZENODO_DOI}>.

## Code availability

The registered protocol, amendment log, source code, Snakemake workflows, configurations, tests, metadata, and environment specifications are archived under CC BY 4.0 at <{ZENODO_DOI}>.

## Artificial-intelligence disclosure

OpenAI Codex was used to assist with language editing and journal-specific document formatting. It was not used to generate data, perform analyses, or determine scientific interpretations or conclusions. The author reviewed and accepts responsibility for all content.

## Suggested licence

CC BY 4.0, unless an institutional or funder requirement calls for another licence offered by the journal.

## Open-access charge note

The journal currently states that Springer Nature will cover the APC only for articles accepted by 31 December 2026. Acceptance—not submission—must occur by that date. Confirm this in the portal before final submission.
"""


def build_readme() -> str:
    return f"""# Discover Plants submission package

Prepared for the Springer Nature Snapp submission system from `../manuscript_v2.tex`.

## Recommended upload files (Word route)

1. `Discover_Plants_manuscript.docx` — **Manuscript**. This editable Word file contains the complete article, Figures 1–6, Tables 1–4, references, and declarations.
2. `Discover_Plants_cover_letter.docx` — **Cover letter**.
3. `supplementary_information/Discover_Plants_Supplementary_Information.docx` — **Supplementary Information 1**. It contains the supplement inventory and Figures S1–S3.
4. `supplementary_information/Discover_Plants_Data_and_Code.zip` — **Supplementary Data 2**. It contains Tables S1–S18, figure source data, the protocol, code, workflows, tests, metadata, and environments.

Do not upload both Word and LaTeX/PDF versions of the same item unless the portal explicitly asks. The matching PDFs and `Discover_Plants_LaTeX_source.zip` remain in this directory as visual previews and alternate source files.

## Portal selections

- Journal: Discover Plants
- Article type: Research
- Publishing model: Open access
- Suggested licence: CC BY 4.0 unless another licence is required

Copy title, abstract, keywords, author details, declarations, and scope terms from `submission_metadata.md`.

## Author checks before upload

- Confirm that the affiliation `{AFFILIATION}` is sufficiently complete. Add department/division and postal code if desired.
- Confirm the corresponding-author email `{EMAIL}` and ORCID `{ORCID}`.
- Confirm that the manuscript is no longer under consideration at any other journal.
- Review the AI-use disclosure and revise it if it does not accurately describe the author's use.
- Inspect every DOCX in Microsoft Word or the submission portal's converted proof, especially figure labels, wide tables, mathematical symbols, and page breaks.
- Confirm the current APC waiver in the live portal. The published journal page says the waiver applies only if the article is accepted by 31 December 2026.
- Do not describe this submission as a Registered Report. It is a Research article containing a prospectively registered confirmatory analysis.

## Source-package notes

- The recommended manuscript and supplementary-information uploads are editable Word files. Both contain embedded high-resolution figures; no separate figure upload should be needed unless the portal requests one.
- The Word manuscript uses continuous line numbering (displayed every five lines), double-spaced body text, Word heading styles, editable Word tables, numbered references, and accessible figure descriptions.
- The package uses Springer Nature's official `sn-jnl.cls` version 3.1 (December 2024) in `pdflatex` mode.
- Snapp requires all LaTeX files and figures in one ZIP directory; the generated ZIP is flat.
- References are included directly in the TeX source to avoid bibliography-conversion failures.
- Main and supplementary materials cite the current version-specific Zenodo record: <{ZENODO_DOI}>.
- The Zenodo record retains an earlier manuscript title. Before submission, either update its metadata to the current title or confirm that retaining the earlier title is acceptable; the DOI and archived contents are otherwise current and verified.
- `VALIDATION_REPORT.txt` records automated compilation and package checks.
- `UPLOAD_FILE_MANIFEST_SHA256.tsv` records exact checksums for the recommended Word-route upload files.
"""


def compile_tex(tex_path: Path) -> str:
    final_output = ""
    for _ in range(2):
        result = run(
            ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex_path.name],
            cwd=tex_path.parent,
        )
        final_output = result.stdout
    return final_output


def make_flat_zip(zip_path: Path, files: list[Path]) -> None:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, arcname=path.name)


def validate_source_zip(zip_path: Path) -> str:
    with TemporaryDirectory(prefix="discover_plants_validate_") as temp_name:
        temp_dir = Path(temp_name)
        with zipfile.ZipFile(zip_path) as archive:
            archive.extractall(temp_dir)
            names = archive.namelist()
        if any("/" in name.rstrip("/") for name in names):
            raise ValueError("LaTeX ZIP is not flat")
        result = None
        for _ in range(2):
            result = run(
                [
                    "pdflatex",
                    "-interaction=nonstopmode",
                    "-halt-on-error",
                    "Discover_Plants_manuscript.tex",
                ],
                cwd=temp_dir,
            )
        assert result is not None
        return result.stdout


def clean_latex_auxiliary(directory: Path, stems: list[str]) -> None:
    extensions = [".aux", ".log", ".out"]
    for stem in stems:
        for extension in extensions:
            path = directory / f"{stem}{extension}"
            if path.exists():
                path.unlink()


def main() -> None:
    required = [
        SOURCE_TEX,
        TEMPLATE_DIR / "sn-jnl.cls",
        TEMPLATE_DIR / "sn-basic.bst",
        SOURCE_SUPPLEMENT,
        *(path for path, _ in MAIN_FIGURES),
        *(path for path, _, _ in SUPPLEMENTARY_FIGURES),
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required files:\n" + "\n".join(missing))

    if md5(SOURCE_SUPPLEMENT) != EXPECTED_SUPPLEMENT_MD5:
        raise ValueError("Local supplement does not match the current Zenodo archive MD5")

    LATEX_DIR.mkdir(parents=True, exist_ok=True)
    SUPPLEMENT_DIR.mkdir(parents=True, exist_ok=True)

    manuscript_tex, abstract_words = build_manuscript_tex()
    if abstract_words >= 250:
        raise ValueError(f"Abstract has {abstract_words} words; journal limit is <250")

    manuscript_path = LATEX_DIR / "Discover_Plants_manuscript.tex"
    manuscript_path.write_text(manuscript_tex, encoding="ascii")
    shutil.copy2(TEMPLATE_DIR / "sn-jnl.cls", LATEX_DIR / "sn-jnl.cls")
    shutil.copy2(TEMPLATE_DIR / "sn-basic.bst", LATEX_DIR / "sn-basic.bst")
    for source, filename in MAIN_FIGURES:
        shutil.copy2(source, LATEX_DIR / filename)

    main_compile_log = compile_tex(manuscript_path)
    manuscript_pdf = LATEX_DIR / "Discover_Plants_manuscript.pdf"
    shutil.copy2(manuscript_pdf, OUTPUT_DIR / manuscript_pdf.name)

    source_files = [
        manuscript_path,
        LATEX_DIR / "sn-jnl.cls",
        LATEX_DIR / "sn-basic.bst",
        *(LATEX_DIR / filename for _, filename in MAIN_FIGURES),
    ]
    source_zip = OUTPUT_DIR / "Discover_Plants_LaTeX_source.zip"
    make_flat_zip(source_zip, source_files)
    zip_compile_log = validate_source_zip(source_zip)

    supplement_tex_path = SUPPLEMENT_DIR / "Discover_Plants_Supplementary_Information.tex"
    supplement_tex_path.write_text(build_supplement_tex(), encoding="ascii")
    for source, filename, _ in SUPPLEMENTARY_FIGURES:
        shutil.copy2(source, SUPPLEMENT_DIR / filename)
    supplement_compile_log = compile_tex(supplement_tex_path)

    supplement_archive = SUPPLEMENT_DIR / "Discover_Plants_Data_and_Code.zip"
    shutil.copy2(SOURCE_SUPPLEMENT, supplement_archive)

    cover_tex_path = OUTPUT_DIR / "Discover_Plants_cover_letter.tex"
    cover_tex_path.write_text(build_cover_letter_tex(), encoding="ascii")
    cover_compile_log = compile_tex(cover_tex_path)
    build_cover_letter_docx(OUTPUT_DIR / "Discover_Plants_cover_letter.docx")

    from build_word_submission import build_word_files

    manuscript_docx = OUTPUT_DIR / "Discover_Plants_manuscript.docx"
    supplement_docx = SUPPLEMENT_DIR / "Discover_Plants_Supplementary_Information.docx"
    word_validation = build_word_files(
        manuscript_tex=manuscript_path,
        manuscript_docx=manuscript_docx,
        supplement_docx=supplement_docx,
        main_figures=MAIN_FIGURES,
        supplementary_figures=SUPPLEMENTARY_FIGURES,
        title_tex=TITLE_TEX,
        author=AUTHOR,
        affiliation=AFFILIATION,
        email=EMAIL,
        orcid=ORCID,
        zenodo_doi=ZENODO_DOI,
    )

    abstract = extract_abstract(SOURCE_TEX.read_text(encoding="utf-8"))
    (OUTPUT_DIR / "submission_metadata.md").write_text(
        build_metadata(abstract, abstract_words), encoding="utf-8"
    )
    (OUTPUT_DIR / "README_submission.md").write_text(
        build_readme(), encoding="utf-8"
    )

    upload_files = [
        manuscript_docx,
        OUTPUT_DIR / "Discover_Plants_cover_letter.docx",
        supplement_docx,
        supplement_archive,
    ]
    manifest_lines = ["sha256\tbytes\tfile"]
    for path in upload_files:
        manifest_lines.append(
            f"{sha256(path)}\t{path.stat().st_size}\t{path.relative_to(OUTPUT_DIR)}"
        )
    (OUTPUT_DIR / "UPLOAD_FILE_MANIFEST_SHA256.tsv").write_text(
        "\n".join(manifest_lines) + "\n", encoding="utf-8"
    )

    all_logs = "\n".join([main_compile_log, zip_compile_log, supplement_compile_log, cover_compile_log])
    failure_markers = ["LaTeX Error", "Undefined control sequence", "undefined references"]
    found_markers = [marker for marker in failure_markers if marker.lower() in all_logs.lower()]
    validation_lines = [
        "Discover Plants submission validation",
        f"Build date: {date.today().isoformat()}",
        f"Source manuscript: {SOURCE_TEX}",
        f"Article type: {ARTICLE_TYPE}",
        f"Abstract words (LaTeX-aware approximate count): {abstract_words} (<250 required)",
        "Main manuscript compiled with pdflatex: PASS",
        "Flat LaTeX source ZIP compiled independently: PASS",
        "Supplementary Information PDF compiled with pdflatex: PASS",
        "Cover letter PDF compiled with pdflatex: PASS",
        (
            "Word manuscript opened and structurally validated: PASS "
            f"({word_validation['manuscript']['images']} figures; "
            f"{word_validation['manuscript']['tables']} tables; "
            f"{word_validation['manuscript']['paragraphs']} paragraphs)"
        ),
        (
            "Word supplementary information opened and structurally validated: PASS "
            f"({word_validation['supplement']['images']} figures; "
            f"{word_validation['supplement']['paragraphs']} paragraphs)"
        ),
        f"Full supplement MD5 matches Zenodo record: PASS ({EXPECTED_SUPPLEMENT_MD5})",
        f"LaTeX failure markers: {'PASS' if not found_markers else 'CHECK ' + ', '.join(found_markers)}",
        f"Main source ZIP entries: {len(source_files)}",
        f"Upload files checksummed: {len(upload_files)}",
        "Manual visual inspection remains required before submission.",
    ]
    (OUTPUT_DIR / "VALIDATION_REPORT.txt").write_text(
        "\n".join(validation_lines) + "\n", encoding="utf-8"
    )

    clean_latex_auxiliary(LATEX_DIR, ["Discover_Plants_manuscript"])
    clean_latex_auxiliary(
        SUPPLEMENT_DIR, ["Discover_Plants_Supplementary_Information"]
    )
    clean_latex_auxiliary(OUTPUT_DIR, ["Discover_Plants_cover_letter"])

    print("Built Discover Plants submission package")
    print(f"Abstract words: {abstract_words}")
    for path in upload_files:
        print(f"{path.relative_to(OUTPUT_DIR)}\t{path.stat().st_size} bytes")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
